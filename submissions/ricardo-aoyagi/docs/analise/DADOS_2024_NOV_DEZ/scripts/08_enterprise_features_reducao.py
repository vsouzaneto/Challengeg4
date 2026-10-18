# Plano × features  +  dentro do Enterprise: reduções por feature, país, canal, indústria, cobrança e seats.
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada.
import sys
from pathlib import Path
from math import lgamma, exp
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '08_Plano_Features_Reducao_Enterprise.xlsx'
INI, FIM = pd.Timestamp('2024-11-01'), pd.Timestamp('2024-12-31'); JP = 7
FAIXAS = [0, 5, 10, 20, 40, 80, 10**4]; LF = ['01–05', '06–10', '11–20', '21–40', '41–80', '81 ou mais']
a = pd.read_csv(DADOS/'ravenstack_accounts.csv'); s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date']); f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv')
def fisher_p(a_, b_, c_, d_):
    n1, n2, m1 = a_ + b_, c_ + d_, a_ + c_; N = n1 + n2
    if N == 0 or n1 == 0 or n2 == 0 or m1 == 0 or m1 == N: return 1.0
    lc = lambda n, k: lgamma(n + 1) - lgamma(k + 1) - lgamma(n - k + 1)
    pmf = lambda x: exp(lc(n1, x) + lc(n2, m1 - x) - lc(N, m1))
    p0 = pmf(a_); lo, hi = max(0, m1 - n2), min(m1, n1)
    return min(1.0, sum(pmf(x) for x in range(lo, hi + 1) if pmf(x) <= p0 * (1 + 1e-7)))
def conf(k1, n1, k2, n2):
    if k1 + k2 < 10: return 'Poucos casos para concluir'
    return 'Sim' if fisher_p(k1, n1 - k1, k2, n2 - k2) < 0.05 else 'Não (pode ser coincidência)'
def tipo(r):
    o = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
    ativas = o[(o.start_date <= r.end_date) & (o.end_date.isna() | (o.end_date > r.end_date))]
    perto = o[(o.start_date - r.end_date).abs().dt.days <= JP]; depois = o[(o.start_date > r.end_date) & (o.start_date <= r.end_date + pd.Timedelta(days=JP))]
    if len(ativas) == 0 and len(depois) == 0: return 'Saída do cliente'
    return 'Migração' if len(perto) else 'Redução'
s['situacao'] = 'Ativa'; m = s.end_date.notna(); s.loc[m, 'situacao'] = s[m].apply(tipo, axis=1)
S = s.merge(a[['account_id', 'industry', 'country', 'referral_source']], on='account_id')
S['cobranca'] = S.billing_frequency.map({'annual': 'Anual', 'monthly': 'Mensal'}); S['faixa_seats'] = pd.cut(S.seats, FAIXAS, labels=LF).astype(str)
S['encerrada_novdez'] = np.where(S.end_date.between(INI, FIM), 'Sim', 'Não'); S['reduzida'] = np.where(S.situacao == 'Redução', 1, 0)
S['mrr_reduzido'] = S.mrr_amount * S.reduzida
fu = f.merge(S[['subscription_id', 'plan_tier', 'situacao']], on='subscription_id')
FS = fu.groupby(['subscription_id', 'plan_tier', 'situacao', 'feature_name']).agg(registros=('usage_id', 'size'), registros_com_erro=('error_count', lambda x: int((x > 0).sum())),
     registros_beta=('is_beta_feature', 'sum')).reset_index()
FS['feature_num'] = FS.feature_name.str.replace('feature_', '').astype(int)
FEATS = sorted(FS.feature_name.unique(), key=lambda z: int(z.split('_')[1])); PLANOS = ['Basic', 'Pro', 'Enterprise']

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
HD = PatternFill('solid', fgColor='1F2937'); HB = PatternFill('solid', fgColor='2A78D6'); HP = PatternFill('solid', fgColor='7C3AED'); HE = PatternFill('solid', fgColor='0F766E'); HR = PatternFill('solid', fgColor='B45309'); HX = PatternFill('solid', fgColor='374151')
YEL = PatternFill('solid', fgColor='FFF3CD'); GRY = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '"US$ "#,##0;-"US$ "#,##0;"-"'
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); p1 = S_('1_Plano_x_Feature'); p2 = S_('2_Reducao_por_plano'); p3 = S_('3_Enterprise_Features'); p4 = S_('4_Enterprise_Segmentos')
BS = {}
def py(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v, 'item') else v
def base(nome, df, cols, fmts):
    ws = wb.create_sheet(nome); ws.append(cols)
    for row in df[cols].itertuples(index=False): ws.append([py(v) for v in row])
    n = len(df) + 1; BS[nome] = (cols, n)
    for i, cn in enumerate(cols, 1):
        h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(22, len(cn) + 2))
        if cn in fmts:
            for r in range(2, n + 1): ws.cell(r, i).number_format = fmts[cn]
    tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
def Rg(b, col): cols, n = BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
base('Base_Assinaturas', S, ['subscription_id', 'account_id', 'industry', 'country', 'referral_source', 'plan_tier', 'cobranca', 'seats', 'faixa_seats', 'mrr_amount', 'start_date', 'end_date', 'situacao', 'reduzida', 'mrr_reduzido', 'encerrada_novdez'],
     {'start_date': 'dd/mm/yyyy', 'end_date': 'dd/mm/yyyy', 'mrr_amount': USD, 'mrr_reduzido': USD})
base('Base_Uso_Feature', FS, ['subscription_id', 'plan_tier', 'situacao', 'feature_name', 'feature_num', 'registros', 'registros_com_erro', 'registros_beta'], {})
def H(ws, row, labels, fills, h=44):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fill=None):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX; x.alignment = Alignment(vertical='center', wrap_text=True)
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x
def title(ws, t1, t2, merge='J'):
    ws['A1'] = t1; ws['A1'].font = Fn(bold=True, size=15); ws['A2'] = t2; ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells(f'A2:{merge}2'); ws.row_dimensions[2].height = 30
def note(ws, r, tx, merge='J', h=32, color=GRAY):
    ws.cell(r, 1, tx).font = Fn(italic=True, size=9, color=color); ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(f'A{r}:{merge}{r}'); ws.row_dimensions[r].height = h
UF = 'Base_Uso_Feature'; PL, SI, FN = Rg(UF, 'plan_tier'), Rg(UF, 'situacao'), Rg(UF, 'feature_name'); RG, RE, RBT = Rg(UF, 'registros'), Rg(UF, 'registros_com_erro'), Rg(UF, 'registros_beta')
BA = 'Base_Assinaturas'; BPL = Rg(BA, 'plan_tier'); BSI = Rg(BA, 'situacao')

# ---------- 1_Plano_x_Feature ----------
ws = p1; title(ws, 'Cada plano usa features diferentes? Erra mais em alguma?', 'Uso ligado à assinatura (sem olhar data, que não é confiável). "% das assinaturas" = quantas assinaturas do plano têm ao menos 1 registro da feature.')
H(ws, 4, ['Resumo', 'Basic', 'Pro', 'Enterprise'], [HD, HB, HP, HE]); r = 5
for lab, fx, fm in [('Assinaturas', lambda p: f'=COUNTIFS({BPL},"{p}")', '#,##0'), ('Registros de uso por assinatura', lambda p: f'=SUMIFS({RG},{PL},"{p}")/COUNTIFS({BPL},"{p}")', '0.0'),
                    ('% dos registros com erro', lambda p: f'=SUMIFS({RE},{PL},"{p}")/SUMIFS({RG},{PL},"{p}")', '0.0%'), ('% dos registros em feature beta', lambda p: f'=SUMIFS({RBT},{PL},"{p}")/SUMIFS({RG},{PL},"{p}")', '0.0%')]:
    C(ws, r, 1, lab, bold=True)
    for j, p in enumerate(PLANOS, 2): C(ws, r, j, fx(p), fm)
    r += 1
r += 1; ws.cell(r, 1, 'Por feature').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Feature', 'Basic: % das assinaturas', 'Pro: % das assinaturas', 'Enterprise: % das assinaturas', 'Basic: % registros com erro', 'Pro: % registros com erro', 'Enterprise: % registros com erro'], [HD, HB, HP, HE, HB, HP, HE]); hr = r; r += 1; f0 = r
for fe in FEATS:
    C(ws, r, 1, fe, bold=True)
    for j, p in enumerate(PLANOS):
        C(ws, r, 2 + j, f'=COUNTIFS({PL},"{p}",{FN},$A{r})/COUNTIFS({BPL},"{p}")', '0.0%')
        C(ws, r, 5 + j, f'=SUMIFS({RE},{PL},"{p}",{FN},$A{r})/SUMIFS({RG},{PL},"{p}",{FN},$A{r})', '0%')
    r += 1
fl = r - 1
note(ws, r + 1, 'Leitura: os três planos usam as mesmas features na mesma proporção (cada feature aparece em ~10–14% das assinaturas de qualquer plano) e erram na mesma proporção (~30% dos registros). Não há feature "do Enterprise". A marca de feature beta cai em ~10% dos registros de TODAS as features, sem identificar uma feature nova específica.', merge='G', h=44)
for col_, w in zip('ABCDEFG', [30, 12, 12, 12, 12, 12, 12]): ws.column_dimensions[col_].width = w
ws.freeze_panes = f'B{hr+1}'

# ---------- 2_Reducao_por_plano ----------
ws = p2; title(ws, 'O Enterprise reduz MAIS, ou cada redução Enterprise só VALE mais?', 'Período todo (jan/23–dez/24) e só nov/dez/24. Situação de cada assinatura pela mesma regra das planilhas 04/06/07.')
H(ws, 4, ['Plano', 'Assinaturas', 'Encerradas em redução', 'Reduzidas a cada 100', 'Reduções: diferente dos outros planos?', 'MRR médio por assinatura', 'MRR perdido em reduções', 'Nov/dez: reduções', 'Nov/dez: MRR perdido em reduções'],
  [HD, HD, HR, HR, HX, HD, HR, HR, HR]); r = 5
for p in PLANOS:
    k1 = int(((S.plan_tier == p) & (S.reduzida == 1)).sum()); n1 = int((S.plan_tier == p).sum()); k2 = int(S.reduzida.sum()) - k1; n2 = len(S) - n1
    C(ws, r, 1, p, bold=True); C(ws, r, 2, f'=COUNTIFS({BPL},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({BPL},A{r},{BSI},"Redução")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}*100', '0.0')
    C(ws, r, 5, conf(k1, n1, k2, n2)); C(ws, r, 6, f'=AVERAGEIFS({Rg(BA,"mrr_amount")},{BPL},A{r},{Rg(BA,"mrr_amount")},">0")', USD); C(ws, r, 7, f'=SUMIFS({Rg(BA,"mrr_reduzido")},{BPL},A{r})', USD)
    C(ws, r, 8, f'=COUNTIFS({BPL},A{r},{BSI},"Redução",{Rg(BA,"encerrada_novdez")},"Sim")', '#,##0'); C(ws, r, 9, f'=SUMIFS({Rg(BA,"mrr_reduzido")},{BPL},A{r},{Rg(BA,"encerrada_novdez")},"Sim")', USD); r += 1
C(ws, r, 1, 'Total', bold=True, fill=GRY)
for j, fm in [(2, '#,##0'), (3, '#,##0'), (7, USD), (8, '#,##0'), (9, USD)]: C(ws, r, j, f'=SUM({L(j)}5:{L(j)}{r-1})', fm, True, GRY)
C(ws, r, 4, f'=C{r}/B{r}*100', '0.0', True, GRY)
note(ws, r + 2, 'Leitura: Enterprise reduz na MESMA frequência que Basic e Pro (~5 a cada 100 assinaturas). A perda em US$ é maior só porque cada assinatura Enterprise vale muito mais (preço por seat de US$ 199 contra 49 e 19). O problema do Enterprise é de VALOR, não de frequência.', merge='I', h=40)
for col_, w in zip('ABCDEFGHI', [12, 11, 12, 11, 24, 14, 15, 11, 15]): ws.column_dimensions[col_].width = w
ch = BarChart(); ch.type = 'col'; ch.title = 'Reduzidas a cada 100 assinaturas, por plano'; ch.height = 6.5; ch.width = 12; ch.legend = None
ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=7), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7)); ch.series[0].graphicalProperties.solidFill = 'B45309'; ws.add_chart(ch, 'A12')
ch = BarChart(); ch.type = 'col'; ch.title = 'MRR perdido em reduções, por plano (US$)'; ch.height = 6.5; ch.width = 12; ch.legend = None
ch.add_data(Reference(ws, min_col=7, min_row=4, max_row=7), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7)); ch.series[0].graphicalProperties.solidFill = 'B45309'; ch.y_axis.number_format = '#,##0'; ws.add_chart(ch, 'F12')

# ---------- 3_Enterprise_Features ----------
EN = S[S.plan_tier == 'Enterprise']; nR = int((EN.situacao == 'Redução').sum()); nA = int((EN.situacao == 'Ativa').sum())
FE = FS[FS.plan_tier == 'Enterprise']
ws = p3; title(ws, 'Enterprise: as assinaturas que reduziram usavam alguma feature diferente? Erravam mais?', f'Compara as {nR} assinaturas Enterprise encerradas em redução (período todo) com as {nA} Enterprise ativas. Ordene pelas colunas para achar destaques.')
H(ws, 4, ['Feature', 'Reduzidas que usaram', '% das reduzidas', '% das ativas', 'Uso: reduzidas ÷ ativas', 'Uso: diferença confiável?', '% registros com erro (reduzidas)', '% registros com erro (ativas)', 'Erro: reduzidas ÷ ativas', 'Erro: diferença confiável?'],
  [HD, HR, HR, HE, HX, HX, HR, HE, HX, HX]); r = 5; f0 = r
for fe in FEATS:
    x = FE[FE.feature_name == fe]; uR = int(x[x.situacao == 'Redução'].subscription_id.nunique()); uA = int(x[x.situacao == 'Ativa'].subscription_id.nunique())
    rR, eR = int(x[x.situacao == 'Redução'].registros.sum()), int(x[x.situacao == 'Redução'].registros_com_erro.sum()); rA, eA = int(x[x.situacao == 'Ativa'].registros.sum()), int(x[x.situacao == 'Ativa'].registros_com_erro.sum())
    C(ws, r, 1, fe, bold=True)
    C(ws, r, 2, f'=COUNTIFS({PL},"Enterprise",{SI},"Redução",{FN},A{r})', '#,##0'); C(ws, r, 3, f'=B{r}/COUNTIFS({BPL},"Enterprise",{BSI},"Redução")', '0%')
    C(ws, r, 4, f'=COUNTIFS({PL},"Enterprise",{SI},"Ativa",{FN},A{r})/COUNTIFS({BPL},"Enterprise",{BSI},"Ativa")', '0%'); C(ws, r, 5, f'=IF(D{r}>0,C{r}/D{r},"")', '0.00"x"'); C(ws, r, 6, conf(uR, nR, uA, nA))
    C(ws, r, 7, f'=IFERROR(SUMIFS({RE},{PL},"Enterprise",{SI},"Redução",{FN},A{r})/SUMIFS({RG},{PL},"Enterprise",{SI},"Redução",{FN},A{r}),"")', '0%')
    C(ws, r, 8, f'=SUMIFS({RE},{PL},"Enterprise",{SI},"Ativa",{FN},A{r})/SUMIFS({RG},{PL},"Enterprise",{SI},"Ativa",{FN},A{r})', '0%')
    C(ws, r, 9, f'=IF(AND(ISNUMBER(G{r}),H{r}>0),G{r}/H{r},"")', '0.00"x"'); C(ws, r, 10, conf(eR, rR, eA, rA)); r += 1
fl = r - 1
for cc in 'FJ': ws.conditional_formatting.add(f'{cc}{f0}:{cc}{fl}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
ws.auto_filter.ref = f'A4:J{fl}'
rr = r + 1; H(ws, rr, ['Resumo', 'Reduzidas', 'Ativas', '', '', 'Diferença confiável?'], [HD, HR, HE, HD, HD, HX]); rr += 1
for lab, fx, col in [('% dos registros com erro', lambda t: f'=SUMIFS({RE},{PL},"Enterprise",{SI},"{t}")/SUMIFS({RG},{PL},"Enterprise",{SI},"{t}")', 'erro'), ('% dos registros em beta', lambda t: f'=SUMIFS({RBT},{PL},"Enterprise",{SI},"{t}")/SUMIFS({RG},{PL},"Enterprise",{SI},"{t}")', 'beta')]:
    C(ws, rr, 1, lab, bold=True); C(ws, rr, 2, fx('Redução'), '0.0%'); C(ws, rr, 3, fx('Ativa'), '0.0%')
    kR = int(FE[FE.situacao == 'Redução']['registros_com_erro' if col == 'erro' else 'registros_beta'].sum()); tR = int(FE[FE.situacao == 'Redução'].registros.sum())
    kA = int(FE[FE.situacao == 'Ativa']['registros_com_erro' if col == 'erro' else 'registros_beta'].sum()); tA = int(FE[FE.situacao == 'Ativa'].registros.sum())
    C(ws, rr, 6, conf(kR, tR, kA, tA)); rr += 1
note(ws, rr + 1, 'Leitura: com ~90 assinaturas reduzidas, cada feature tem 3 a 17 casos. Algumas aparecem mais (ou menos) nas reduzidas, mas em 40 features é esperado que 2 deem "Sim" por acaso. Ver a coluna "diferença confiável?" antes de apontar uma feature como causa. No geral, reduzidas e ativas erram na mesma proporção e usam beta na mesma proporção.', merge='J', h=44)
for col_, w in zip('ABCDEFGHIJ', [24, 11, 10, 10, 10, 22, 12, 12, 10, 22]): ws.column_dimensions[col_].width = w
ws.freeze_panes = 'B5'
featsim = [(fe, 'uso') for fe, v in zip(FEATS, [ws.cell(i, 6).value for i in range(f0, fl + 1)]) if v == 'Sim'] + [(fe, 'erro') for fe, v in zip(FEATS, [ws.cell(i, 10).value for i in range(f0, fl + 1)]) if v == 'Sim']

# ---------- 4_Enterprise_Segmentos ----------
ws = p4; title(ws, 'Enterprise: algum país, canal, indústria, cobrança ou faixa de seats reduz mais?', 'Reduzidas a cada 100 assinaturas Enterprise. "Confiável?" compara a categoria com as demais categorias do Enterprise. Período todo e nov/dez.')
r = 4; segsim = []
for col, lab, vals in [('country', 'País', sorted(EN.country.unique())), ('referral_source', 'Canal de aquisição', sorted(EN.referral_source.unique())), ('industry', 'Indústria', sorted(EN.industry.unique())),
                       ('cobranca', 'Cobrança', ['Anual', 'Mensal']), ('faixa_seats', 'Faixa de seats', LF)]:
    ws.cell(r, 1, lab).font = Fn(bold=True, size=12); r += 1
    H(ws, r, [lab, 'Assinaturas Enterprise', 'Reduzidas', 'Reduzidas a cada 100', 'Diferente das demais?', 'MRR perdido em reduções', 'Nov/dez: reduzidas', 'Nov/dez: MRR perdido'], [HD, HE, HR, HR, HX, HR, HR, HR]); r += 1; r0 = r
    for v in vals:
        msk = EN[col] == v; k1 = int((msk & (EN.reduzida == 1)).sum()); n1 = int(msk.sum()); k2 = int(EN.reduzida.sum()) - k1; n2 = len(EN) - n1
        cd = f'{BPL},"Enterprise",{Rg(BA, col)},"{v}"'
        C(ws, r, 1, v, bold=True); C(ws, r, 2, f'=COUNTIFS({cd})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({cd},{BSI},"Redução")', '#,##0'); C(ws, r, 4, f'=IF(B{r}>0,C{r}/B{r}*100,"")', '0.0')
        cv = conf(k1, n1, k2, n2); C(ws, r, 5, cv); C(ws, r, 6, f'=SUMIFS({Rg(BA,"mrr_reduzido")},{cd})', USD)
        C(ws, r, 7, f'=COUNTIFS({cd},{BSI},"Redução",{Rg(BA,"encerrada_novdez")},"Sim")', '#,##0'); C(ws, r, 8, f'=SUMIFS({Rg(BA,"mrr_reduzido")},{cd},{Rg(BA,"encerrada_novdez")},"Sim")', USD)
        if cv == 'Sim': segsim.append(f'{lab} {v}: {k1/n1*100:.1f} a cada 100 (demais: {k2/n2*100:.1f})')
        r += 1
    C(ws, r, 1, 'Enterprise total', bold=True, fill=GRY); C(ws, r, 2, f'=SUM(B{r0}:B{r-1})', '#,##0', True, GRY); C(ws, r, 3, f'=SUM(C{r0}:C{r-1})', '#,##0', True, GRY); C(ws, r, 4, f'=C{r}/B{r}*100', '0.0', True, GRY)
    C(ws, r, 6, f'=SUM(F{r0}:F{r-1})', USD, True, GRY); C(ws, r, 7, f'=SUM(G{r0}:G{r-1})', '#,##0', True, GRY); C(ws, r, 8, f'=SUM(H{r0}:H{r-1})', USD, True, GRY); r += 2
ws.conditional_formatting.add(f'E4:E{r}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
note(ws, r, 'Leitura: veja as linhas com "Sim" (destaques confiáveis). Em nov/dez, as reduções Enterprise por categoria são poucas (de 0 a ~10): servem para localizar onde está o dinheiro perdido, não para provar causa.', merge='H')
for col_, w in zip('ABCDEFGH', [18, 12, 10, 11, 24, 15, 11, 14]): ws.column_dimensions[col_].width = w

# ---------- LEIA-ME ----------
red = S.groupby('plan_tier').reduzida.mean() * 100
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
Lm = [('Plano × features e o padrão das reduções no Enterprise', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. "Redução" = assinatura encerrada sem reposição em até 7 dias, com a conta seguindo cliente (mesma regra das planilhas 04, 06 e 07).', dict(italic=True, color=GRAY)),
      ('', {}), ('PERGUNTAS', dict(bold=True, size=12)),
      ('1. Os planos usam features diferentes? Alguma feature (nova ou com erro) pode estar levando o Enterprise a reduzir?  → abas 1 e 3', {}),
      ('2. O Enterprise reduz mais que os outros planos?  → aba 2', {}),
      ('3. Dentro do Enterprise, algum país, canal, indústria, cobrança ou faixa de seats reduz mais?  → aba 4', {}),
      ('', {}), ('RESPOSTAS', dict(bold=True, size=12)),
      ('1. Os três planos usam as mesmas features, na mesma proporção, e erram na mesma proporção (~30% dos registros). A marca "beta" aparece em ~10% dos registros de todas as features, então não dá para identificar uma "feature nova" nesta base.', {}),
      ('   Entre as Enterprise reduzidas vs ativas: feature_18, feature_24 e feature_28 aparecem MENOS nas reduzidas (não são causa); feature_13 e feature_39 têm mais registros com erro nas reduzidas, mas com poucos casos. Em 80 comparações (40 features × uso e erro), cerca de 4 "Sim" são esperados só por acaso, que é o que aparece. No geral, reduzidas e ativas erram e usam beta na mesma proporção.', {}),
      (f'2. NÃO. Reduzidas a cada 100 assinaturas: Basic {red["Basic"]:.1f}, Pro {red["Pro"]:.1f}, Enterprise {red["Enterprise"]:.1f}. O Enterprise perde mais dinheiro porque cada assinatura vale muito mais, não porque reduz com mais frequência.', {}),
      ('3. Dentro do Enterprise, nenhum país, canal, indústria, cobrança ou faixa de seats reduz de forma diferente das demais com segurança. Canadá e Índia aparecem abaixo e Austrália, Alemanha, França e Reino Unido um pouco acima, e cobrança mensal um pouco acima da anual, mas todas essas diferenças podem ser coincidência. Em US$, a perda acompanha onde estão as assinaturas Enterprise (sobretudo EUA).', {}),
      ('', {}), ('IMPLICAÇÃO', dict(bold=True, size=12)),
      ('Nesta base não aparece uma feature "culpada" pela redução do Enterprise. O que pesa é o valor: cada redução Enterprise custa caro. A ação prioritária é proteger as contas Enterprise (acompanhamento próximo antes de renovações/reconfigurações), mais do que mexer em uma feature específica.', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      ('• Base fictícia. Datas de uso não são confiáveis (uso ligado à assinatura sem olhar data). Grupos pequenos por feature e por categoria.', {})]
for i, (tx, st_) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st_) if st_ else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT)
print(red.round(2).to_dict(), 'featsim', featsim, 'segsim', segsim)
print(S.groupby('plan_tier').mrr_reduzido.sum().to_dict())
