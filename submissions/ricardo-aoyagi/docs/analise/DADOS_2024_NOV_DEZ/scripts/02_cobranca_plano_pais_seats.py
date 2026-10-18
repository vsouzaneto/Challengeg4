# Nov/Dez 2024 · Hipóteses: (1) faixa de seats muda o cancelamento rápido; (2) anual cancela mais/menos conforme plano e país.
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada. Só quantidades.
import sys
from pathlib import Path
from math import lgamma, exp
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '02_Cobranca_Plano_Pais_Seats.xlsx'
INI, FIM = pd.Timestamp('2024-11-01'), pd.Timestamp('2024-12-31'); LIMITE = 30
FAIXAS = [0, 5, 10, 20, 40, 80, 10**4]; LF = ['01–05', '06–10', '11–20', '21–40', '41–80', '81 ou mais']

a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])

def fisher_p(a_, b_, c_, d_):
    n1, n2, m1 = a_ + b_, c_ + d_, a_ + c_; N = n1 + n2
    if N == 0 or n1 == 0 or n2 == 0 or m1 == 0 or m1 == N: return 1.0
    lc = lambda n, k: lgamma(n + 1) - lgamma(k + 1) - lgamma(n - k + 1)
    pmf = lambda x: exp(lc(n1, x) + lc(n2, m1 - x) - lc(N, m1))
    p0 = pmf(a_); lo, hi = max(0, m1 - n2), min(m1, n1)
    return min(1.0, sum(pmf(x) for x in range(lo, hi + 1) if pmf(x) <= p0 * (1 + 1e-7)))
def veredito(k1, n1, k2, n2):
    if k1 + k2 < 10: return 'Poucos casos para concluir'
    return 'Sim' if fisher_p(k1, n1 - k1, k2, n2 - k2) < 0.05 else 'Não (pode ser coincidência)'

GR, GL, GA = f'Rápida (< {LIMITE} dias)', f'Lenta ({LIMITE}+ dias)', 'Ativa em 31/12'
S = s.merge(a[['account_id', 'country', 'plan_tier']].rename(columns={'plan_tier': 'plano_inicial_conta'}), on='account_id')
S['dias_de_vida'] = (S.end_date - S.start_date).dt.days
enc = S.end_date.between(INI, FIM)
S['grupo'] = np.select([enc & (S.dias_de_vida < LIMITE), enc, S.end_date.isna()], [GR, GL, GA], 'Fora')
S = S[S.grupo != 'Fora'].copy()
S['rapida'] = (S.grupo == GR).astype(int); S['cancelou_novdez'] = (S.grupo != GA).astype(int)
S['faixa_seats'] = pd.cut(S.seats, FAIXAS, labels=LF).astype(str)
S['cobranca'] = S.billing_frequency.map({'annual': 'Anual', 'monthly': 'Mensal'})
S = S.rename(columns={'plan_tier': 'plano_assinatura', 'country': 'pais'})

# Fonte B (eventos nov/dez): atributos da conta
B = c[c.churn_date.between(INI, FIM)].merge(a[['account_id', 'country', 'plan_tier', 'seats', 'signup_date']], on='account_id')
B['grupo'] = np.where((B.churn_date - B.signup_date).dt.days < LIMITE, f'Conta nova (< {LIMITE} dias)', f'Conta antiga ({LIMITE}+ dias)')
BN, BV = f'Conta nova (< {LIMITE} dias)', f'Conta antiga ({LIMITE}+ dias)'
B['faixa_seats_conta'] = pd.cut(B.seats, FAIXAS, labels=LF).astype(str)
def cob_pred(r):
    x = s[(s.account_id == r.account_id) & (s.start_date <= r.churn_date) & (s.end_date.isna() | (s.end_date >= r.churn_date))]
    if not len(x): x = s[s.account_id == r.account_id]
    v = x.billing_frequency.value_counts()
    return 'Empate' if len(v) > 1 and v.iloc[0] == v.iloc[1] else {'annual': 'Anual', 'monthly': 'Mensal'}[v.index[0]]
B['cobranca_predominante'] = B.apply(cob_pred, axis=1)
B = B.rename(columns={'plan_tier': 'plano_inicial_conta', 'country': 'pais'})

# ===================== planilha =====================
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
HD = PatternFill('solid', fgColor='1F2937'); HAn = PatternFill('solid', fgColor='7C3AED'); HMe = PatternFill('solid', fgColor='0F766E'); HX = PatternFill('solid', fgColor='374151'); HRp = PatternFill('solid', fgColor='C2410C'); HLt = PatternFill('solid', fgColor='9A3412')
YEL = PatternFill('solid', fgColor='FFF3CD'); GRY = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); h1 = S_('H1_Seats'); h2 = S_('H2_Cobranca'); h3 = S_('H2_Cobranca_Plano_Pais'); fb = S_('Fonte_B_Eventos')
BS = {}
def py(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v, 'item') else v
def base(nome, df, cols, fmts):
    ws = wb.create_sheet(nome); ws.append(cols)
    for row in df[cols].itertuples(index=False): ws.append([py(v) for v in row])
    n = len(df) + 1; BS[nome] = (cols, n)
    for i, cn in enumerate(cols, 1):
        h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(24, len(cn) + 2))
        if cn in fmts:
            for r in range(2, n + 1): ws.cell(r, i).number_format = fmts[cn]
    tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
def Rg(b, col): cols, n = BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
D = 'dd/mm/yyyy'
base('Base_Assinaturas', S, ['subscription_id', 'account_id', 'grupo', 'rapida', 'cancelou_novdez', 'cobranca', 'plano_assinatura', 'plano_inicial_conta', 'pais', 'seats', 'faixa_seats', 'start_date', 'end_date', 'dias_de_vida'], {'start_date': D, 'end_date': D})
base('Base_Eventos', B, ['churn_event_id', 'account_id', 'grupo', 'churn_date', 'plano_inicial_conta', 'pais', 'seats', 'faixa_seats_conta', 'cobranca_predominante'], {'churn_date': D})
def H(ws, row, labels, fills, h=42):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fill=None):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x
def title(ws, t1, t2):
    ws['A1'] = t1; ws['A1'].font = Fn(bold=True, size=15); ws['A2'] = t2; ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True)
rp, cn_ = Rg('Base_Assinaturas', 'rapida'), Rg('Base_Assinaturas', 'cancelou_novdez')
SUB = lambda conds: ','.join(f'{Rg("Base_Assinaturas", k)},"{v}"' for k, v in conds)
NOTA = 'Base = assinaturas que estavam na carteira em nov/dez (encerradas em nov/dez + ativas em 31/12). "Rápidas a cada 100" = quantas dessas cancelaram com menos de 30 dias de vida.'

# ---------- H1_Seats ----------
ws = h1; title(ws, 'Hipótese 1 · Assinaturas com mais ou menos seats cancelam mais?', NOTA); ws.row_dimensions[2].height = 26
H(ws, 4, ['Faixa de seats', 'Assinaturas na base', 'Rápidas (< 30 dias)', 'Rápidas a cada 100', 'Canceladas em nov/dez (qualquer tempo)', 'Canceladas a cada 100', 'Rápidas: faixa diferente das demais?'], [HD, HD, HRp, HRp, HLt, HLt, HX])
r = 5
for lf in LF:
    k1 = int(S[(S.faixa_seats == lf)].rapida.sum()); n1 = int((S.faixa_seats == lf).sum()); k2 = int(S.rapida.sum()) - k1; n2 = len(S) - n1
    C(ws, r, 1, lf, bold=True); C(ws, r, 2, f'=COUNTIFS({SUB([("faixa_seats", lf)])})', '#,##0'); C(ws, r, 3, f'=SUMIFS({rp},{SUB([("faixa_seats", lf)])})', '#,##0')
    C(ws, r, 4, f'=C{r}/B{r}*100', '0.0'); C(ws, r, 5, f'=SUMIFS({cn_},{SUB([("faixa_seats", lf)])})', '#,##0'); C(ws, r, 6, f'=E{r}/B{r}*100', '0.0'); C(ws, r, 7, veredito(k1, n1, k2, n2)); r += 1
C(ws, r, 1, 'Todas', bold=True, fill=GRY); C(ws, r, 2, f'=SUM(B5:B{r-1})', '#,##0', True, GRY); C(ws, r, 3, f'=SUM(C5:C{r-1})', '#,##0', True, GRY); C(ws, r, 4, f'=C{r}/B{r}*100', '0.0', True, GRY)
C(ws, r, 5, f'=SUM(E5:E{r-1})', '#,##0', True, GRY); C(ws, r, 6, f'=E{r}/B{r}*100', '0.0', True, GRY); C(ws, r, 7, '', fill=GRY); rt1 = r
ch = BarChart(); ch.type = 'col'; ch.title = 'Rápidas a cada 100 assinaturas, por faixa de seats'; ch.height = 7; ch.width = 15; ch.legend = None
ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=rt1 - 1), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rt1 - 1)); ch.series[0].graphicalProperties.solidFill = 'C2410C'; ws.add_chart(ch, 'I4')
r = rt1 + 3; ws.cell(r, 1, 'Faixa de seats × plano da assinatura: rápidas a cada 100 (entre parênteses: rápidas / base)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Faixa de seats'] + ['Basic', 'Pro', 'Enterprise'], [HD] * 4); r += 1; m0 = r
for lf in LF:
    C(ws, r, 1, lf, bold=True)
    for j, pl in enumerate(['Basic', 'Pro', 'Enterprise'], 2):
        cond = SUB([('faixa_seats', lf), ('plano_assinatura', pl)])
        C(ws, r, j, f'=IFERROR(TEXT(SUMIFS({rp},{cond})/COUNTIFS({cond})*100,"0.0")&"  ("&SUMIFS({rp},{cond})&" / "&COUNTIFS({cond})&")","")')
    r += 1
ws.cell(r + 1, 1, 'Leitura: nenhuma faixa de seats se destaca; as taxas ficam entre ~3 e ~4,5 a cada 100. Nas combinações com plano, as caixas têm poucas rápidas (muitas com 1 a 7), então diferenças grandes (ex.: Pro 01–05) vêm de poucos casos.').font = Fn(italic=True, size=9, color=GRAY)
ws.column_dimensions['A'].width = 16
for c_ in 'BCDEFG': ws.column_dimensions[c_].width = 16
ws.column_dimensions['G'].width = 26

# ---------- H2_Cobranca ----------
ws = h2; title(ws, 'Hipótese 2 · Anual cancela mais ou menos que mensal? Depende do plano? Do país?', NOTA); ws.row_dimensions[2].height = 26
def bloco(ws, r, titulo, col, valores, rotulo):
    ws.cell(r, 1, titulo).font = Fn(bold=True, size=12); r += 1
    H(ws, r, [rotulo, 'Anual: base', 'Anual: rápidas', 'Anual: a cada 100', 'Mensal: base', 'Mensal: rápidas', 'Mensal: a cada 100', 'Anual − Mensal', 'Diferença confiável?', 'Anual: canceladas a cada 100 (qualquer tempo)', 'Mensal: canceladas a cada 100 (qualquer tempo)'],
      [HD, HAn, HAn, HAn, HMe, HMe, HMe, HX, HX, HAn, HMe]); r += 1; r0 = r
    for v in valores:
        cA = [('cobranca', 'Anual')] + ([(col, v)] if col else []); cM = [('cobranca', 'Mensal')] + ([(col, v)] if col else [])
        msk = (S[col] == v) if col else pd.Series(True, index=S.index)
        kA = int(S[msk & (S.cobranca == 'Anual')].rapida.sum()); nA = int((msk & (S.cobranca == 'Anual')).sum()); kM = int(S[msk & (S.cobranca == 'Mensal')].rapida.sum()); nM = int((msk & (S.cobranca == 'Mensal')).sum())
        C(ws, r, 1, v if col else 'Todas as assinaturas', bold=True)
        C(ws, r, 2, f'=COUNTIFS({SUB(cA)})', '#,##0'); C(ws, r, 3, f'=SUMIFS({rp},{SUB(cA)})', '#,##0'); C(ws, r, 4, f'=IF(B{r}>0,C{r}/B{r}*100,"")', '0.0')
        C(ws, r, 5, f'=COUNTIFS({SUB(cM)})', '#,##0'); C(ws, r, 6, f'=SUMIFS({rp},{SUB(cM)})', '#,##0'); C(ws, r, 7, f'=IF(E{r}>0,F{r}/E{r}*100,"")', '0.0')
        C(ws, r, 8, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(G{r})),D{r}-G{r},"")', '+0.0;-0.0;0.0'); C(ws, r, 9, veredito(kA, nA, kM, nM))
        C(ws, r, 10, f'=IF(B{r}>0,SUMIFS({cn_},{SUB(cA)})/B{r}*100,"")', '0.0'); C(ws, r, 11, f'=IF(E{r}>0,SUMIFS({cn_},{SUB(cM)})/E{r}*100,"")', '0.0'); r += 1
    ws.conditional_formatting.add(f'I{r0}:I{r-1}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
    return r + 1, r0
r = 4
r, rg = bloco(ws, r, '1. Geral', None, [None], 'Recorte')
r, rpl = bloco(ws, r, '2. Por plano da assinatura', 'plano_assinatura', ['Basic', 'Pro', 'Enterprise'], 'Plano da assinatura')
r, rpi = bloco(ws, r, '3. Por plano inicial contratado (cadastro da conta)', 'plano_inicial_conta', ['Basic', 'Pro', 'Enterprise'], 'Plano inicial')
ws.cell(r - 1, 1, 'Atenção: o "plano inicial" do cadastro bate com o plano da primeira assinatura em só 1 a cada 3 contas (ver DADOS_RECONCILIA). Use com cautela.').font = Fn(italic=True, size=9, color='B45309'); r += 1
r, rpa = bloco(ws, r, '4. Por país', 'pais', sorted(S.pais.unique()), 'País')
ws.cell(r, 1, 'Leitura: no geral, assinaturas anuais têm um pouco mais de cancelamento rápido que as mensais (cerca de 1 a cada 100 a mais), mas a diferença pode ser coincidência. Por plano e por país, a direção se repete na maioria das linhas (anual ≥ mensal), sem nenhuma diferença confiável isolada; vários países têm poucos casos.').font = Fn(italic=True, size=9, color=GRAY)
ws.cell(r + 1, 1, 'Ponto de atenção para o negócio: uma assinatura ANUAL cancelada com menos de 30 dias costuma significar arrependimento logo após a compra (e possível reembolso). Vale olhar essas 89 assinaturas à parte.').font = Fn(italic=True, size=9, color='B45309')
for col_, w in zip('ABCDEFGHIJK', [26, 10, 10, 10, 10, 10, 10, 10, 24, 14, 14]): ws.column_dimensions[col_].width = w
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.title = 'Rápidas a cada 100: anual vs mensal, por país'; ch.height = 8; ch.width = 18
ch.add_data(Reference(ws, min_col=4, min_row=rpa - 1, max_row=rpa + S.pais.nunique() - 1), titles_from_data=True); ch.add_data(Reference(ws, min_col=7, min_row=rpa - 1, max_row=rpa + S.pais.nunique() - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=rpa, max_row=rpa + S.pais.nunique() - 1)); ch.series[0].graphicalProperties.solidFill = '7C3AED'; ch.series[1].graphicalProperties.solidFill = '0F766E'; ch.legend.position = 'b'; ws.add_chart(ch, 'M4')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.title = 'Rápidas a cada 100: anual vs mensal, por plano da assinatura'; ch.height = 7; ch.width = 14
ch.add_data(Reference(ws, min_col=4, min_row=rpl - 1, max_row=rpl + 2), titles_from_data=True); ch.add_data(Reference(ws, min_col=7, min_row=rpl - 1, max_row=rpl + 2), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=rpl, max_row=rpl + 2)); ch.series[0].graphicalProperties.solidFill = '7C3AED'; ch.series[1].graphicalProperties.solidFill = '0F766E'; ch.legend.position = 'b'; ws.add_chart(ch, 'M22')

# ---------- H2_Cobranca_Plano_Pais (3 dimensões) ----------
ws = h3; title(ws, 'Cobrança × plano × país: rápidas a cada 100 (entre parênteses: rápidas / base)', NOTA + ' Combinação com 3 dimensões: a maioria das caixas tem 0 a 5 rápidas; trate como mapa para levantar perguntas, não como conclusão.'); ws.row_dimensions[2].height = 30
paises = sorted(S.pais.unique())
H(ws, 4, ['País'] + [f'{pl} · {cb}' for pl in ['Basic', 'Pro', 'Enterprise'] for cb in ['Anual', 'Mensal']], [HD] + [HAn, HMe] * 3, 30)
r = 5
for pa in paises:
    C(ws, r, 1, pa, bold=True); j = 2
    for pl in ['Basic', 'Pro', 'Enterprise']:
        for cb in ['Anual', 'Mensal']:
            cond = SUB([('pais', pa), ('plano_assinatura', pl), ('cobranca', cb)])
            C(ws, r, j, f'=IFERROR(SUMIFS({rp},{cond})/COUNTIFS({cond})*100,"")', '0.0'); j += 1
    r += 1
ws.conditional_formatting.add(f'B5:G{r-1}', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='max', end_color='F59E0B'))
r += 1; ws.cell(r, 1, 'Mesma tabela em quantidades (rápidas / base)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['País'] + [f'{pl} · {cb}' for pl in ['Basic', 'Pro', 'Enterprise'] for cb in ['Anual', 'Mensal']], [HD] + [HAn, HMe] * 3, 30); r += 1
for pa in paises:
    C(ws, r, 1, pa, bold=True); j = 2
    for pl in ['Basic', 'Pro', 'Enterprise']:
        for cb in ['Anual', 'Mensal']:
            cond = SUB([('pais', pa), ('plano_assinatura', pl), ('cobranca', cb)])
            C(ws, r, j, f'=SUMIFS({rp},{cond})&" / "&COUNTIFS({cond})'); j += 1
    r += 1
ws.column_dimensions['A'].width = 10
for j in range(2, 8): ws.column_dimensions[L(j)].width = 15

# ---------- Fonte B ----------
ws = fb; title(ws, 'Fonte B · Eventos de churn de nov/dez: conta nova vs conta antiga', 'Distribuição (% do grupo). Atributos da conta: plano inicial, país, seats do cadastro e cobrança predominante das assinaturas ativas na data do evento.')
gB = Rg('Base_Eventos', 'grupo')
H(ws, 4, ['Característica', 'Categoria', 'Conta nova: qtd', 'Conta nova: %', 'Conta antiga: qtd', 'Conta antiga: %', 'Nova − Antiga (pontos %)', 'Diferença confiável?'], [HD, HD, HRp, HRp, HLt, HLt, HX, HX])
r = 5; nBN, nBV = int((B.grupo == BN).sum()), int((B.grupo == BV).sum())
for col, lab, vals in [('cobranca_predominante', 'Cobrança predominante', sorted(B.cobranca_predominante.unique())), ('plano_inicial_conta', 'Plano inicial', ['Basic', 'Pro', 'Enterprise']),
                       ('pais', 'País', sorted(B.pais.unique())), ('faixa_seats_conta', 'Seats da conta (cadastro)', [x for x in LF if x in set(B.faixa_seats_conta)])]:
    first = r
    for v in vals:
        k1 = int(((B.grupo == BN) & (B[col] == v)).sum()); k2 = int(((B.grupo == BV) & (B[col] == v)).sum())
        C(ws, r, 1, lab if r == first else '', bold=True); C(ws, r, 2, v, bold=True)
        C(ws, r, 3, f'=COUNTIFS({gB},"{BN}",{Rg("Base_Eventos", col)},"{v}")', '#,##0'); C(ws, r, 4, f'=C{r}/COUNTIFS({gB},"{BN}")', '0%')
        C(ws, r, 5, f'=COUNTIFS({gB},"{BV}",{Rg("Base_Eventos", col)},"{v}")', '#,##0'); C(ws, r, 6, f'=E{r}/COUNTIFS({gB},"{BV}")', '0%')
        C(ws, r, 7, f'=(D{r}-F{r})*100', '+0;-0;0'); C(ws, r, 8, veredito(k1, nBN, k2, nBV)); r += 1
    r += 1
ws.conditional_formatting.add(f'H5:H{r}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
ws.cell(r, 1, 'Leitura: com 49 eventos de conta nova, quase todas as categorias têm poucos casos. Nenhuma diferença confiável em cobrança, plano, país ou seats.').font = Fn(italic=True, size=9, color=GRAY)
for col_, w in zip('ABCDEFGH', [24, 16, 11, 11, 11, 11, 12, 26]): ws.column_dimensions[col_].width = w

# ---------- LEIA-ME ----------
kA = int(S[S.cobranca == 'Anual'].rapida.sum()); nA = int((S.cobranca == 'Anual').sum()); kM = int(S[S.cobranca == 'Mensal'].rapida.sum()); nM = int((S.cobranca == 'Mensal').sum())
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
Lm = [('Nov/dez 2024 · Cobrança, plano inicial, país e seats: alguma combinação cancela mais rápido?', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. Só quantidades. Mesmo recorte da planilha 01 (assinaturas encerradas em nov/dez + ativas em 31/12).', dict(italic=True, color=GRAY)),
      ('', {}), ('HIPÓTESES', dict(bold=True, size=12)),
      ('1. Assinaturas com mais ou menos seats cancelam mais.  → aba H1_Seats (faixas: 1–5, 6–10, 11–20, 21–40, 41–80, 81+)', {}),
      ('2. Planos anuais cancelam mais ou menos dependendo do plano contratado e do país.  → abas H2_Cobranca e H2_Cobranca_Plano_Pais', {}),
      ('Fonte B (eventos de churn) lado a lado  → aba Fonte_B_Eventos', {}),
      ('', {}), ('COMO LER', dict(bold=True, size=12)),
      ('• "Rápidas a cada 100": de cada 100 assinaturas daquele tipo na carteira de nov/dez, quantas cancelaram com menos de 30 dias de vida. Na média geral são cerca de 3 a cada 100.', {}),
      ('• "Diferença confiável?": "Sim" = diferença grande o bastante para não ser coincidência; "Não" = pode ser acaso; "Poucos casos" = menos de 10 cancelamentos rápidos na comparação.', {}),
      ('', {}), ('RESPOSTAS', dict(bold=True, size=12)),
      ('1. Seats: NÃO se confirma. Todas as faixas ficam entre ~3 e ~4,5 rápidas a cada 100, sem faixa diferente das demais. Olhando cancelamento em qualquer tempo, a faixa 81+ aparece um pouco acima (~9,5 a cada 100), mas com só 200 assinaturas.', {}),
      (f'2. Cobrança: anual tem um pouco mais de cancelamento rápido que mensal ({kA} de {nA} = {kA/nA*100:.1f} a cada 100, contra {kM} de {nM} = {kM/nM*100:.1f}). A direção se repete em quase todos os planos e na maioria dos países, mas nenhuma diferença é confiável sozinha. Ou seja: é um indício a acompanhar, não uma conclusão.', {}),
      ('   Por plano: a maior distância anual × mensal está em Basic e Pro; em Enterprise praticamente não há diferença. Por país: UK e US concentram a maior parte dos casos; os demais países têm poucos cancelamentos rápidos para concluir.', {}),
      ('3. Fonte B (eventos): conta nova e conta antiga se distribuem de forma parecida em cobrança, plano, país e seats.', {}),
      ('', {}), ('PONTO DE ATENÇÃO', dict(bold=True, size=12)),
      (f'• Assinatura ANUAL cancelada com menos de 30 dias ({kA} casos em nov/dez) costuma indicar arrependimento logo após a compra e possível reembolso. Mesmo sem diferença confiável, é um grupo que merece ser olhado à parte.', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      ('• Base fictícia. O plano inicial do cadastro não bate com as assinaturas na maioria das contas. Nas combinações de 3 dimensões, cada caixa tem poucos casos.', {})]
for i, (tx, st) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st) if st else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT)
print('anual', kA, nA, 'mensal', kM, nM, veredito(kA, nA, kM, nM), round(fisher_p(kA, nA - kA, kM, nM - kM), 3))
for pl in ['Basic', 'Pro', 'Enterprise']:
    m = S.plano_assinatura == pl; x = [int(S[m & (S.cobranca == cb)].rapida.sum()) for cb in ['Anual', 'Mensal']]; n = [int((m & (S.cobranca == cb)).sum()) for cb in ['Anual', 'Mensal']]
    print(pl, x, n, veredito(x[0], n[0], x[1], n[1]))
print(B.cobranca_predominante.value_counts().to_dict())
