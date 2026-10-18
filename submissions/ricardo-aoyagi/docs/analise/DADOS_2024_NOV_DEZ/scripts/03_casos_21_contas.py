# Nov/Dez 2024 · Caso a caso: contas em que as DUAS fontes concordam que o churn foi rápido
# (assinatura encerrada com < 30 dias em nov/dez  E  evento de churn de conta com < 30 dias de cadastro).
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada. Só quantidades.
import sys
from pathlib import Path
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '03_Casos_21_Contas.xlsx'
INI, FIM = pd.Timestamp('2024-11-01'), pd.Timestamp('2024-12-31'); LIMITE = 30; JAN = 7; SEED = 1

a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv', parse_dates=['usage_date'])
t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv', parse_dates=['submitted_at','closed_at'])
c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])
SN = lambda b: 'Sim' if b else 'Não'
s['dias_de_vida'] = (s.end_date - s.start_date).dt.days
enc = s.end_date.between(INI, FIM)
s['grupo'] = np.select([enc & (s.dias_de_vida < LIMITE), enc], ['Rápida', 'Lenta'], '')
ev = c[c.churn_date.between(INI, FIM)].merge(a[['account_id', 'signup_date']], on='account_id'); ev['dias_de_conta'] = (ev.churn_date - ev.signup_date).dt.days
CONTAS = sorted(set(s[s.grupo == 'Rápida'].account_id) & set(ev[ev.dias_de_conta < LIMITE].account_id))

# ---------- para cada assinatura encerrada em nov/dez: a conta continuou? houve compra nova perto? ----------
def contexto(r):
    outras = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
    perto = outras[(outras.start_date - r.end_date).abs().dt.days <= JAN]
    ativas = outras[(outras.start_date <= r.end_date) & (outras.end_date.isna() | (outras.end_date > r.end_date))]
    return pd.Series({'outras_ativas_no_dia': len(ativas), 'conta_ficou_sem_assinatura': SN(len(ativas) == 0),
                      'nova_assinatura_ate_7d': SN(len(perto) > 0), 'nova_com_mesmos_seats': SN((perto.seats == r.seats).any())})
X = s[s.grupo != ''].copy(); X = pd.concat([X, X.apply(contexto, axis=1)], axis=1)
X['caso_21'] = np.where(X.account_id.isin(CONTAS), 'Sim', 'Não')
# referência: chance de uma conta qualquer abrir assinatura em ±7 dias de uma data qualquer de nov/dez
rng = np.random.default_rng(SEED); acc_all = s.account_id.unique(); hits = []
for _ in range(3000):
    ac = rng.choice(acc_all); d = INI + pd.Timedelta(days=int(rng.integers(0, 61)))
    hits.append(((s.account_id == ac) & ((s.start_date - d).abs().dt.days <= JAN)).any())
REF7 = float(np.mean(hits))

# ---------- resumo por conta ----------
rows = []
for ac in CONTAS:
    A = a[a.account_id == ac].iloc[0]; Sx = s[s.account_id == ac]; Ex = ev[ev.account_id == ac].sort_values('churn_date'); Tx = t[t.account_id == ac]
    Fx = f[f.subscription_id.isin(Sx.subscription_id)]; Rx = X[(X.account_id == ac) & (X.grupo == 'Rápida')]
    e1 = Ex[Ex.dias_de_conta < LIMITE].iloc[0]
    rows.append(dict(conta=ac, industria=A.industry, pais=A.country, canal=A.referral_source, plano_inicial=A.plan_tier, seats_cadastro=int(A.seats), churn_flag=SN(A.churn_flag),
        cadastro=A.signup_date, dias_ate_1a_assinatura=int((Sx.start_date.min() - A.signup_date).days), assinaturas_total=len(Sx),
        assinaturas_abertas_novdez=int(Sx.start_date.between(INI, FIM).sum()), encerradas_rapidas=len(Rx), ativas_em_31_12=int(Sx.end_date.isna().sum()),
        data_evento=e1.churn_date, dias_de_conta_no_evento=int(e1.dias_de_conta), motivo=e1.reason_code, comentario=e1.feedback_text if isinstance(e1.feedback_text, str) else '(sem comentário)',
        reembolso=SN(e1.refund_amount_usd > 0), nova_assinatura_ate_7d_do_cancelamento=SN((Rx.nova_assinatura_ate_7d == 'Sim').any()),
        tickets_total=len(Tx), tickets_datados_antes_do_cadastro=int((Tx.submitted_at < A.signup_date).sum()),
        tickets_entre_cadastro_e_evento=int(((Tx.submitted_at >= A.signup_date) & (Tx.submitted_at <= e1.churn_date)).sum()),
        registros_uso=len(Fx), pct_registros_com_erro=round(float((Fx.error_count > 0).mean()), 3) if len(Fx) else None))
RS = pd.DataFrame(rows)

# ---------- linha do tempo ----------
lt = []
for ac in CONTAS:
    A = a[a.account_id == ac].iloc[0]
    lt.append(dict(conta=ac, data=A.signup_date, tipo='1 Cadastro da conta', detalhe=f'{A.industry} · {A.country} · canal {A.referral_source} · plano inicial {A.plan_tier} · {A.seats} seats', alerta=''))
    for r in s[s.account_id == ac].itertuples():
        d = f'{r.subscription_id} · {r.plan_tier} · {"anual" if r.billing_frequency=="annual" else "mensal"} · {r.seats} seats' + (' · trial' if r.is_trial else '')
        lt.append(dict(conta=ac, data=r.start_date, tipo='2 Assinatura iniciada', detalhe=d, alerta=''))
        if pd.notna(r.end_date):
            lt.append(dict(conta=ac, data=r.end_date, tipo='3 Assinatura encerrada', detalhe=f'{d} · {int(r.dias_de_vida)} dias de vida', alerta='Cancelamento rápido' if r.dias_de_vida < LIMITE and INI <= r.end_date <= FIM else ''))
    for r in c[c.account_id == ac].itertuples():
        lt.append(dict(conta=ac, data=r.churn_date, tipo='4 Evento de churn', detalhe=f'motivo {r.reason_code} · comentário: {r.feedback_text if isinstance(r.feedback_text, str) else "(sem comentário)"}' + (f' · reembolso US$ {r.refund_amount_usd:.2f}' if r.refund_amount_usd > 0 else ''), alerta=''))
    for r in t[t.account_id == ac].itertuples():
        sat = '' if pd.isna(r.satisfaction_score) else f' · satisfação {int(r.satisfaction_score)}'
        lt.append(dict(conta=ac, data=r.submitted_at, tipo='5 Ticket', detalhe=f'{r.ticket_id} · prioridade {r.priority}{sat}' + (' · escalado' if r.escalation_flag else ''), alerta='Data ANTES do cadastro' if r.submitted_at < A.signup_date else ''))
LT = pd.DataFrame(lt).sort_values(['conta', 'data', 'tipo']).reset_index(drop=True)

# ---------- planilha ----------
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
HD = PatternFill('solid', fgColor='1F2937'); HRp = PatternFill('solid', fgColor='C2410C'); HLt = PatternFill('solid', fgColor='9A3412'); HX = PatternFill('solid', fgColor='374151'); HB = PatternFill('solid', fgColor='1D4ED8')
YEL = PatternFill('solid', fgColor='FFF3CD'); RED = PatternFill('solid', fgColor='F8D7DA'); GRY = PatternFill('solid', fgColor='EDEDED'); BLU = PatternFill('solid', fgColor='DBEAFE')
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); tc = S_('Tracos_comuns')
def py(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v, 'item') else v
def tabela(nome, df, nomes, fmts, widths, tname):
    ws = wb.create_sheet(nome); ws.append(nomes)
    for row in df.itertuples(index=False): ws.append([py(v) for v in row])
    n = len(df) + 1
    for i, cn in enumerate(nomes, 1):
        h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; h.alignment = Alignment(wrap_text=True, vertical='center'); ws.column_dimensions[L(i)].width = widths[i - 1]
        for r in range(2, n + 1):
            x = ws.cell(r, i); x.font = Fn(size=9)
            if df.columns[i - 1] in fmts: x.number_format = fmts[df.columns[i - 1]]
    ws.row_dimensions[1].height = 44
    tb = Table(displayName=tname, ref=f'A1:{L(len(nomes))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
    return ws, n
D = 'dd/mm/yyyy'
nomesR = ['Conta', 'Indústria', 'País', 'Canal', 'Plano inicial', 'Seats (cadastro)', 'churn_flag', 'Cadastro', 'Dias até 1ª assinatura', 'Assinaturas (total)', 'Abertas em nov/dez', 'Encerradas rápido', 'Ativas em 31/12',
          'Data do evento', 'Dias de conta no evento', 'Motivo', 'Comentário', 'Reembolso?', 'Nova assinatura até 7 dias do cancelamento?', 'Tickets (total)', 'Tickets com data ANTES do cadastro', 'Tickets entre cadastro e evento', 'Registros de uso', '% registros com erro']
wsR, nR = tabela('Resumo_21', RS, nomesR, {'cadastro': D, 'data_evento': D, 'pct_registros_com_erro': '0%'}, [10, 12, 6, 9, 10, 8, 8, 10, 9, 9, 9, 9, 9, 10, 9, 10, 18, 9, 12, 8, 10, 10, 9, 9], 'tResumo21')
wsR.conditional_formatting.add(f'U2:U{nR}', CellIsRule(operator='greaterThan', formula=['0'], fill=RED))
wsR.conditional_formatting.add(f'M2:M{nR}', CellIsRule(operator='greaterThan', formula=['0'], fill=BLU))
wsL, nL = tabela('Linha_do_tempo', LT, ['Conta', 'Data', 'Tipo', 'Detalhe', 'Alerta'], {'data': D}, [10, 11, 20, 80, 22], 'tLinhaTempo')
wsL.conditional_formatting.add(f'E2:E{nL}', CellIsRule(operator='equal', formula=['"Data ANTES do cadastro"'], fill=RED))
wsL.conditional_formatting.add(f'E2:E{nL}', CellIsRule(operator='equal', formula=['"Cancelamento rápido"'], fill=YEL))
wsL.conditional_formatting.add(f'A2:D{nL}', FormulaRule(formula=[f'$C2="4 Evento de churn"'], fill=PatternFill('solid', fgColor='FDE68A')))
colsX = ['subscription_id', 'account_id', 'grupo', 'caso_21', 'start_date', 'end_date', 'dias_de_vida', 'plan_tier', 'billing_frequency', 'seats', 'outras_ativas_no_dia', 'conta_ficou_sem_assinatura', 'nova_assinatura_ate_7d', 'nova_com_mesmos_seats']
wsX, nX = tabela('Base_Encerradas_NovDez', X[colsX], ['Assinatura', 'Conta', 'Grupo', 'Uma das 21 contas?', 'Início', 'Fim', 'Dias de vida', 'Plano', 'Cobrança', 'Seats', 'Outras assinaturas ativas no dia do fim',
                  'Conta ficou sem nenhuma assinatura?', 'Nova assinatura até 7 dias (antes/depois)?', 'Nova com mesmos seats?'], {'start_date': D, 'end_date': D}, [11, 10, 8, 9, 10, 10, 8, 10, 9, 7, 11, 11, 11, 11], 'tEncerradas')
def RX(col): i = colsX.index(col) + 1; return f'Base_Encerradas_NovDez!${L(i)}$2:${L(i)}${nX}'
def RR(i): return f'Resumo_21!${L(i)}$2:${L(i)}${nR}'

# ---------- Traços comuns ----------
ws = tc
ws['A1'] = 'O que as 21 contas têm em comum (e o que isso diz sobre os cancelamentos rápidos)'; ws['A1'].font = Fn(bold=True, size=15)
ws['A2'] = 'Contas em que as duas fontes concordam: assinatura encerrada com menos de 30 dias em nov/dez E evento de churn de conta com menos de 30 dias de cadastro. Contagens por fórmula sobre Resumo_21 e Base_Encerradas_NovDez.'
ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:E2'); ws.row_dimensions[2].height = 28
def H(row, labels, fills):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = 36
def C(r, c_, v, fmt=None, bold=False):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical='top')
    if fmt: x.number_format = fmt
    return x
r = 4; ws.cell(r, 1, '1. A conta foi embora?').font = Fn(bold=True, size=12); r += 1
H(r, ['Pergunta', 'Resultado', 'De', '%', 'Leitura'], [HD, HRp, HD, HRp, HD]); r += 1
itens = [('Das 21 contas, quantas continuam com assinatura ativa em 31/12?', f'=COUNTIF({RR(13)},">0")', f'=COUNTA({RR(1)})', 'Todas seguem clientes. O "churn" foi de uma assinatura, não da conta.'),
         ('Ativas em 31/12 por conta (média)', f'=AVERAGE({RR(13)})', None, ''),
         ('Assinaturas abertas em nov/dez por conta (média)', f'=AVERAGE({RR(11)})', None, 'São contas novas que abriram várias assinaturas em poucos dias.'),
         ('Dias entre o cadastro e a 1ª assinatura (média)', f'=AVERAGE({RR(9)})', None, ''),
         ('Das 159 assinaturas rápidas de nov/dez, em quantas a conta ficou SEM nenhuma assinatura?', f'=COUNTIFS({RX("grupo")},"Rápida",{RX("conta_ficou_sem_assinatura")},"Sim")', f'=COUNTIFS({RX("grupo")},"Rápida")', 'Praticamente nenhuma: o cancelamento rápido quase nunca é a saída do cliente.'),
         ('Outras assinaturas ativas no dia do cancelamento rápido (média, 159 rápidas)', f'=AVERAGEIFS({RX("outras_ativas_no_dia")},{RX("grupo")},"Rápida")', None, '')]
for lab, v, de, txt in itens:
    C(r, 1, lab, bold=True); C(r, 2, v, '0.0' if de is None else '#,##0'); C(r, 3, de, '#,##0') if de else C(r, 3, ''); C(r, 4, f'=B{r}/C{r}' if de else '', '0%'); C(r, 5, txt); r += 1
r += 1; ws.cell(r, 1, '2. O cancelamento rápido parece troca/ajuste de assinatura?').font = Fn(bold=True, size=12); r += 1
H(r, ['Grupo', 'Assinaturas', 'Com nova assinatura da mesma conta até 7 dias antes/depois', '%', 'Leitura'], [HD, HD, HRp, HRp, HD]); r += 1; r2 = r
for lab, cond in [('As 21 contas (assinaturas rápidas)', f'{RX("grupo")},"Rápida",{RX("caso_21")},"Sim"'), ('Todas as rápidas de nov/dez (159)', f'{RX("grupo")},"Rápida"'), ('Lentas de nov/dez (107)', f'{RX("grupo")},"Lenta"')]:
    C(r, 1, lab, bold=True); C(r, 2, f'=COUNTIFS({cond})', '#,##0'); C(r, 3, f'=COUNTIFS({cond},{RX("nova_assinatura_ate_7d")},"Sim")', '#,##0'); C(r, 4, f'=C{r}/B{r}', '0%'); C(r, 5, ''); r += 1
C(r, 1, 'Referência: uma conta qualquer, numa data qualquer de nov/dez', bold=True); C(r, 2, ''); C(r, 3, ''); C(r, 4, round(REF7, 3), '0%')
C(r, 5, 'Sorteio de 3.000 pares conta × data (calculado no script). Mostra o quanto é comum abrir assinatura nova em qualquer janela de 15 dias.'); r += 1
ws.cell(r, 1, 'Leitura: quem cancela rápido abre outra assinatura perto da data com mais frequência do que o normal (≈ 2 em 3, contra ≈ 2 em 5). Combina com troca de plano/configuração logo após a compra, e não com abandono.').font = Fn(italic=True, size=9, color=GRAY); r += 2
ws.cell(r, 1, '3. Houve reclamação antes?').font = Fn(bold=True, size=12); r += 1
H(r, ['Pergunta', 'Resultado', 'De', '%', 'Leitura'], [HD, HRp, HD, HRp, HD]); r += 1
for lab, v, de, txt in [('Tickets das 21 contas (total)', f'=SUM({RR(20)})', None, ''),
                        ('… com data ANTES de a conta existir', f'=SUM({RR(21)})', f'=SUM({RR(20)})', 'Impossível na vida real: a ligação ticket → conta não é confiável para essas contas.'),
                        ('… entre o cadastro e o evento de churn', f'=SUM({RR(22)})', f'=SUM({RR(20)})', 'Quase nenhuma reclamação no período em que a conta existia.'),
                        ('Contas com ao menos 1 ticket entre cadastro e evento', f'=COUNTIF({RR(22)},">0")', f'=COUNTA({RR(1)})', '')]:
    C(r, 1, lab, bold=True); C(r, 2, v, '#,##0'); C(r, 3, de, '#,##0') if de else C(r, 3, ''); C(r, 4, f'=B{r}/C{r}' if de else '', '0%'); C(r, 5, txt); r += 1
r += 1; ws.cell(r, 1, '4. O que o cliente declarou no evento').font = Fn(bold=True, size=12); r += 1
H(r, ['Motivo / comentário', 'Contas', '', '', ''], [HD, HRp, HD, HD, HD]); r += 1
for v in sorted(RS.motivo.unique()): C(r, 1, f'Motivo: {v}', bold=True); C(r, 2, f'=COUNTIF({RR(16)},"{v}")', '#,##0'); r += 1
for v in sorted(RS.comentario.unique()): C(r, 1, f'Comentário: {v}', bold=True); C(r, 2, f'=COUNTIF({RR(17)},"{v}")', '#,##0'); r += 1
C(r, 1, 'Com reembolso', bold=True); C(r, 2, f'=COUNTIF({RR(18)},"Sim")', '#,##0'); r += 1
ws.cell(r, 1, 'Motivos espalhados, sem um dominante; e o motivo segue sem conversar com o comentário (ver Linha_do_tempo).').font = Fn(italic=True, size=9, color=GRAY); r += 2
ws.cell(r, 1, '5. Perfil das 21 contas').font = Fn(bold=True, size=12); r += 1
H(r, ['Característica', 'Valor mais comum', 'Contas', '', ''], [HD, HD, HRp, HD, HD]); r += 1
for lab, i_, col in [('Indústria', 2, 'industria'), ('País', 3, 'pais'), ('Canal', 4, 'canal'), ('Plano inicial', 5, 'plano_inicial')]:
    top = RS[col].value_counts().index[0]; C(r, 1, lab, bold=True); C(r, 2, top); C(r, 3, f'=COUNTIF({RR(i_)},B{r})', '#,##0'); r += 1
C(r, 1, 'churn_flag = True', bold=True); C(r, 2, 'True'); C(r, 3, f'=COUNTIF({RR(7)},"Sim")', '#,##0'); r += 1
ws.cell(r, 1, 'Perfil espalhado (exceto país: maioria US, que também é a maioria da base). A marcação churn_flag da conta quase nunca está ligada, mesmo nesses casos "limpos".').font = Fn(italic=True, size=9, color=GRAY)
for col_, w in zip('ABCDE', [62, 14, 14, 10, 70]): ws.column_dimensions[col_].width = w

# ---------- LEIA-ME ----------
n_ok = int((X[X.grupo == 'Rápida'].conta_ficou_sem_assinatura == 'Sim').sum())
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
Lm = [('Nov/dez 2024 · Caso a caso: as 21 contas em que as duas fontes concordam que o churn foi rápido', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. Só quantidades.', dict(italic=True, color=GRAY)),
      ('', {}), ('QUEM SÃO', dict(bold=True, size=12)),
      (f'{len(CONTAS)} contas que, ao mesmo tempo, tiveram uma assinatura encerrada com menos de 30 dias em nov/dez (fonte assinaturas) e um evento de churn com menos de 30 dias de cadastro (fonte eventos). São os casos mais "limpos" de churn rápido.', {}),
      ('', {}), ('ABAS', dict(bold=True, size=12)),
      ('• Tracos_comuns: o que as 21 têm em comum, comparado com todas as 159 assinaturas rápidas e com as lentas.', {}),
      ('• Resumo_21: uma linha por conta.', {}),
      ('• Linha_do_tempo: tudo o que aconteceu com cada conta em ordem de data (cadastro, assinaturas, cancelamentos, eventos, tickets). Filtre pela coluna Conta para ler um caso.', {}),
      ('• Base_Encerradas_NovDez: as 266 assinaturas encerradas em nov/dez, com o que havia na conta no dia.', {}),
      ('', {}), ('O QUE OS CASOS MOSTRAM', dict(bold=True, size=12)),
      ('1. Nenhuma das 21 contas foi embora. Todas seguem com várias assinaturas ativas em 31/12 (cerca de 8 por conta). Nas 159 assinaturas rápidas de nov/dez, em só ' + str(n_ok) + ' caso a conta ficou sem nenhuma assinatura.', {}),
      ('2. São contas novas que compram em rajada: abrem várias assinaturas em poucos dias (às vezes no mesmo dia, com o mesmo número de seats) e cancelam uma delas logo em seguida.', {}),
      (f'3. Perto do cancelamento rápido costuma haver uma assinatura nova da mesma conta (≈ 2 em 3 casos, contra ≈ {REF7*100:.0f}% numa data qualquer). Parece troca ou ajuste de plano logo após a compra, e não abandono.', {}),
      ('4. Não há reclamação antes: quase todos os tickets dessas contas têm data ANTERIOR à existência da conta (impossível), e quase nenhum cai entre o cadastro e o evento.', {}),
      ('5. O evento de churn traz motivos variados e comentários que não batem com o motivo; churn_flag quase nunca está ligado.', {}),
      ('', {}), ('IMPLICAÇÃO', dict(bold=True, size=12)),
      ('O "cancelamento rápido" de nov/dez, que puxou o churn para cima, é em grande parte cancelamento de ASSINATURA dentro de contas que continuam ativas, com cara de ajuste no momento da contratação. Antes de tratar como perda de cliente, vale separar: (a) troca/ajuste de plano, (b) redução real de contrato e (c) saída do cliente. Com a regra atual, as três coisas são contadas como churn.', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      ('• Base fictícia; 21 contas é pouco para generalizar, por isso comparamos com as 159 rápidas. "Até 7 dias" é uma janela escolhida para a análise.', {})]
for i, (tx, st) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st) if st else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT, len(CONTAS), 'ref7', round(REF7, 3), 'sem assinatura', n_ok)
print(X.groupby('grupo').nova_assinatura_ate_7d.apply(lambda x: (x == 'Sim').mean()).round(3).to_dict(), X[(X.grupo == 'Rápida') & (X.caso_21 == 'Sim')].nova_assinatura_ate_7d.eq('Sim').mean().round(3))
print(RS[['tickets_total', 'tickets_datados_antes_do_cadastro', 'tickets_entre_cadastro_e_evento']].sum().to_dict(), RS.churn_flag.value_counts().to_dict(), RS.ativas_em_31_12.min())
