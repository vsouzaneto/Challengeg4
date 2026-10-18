# Upgrade/downgrade (flags) vs mudança real de plano/seats  +  compras em rajada por país/segmento.
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada. Só quantidades.
import sys
from pathlib import Path
from math import lgamma, exp
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '05_Upgrade_Downgrade_e_Rajadas.xlsx'
INI, FIM = pd.Timestamp('2024-11-01'), pd.Timestamp('2024-12-31'); LIMITE = 30; JP = 7
RAJ_N, RAJ_D = 3, 7          # rajada = 3 ou mais assinaturas da mesma conta iniciadas dentro de 7 dias
NIVEL = {'Basic': 1, 'Pro': 2, 'Enterprise': 3}

a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
def fisher_p(a_, b_, c_, d_):
    n1, n2, m1 = a_ + b_, c_ + d_, a_ + c_; N = n1 + n2
    if N == 0 or n1 == 0 or n2 == 0 or m1 == 0 or m1 == N: return 1.0
    lc = lambda n, k: lgamma(n + 1) - lgamma(k + 1) - lgamma(n - k + 1)
    pmf = lambda x: exp(lc(n1, x) + lc(n2, m1 - x) - lc(N, m1))
    p0 = pmf(a_); lo, hi = max(0, m1 - n2), min(m1, n1)
    return min(1.0, sum(pmf(x) for x in range(lo, hi + 1) if pmf(x) <= p0 * (1 + 1e-7)))
def conf(k1, n1, k2, n2):
    if k1 + k2 < 10: return 'Poucos casos para concluir'
    return 'Sim' if fisher_p(k1, n1 - k1, k2, n2 - k2) < 0.05 else 'Não (pode ser coincidência)'

S = s.merge(a[['account_id', 'industry', 'country', 'referral_source', 'signup_date']], on='account_id').sort_values(['account_id', 'start_date', 'subscription_id']).reset_index(drop=True)
S['marcacao'] = np.select([S.upgrade_flag & S.downgrade_flag, S.upgrade_flag, S.downgrade_flag], ['Upgrade e downgrade', 'Upgrade', 'Downgrade'], 'Nenhuma')
S['nivel'] = S.plan_tier.map(NIVEL)
g = S.groupby('account_id')
S['assinatura_anterior'] = g.subscription_id.shift(); S['plano_anterior'] = g.plan_tier.shift(); S['seats_anterior'] = g.seats.shift()
nv_ant = S.plano_anterior.map(NIVEL)
S['plano_vs_anterior'] = np.select([S.plano_anterior.isna(), S.nivel > nv_ant, S.nivel < nv_ant], ['Primeira da conta', 'Subiu', 'Desceu'], 'Igual')
S['seats_vs_anterior'] = np.select([S.plano_anterior.isna(), S.seats > S.seats_anterior, S.seats < S.seats_anterior], ['Primeira da conta', 'Subiu', 'Desceu'], 'Igual')
S['subiu_plano_ou_seats'] = np.where(S.plano_anterior.isna(), '', np.where((S.plano_vs_anterior == 'Subiu') | (S.seats_vs_anterior == 'Subiu'), 'Sim', 'Não'))
S['desceu_plano_ou_seats'] = np.where(S.plano_anterior.isna(), '', np.where((S.plano_vs_anterior == 'Desceu') | (S.seats_vs_anterior == 'Desceu'), 'Sim', 'Não'))
# rajada
def rajada(gr):
    d = gr.start_date.values; f = np.zeros(len(gr), bool)
    for i in range(len(gr)):
        w = (d >= d[i]) & (d <= d[i] + np.timedelta64(RAJ_D - 1, 'D'))
        if w.sum() >= RAJ_N: f |= w
    return pd.Series(f, index=gr.index)
S['em_rajada'] = np.where(S.groupby('account_id', group_keys=False).apply(rajada), 'Sim', 'Não')
S['periodo_inicio'] = np.where(S.start_date >= INI, 'Nov–Dez/24', np.where(S.start_date >= pd.Timestamp('2024-07-01'), 'Jul–Out/24', 'Até Jun/24'))
S['trimestre_inicio'] = S.start_date.dt.to_period('Q').astype(str)
S['dias_de_conta_no_inicio'] = (S.start_date - S.signup_date).dt.days
S['conta_nova_no_inicio'] = np.where(S.dias_de_conta_no_inicio < LIMITE, 'Sim', 'Não')
# encerradas em nov/dez: tipo (mesma regra da planilha 04) e par da troca
S['dias_de_vida'] = (S.end_date - S.start_date).dt.days
def tipo(r):
    o = S[(S.account_id == r.account_id) & (S.subscription_id != r.subscription_id)]
    ativas = o[(o.start_date <= r.end_date) & (o.end_date.isna() | (o.end_date > r.end_date))]
    perto = o[(o.start_date - r.end_date).abs().dt.days <= JP]
    depois = o[(o.start_date > r.end_date) & (o.start_date <= r.end_date + pd.Timedelta(days=JP))]
    if len(ativas) == 0 and len(depois) == 0: return pd.Series({'tipo_cancelamento': 'Saída do cliente', 'nova_mais_proxima': ''})
    if len(perto): return pd.Series({'tipo_cancelamento': 'Troca/ajuste', 'nova_mais_proxima': perto.iloc[(perto.start_date - r.end_date).abs().argsort()].iloc[0].subscription_id})
    return pd.Series({'tipo_cancelamento': 'Redução', 'nova_mais_proxima': ''})
E = S[S.end_date.between(INI, FIM)].copy(); E = pd.concat([E, E.apply(tipo, axis=1)], axis=1)
E['velocidade'] = np.where(E.dias_de_vida < LIMITE, f'Rápida (< {LIMITE} dias)', f'Lenta ({LIMITE}+ dias)')
nova = S.set_index('subscription_id')
def par(r):
    if not r.nova_mais_proxima: return pd.Series({'marcacao_da_nova': '', 'plano_nova_vs_cancelada': '', 'seats_nova_vs_cancelada': ''})
    n_ = nova.loc[r.nova_mais_proxima]; dp = NIVEL[n_.plan_tier] - NIVEL[r.plan_tier]
    return pd.Series({'marcacao_da_nova': n_.marcacao, 'plano_nova_vs_cancelada': 'Subiu' if dp > 0 else ('Desceu' if dp < 0 else 'Igual'),
                      'seats_nova_vs_cancelada': 'Subiu' if n_.seats > r.seats else ('Desceu' if n_.seats < r.seats else 'Igual')})
E = pd.concat([E, E.apply(par, axis=1)], axis=1)
E['par_subiu'] = np.where(E.nova_mais_proxima == '', '', np.where((E.plano_nova_vs_cancelada == 'Subiu') | (E.seats_nova_vs_cancelada == 'Subiu'), 'Sim', 'Não'))
E['par_desceu'] = np.where(E.nova_mais_proxima == '', '', np.where((E.plano_nova_vs_cancelada == 'Desceu') | (E.seats_nova_vs_cancelada == 'Desceu'), 'Sim', 'Não'))
# contas
K = a[['account_id', 'industry', 'country', 'referral_source', 'signup_date']].copy()
K['assinaturas'] = K.account_id.map(S.groupby('account_id').size())
K['assinaturas_em_rajada'] = K.account_id.map(S[S.em_rajada == 'Sim'].groupby('account_id').size()).fillna(0).astype(int)
K['teve_rajada'] = np.where(K.assinaturas_em_rajada > 0, 'Sim', 'Não')
K['teve_rajada_novdez'] = np.where(K.account_id.isin(S[(S.em_rajada == 'Sim') & (S.periodo_inicio == 'Nov–Dez/24')].account_id), 'Sim', 'Não')
K['cadastro_2024_S2'] = np.where(K.signup_date >= pd.Timestamp('2024-07-01'), 'Sim', 'Não')

# ===================== planilha =====================
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
HD = PatternFill('solid', fgColor='1F2937'); HU = PatternFill('solid', fgColor='15803D'); HDn = PatternFill('solid', fgColor='B91C1C'); HN = PatternFill('solid', fgColor='6B7280'); HB = PatternFill('solid', fgColor='7C3AED'); HX = PatternFill('solid', fgColor='374151')
YEL = PatternFill('solid', fgColor='FFF3CD'); GRY = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); f1 = S_('1_Upgrade_Downgrade'); f2 = S_('2_Marcacoes_nas_trocas'); r3 = S_('3_Rajadas_no_tempo'); r4 = S_('4_Rajadas_por_segmento'); r5 = S_('5_Rajadas_e_cancelamento')
BS = {}
def py(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v, 'item') else v
def base(nome, df, cols, fmts):
    ws = wb.create_sheet(nome); ws.append(cols)
    for row in df[cols].itertuples(index=False): ws.append([py(v) for v in row])
    n = len(df) + 1; BS[nome] = (cols, n)
    for i, cn in enumerate(cols, 1):
        h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(22, len(cn) + 2))
        if cn in fmts:
            for r in range(2, n + 1): ws.cell(r, i).number_format = fmts[cn]
    tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
def Rg(b, col): cols, n = BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
D = 'dd/mm/yyyy'
base('Base_Assinaturas', S, ['subscription_id', 'account_id', 'industry', 'country', 'referral_source', 'start_date', 'end_date', 'plan_tier', 'seats', 'billing_frequency', 'marcacao',
     'assinatura_anterior', 'plano_anterior', 'seats_anterior', 'plano_vs_anterior', 'seats_vs_anterior', 'subiu_plano_ou_seats', 'desceu_plano_ou_seats', 'em_rajada', 'periodo_inicio', 'trimestre_inicio', 'conta_nova_no_inicio'], {'start_date': D, 'end_date': D})
base('Base_Encerradas_NovDez', E, ['subscription_id', 'account_id', 'industry', 'country', 'velocidade', 'tipo_cancelamento', 'plan_tier', 'seats', 'marcacao', 'em_rajada', 'nova_mais_proxima', 'marcacao_da_nova',
     'plano_nova_vs_cancelada', 'seats_nova_vs_cancelada', 'par_subiu', 'par_desceu'], {})
base('Base_Contas', K, ['account_id', 'industry', 'country', 'referral_source', 'signup_date', 'cadastro_2024_S2', 'assinaturas', 'assinaturas_em_rajada', 'teve_rajada', 'teve_rajada_novdez'], {'signup_date': D})
def H(ws, row, labels, fills, h=40):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fill=None):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX; x.alignment = Alignment(wrap_text=True, vertical='center')
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x
def title(ws, t1, t2, merge='H'):
    ws['A1'] = t1; ws['A1'].font = Fn(bold=True, size=15); ws['A2'] = t2; ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells(f'A2:{merge}2'); ws.row_dimensions[2].height = 30
def note(ws, r, tx, color=GRAY, merge='H'):
    ws.cell(r, 1, tx).font = Fn(italic=True, size=9, color=color); ws.cell(r, 1).alignment = Alignment(wrap_text=True); ws.merge_cells(f'A{r}:{merge}{r}'); ws.row_dimensions[r].height = 30
MK = Rg('Base_Assinaturas', 'marcacao'); MARC = ['Upgrade', 'Downgrade', 'Upgrade e downgrade', 'Nenhuma']; FILL_M = [HU, HDn, HB, HN]

# ---------- 1_Upgrade_Downgrade ----------
ws = f1; title(ws, 'As marcações de upgrade/downgrade batem com mudança real de plano e/ou seats?',
    'Cada assinatura é comparada com a assinatura ANTERIOR da mesma conta (por data de início). A tabela só guarda um plano por assinatura; se a marcação se refere a uma mudança dentro da própria assinatura, ela não aparece nos dados.')
r = 4; ws.cell(r, 1, '1. Resumo: entre as assinaturas marcadas, em quantas o plano e/ou os seats realmente subiram (ou desceram)?').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Marcação', 'Assinaturas (com anterior)', 'Plano subiu', '%', 'Seats subiram', '%', 'Plano E/OU seats subiram', '%', 'Plano E/OU seats desceram', '%'], [HD, HD, HU, HU, HU, HU, HU, HU, HDn, HDn]); r += 1; r1 = r
PV, SV, SU, SD = Rg('Base_Assinaturas', 'plano_vs_anterior'), Rg('Base_Assinaturas', 'seats_vs_anterior'), Rg('Base_Assinaturas', 'subiu_plano_ou_seats'), Rg('Base_Assinaturas', 'desceu_plano_ou_seats')
for m, fl in zip(MARC, FILL_M):
    C(ws, r, 1, m, bold=True); ws.cell(r, 1).fill = fl; ws.cell(r, 1).font = Fn(size=9, bold=True, color='FFFFFF')
    C(ws, r, 2, f'=COUNTIFS({MK},A{r},{PV},"<>Primeira da conta")', '#,##0')
    C(ws, r, 3, f'=COUNTIFS({MK},A{r},{PV},"Subiu")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}', '0%')
    C(ws, r, 5, f'=COUNTIFS({MK},A{r},{SV},"Subiu")', '#,##0'); C(ws, r, 6, f'=E{r}/B{r}', '0%')
    C(ws, r, 7, f'=COUNTIFS({MK},A{r},{SU},"Sim")', '#,##0'); C(ws, r, 8, f'=G{r}/B{r}', '0%')
    C(ws, r, 9, f'=COUNTIFS({MK},A{r},{SD},"Sim")', '#,##0'); C(ws, r, 10, f'=I{r}/B{r}', '0%'); r += 1
X = S[S.plano_anterior.notna()]
ku = int(((X.marcacao == 'Upgrade') & (X.subiu_plano_ou_seats == 'Sim')).sum()); nu = int((X.marcacao == 'Upgrade').sum()); kn = int(((X.marcacao == 'Nenhuma') & (X.subiu_plano_ou_seats == 'Sim')).sum()); nn = int((X.marcacao == 'Nenhuma').sum())
kd = int(((X.marcacao == 'Downgrade') & (X.desceu_plano_ou_seats == 'Sim')).sum()); nd = int((X.marcacao == 'Downgrade').sum()); kn2 = int(((X.marcacao == 'Nenhuma') & (X.desceu_plano_ou_seats == 'Sim')).sum())
r += 1
C(ws, r, 1, 'Upgrade sobe mais que "Nenhuma"?', bold=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); C(ws, r, 7, conf(ku, nu, kn, nn)); ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=10); r += 1
C(ws, r, 1, 'Downgrade desce mais que "Nenhuma"?', bold=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); C(ws, r, 7, conf(kd, nd, kn2, nn)); ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=10); r += 1
note(ws, r, 'Leitura: se a marcação de upgrade fosse fiel, quase todas as assinaturas marcadas teriam plano e/ou seats acima da anterior, bem mais do que as sem marcação. Não é o que acontece: os percentuais de "subiu" e "desceu" são praticamente os mesmos em todas as marcações. As flags de upgrade/downgrade não refletem o plano e os seats registrados.', merge='J'); r += 2
ws.cell(r, 1, '2. Detalhe: plano × seats em relação à assinatura anterior').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Plano vs anterior', 'Seats vs anterior'] + MARC, [HD, HD] + FILL_M); r += 1
for pv in ['Subiu', 'Igual', 'Desceu']:
    for sv in ['Subiu', 'Igual', 'Desceu']:
        C(ws, r, 1, pv, bold=True); C(ws, r, 2, sv, bold=True)
        for j, m in enumerate(MARC, 3): C(ws, r, j, f'=COUNTIFS({MK},"{m}",{PV},A{r},{SV},B{r})', '#,##0')
        r += 1
C(ws, r, 1, 'Primeira da conta', bold=True); C(ws, r, 2, '—')
for j, m in enumerate(MARC, 3): C(ws, r, j, f'=COUNTIFS({MK},"{m}",{PV},"Primeira da conta")', '#,##0')
r += 1; note(ws, r, 'Atenção também: há assinaturas marcadas como upgrade/downgrade que são a PRIMEIRA da conta (não existe plano anterior para subir ou descer).', merge='J')
for col_, w in zip('ABCDEFGHIJ', [22, 14, 11, 7, 11, 7, 12, 7, 12, 7]): ws.column_dimensions[col_].width = w
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.title = '% em que plano e/ou seats subiram vs desceram, por marcação'; ch.height = 7; ch.width = 15
ch.add_data(Reference(ws, min_col=8, min_row=r1 - 1, max_row=r1 + 3), titles_from_data=True); ch.add_data(Reference(ws, min_col=10, min_row=r1 - 1, max_row=r1 + 3), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r1, max_row=r1 + 3)); ch.series[0].graphicalProperties.solidFill = '15803D'; ch.series[1].graphicalProperties.solidFill = 'B91C1C'; ch.legend.position = 'b'; ch.y_axis.number_format = '0%'; ws.add_chart(ch, 'L4')

# ---------- 2_Marcacoes_nas_trocas ----------
ws = f2; title(ws, 'Nas trocas de nov/dez: a marcação acompanha o que mudou entre a assinatura cancelada e a nova?', f'Só as trocas/ajustes (assinatura nova até {JP} dias do cancelamento, mesma regra da planilha 04). Compara a nova com a cancelada.')
EB = 'Base_Encerradas_NovDez'; TP = Rg(EB, 'tipo_cancelamento'); MN = Rg(EB, 'marcacao_da_nova'); PS = Rg(EB, 'par_subiu'); PD_ = Rg(EB, 'par_desceu')
r = 4; H(ws, r, ['Marcação da assinatura NOVA', 'Trocas', 'Plano e/ou seats subiram', '%', 'Plano e/ou seats desceram', '%'], [HD, HD, HU, HU, HDn, HDn]); r += 1
for m, fl in zip(MARC, FILL_M):
    C(ws, r, 1, m, bold=True); ws.cell(r, 1).fill = fl; ws.cell(r, 1).font = Fn(size=9, bold=True, color='FFFFFF')
    C(ws, r, 2, f'=COUNTIFS({TP},"Troca/ajuste",{MN},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({TP},"Troca/ajuste",{MN},A{r},{PS},"Sim")', '#,##0'); C(ws, r, 4, f'=IF(B{r}>0,C{r}/B{r},"")', '0%')
    C(ws, r, 5, f'=COUNTIFS({TP},"Troca/ajuste",{MN},A{r},{PD_},"Sim")', '#,##0'); C(ws, r, 6, f'=IF(B{r}>0,E{r}/B{r},"")', '0%'); r += 1
C(ws, r, 1, 'Total', bold=True, fill=GRY); C(ws, r, 2, f'=SUM(B5:B{r-1})', '#,##0', True, GRY); C(ws, r, 3, f'=SUM(C5:C{r-1})', '#,##0', True, GRY); C(ws, r, 4, f'=C{r}/B{r}', '0%', True, GRY); C(ws, r, 5, f'=SUM(E5:E{r-1})', '#,##0', True, GRY); C(ws, r, 6, f'=E{r}/B{r}', '0%', True, GRY)
r += 2; note(ws, r, 'Leitura: a maioria das trocas não tem marcação nenhuma, mesmo quando o plano e/ou os seats mudaram. E nas poucas marcadas, "upgrade" aparece tanto em trocas que subiram quanto nas que desceram. Para entender upgrade/downgrade de verdade, a fonte confiável é comparar plano e seats entre assinaturas, não a marcação.', merge='F')
for col_, w in zip('ABCDEF', [26, 10, 14, 8, 14, 8]): ws.column_dimensions[col_].width = w

# ---------- 3_Rajadas_no_tempo ----------
ws = r3; title(ws, f'Compras em rajada: {RAJ_N} ou mais assinaturas da mesma conta iniciadas dentro de {RAJ_D} dias', 'Por trimestre de início. "Em rajada" = a assinatura faz parte de um grupo desses.')
ER, TR = Rg('Base_Assinaturas', 'em_rajada'), Rg('Base_Assinaturas', 'trimestre_inicio')
H(ws, 4, ['Trimestre de início', 'Assinaturas iniciadas', 'Em rajada', '% em rajada'], [HD, HD, HB, HB]); r = 5
for q in sorted(S.trimestre_inicio.unique()):
    C(ws, r, 1, q, bold=True); C(ws, r, 2, f'=COUNTIFS({TR},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({TR},A{r},{ER},"Sim")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}', '0%'); r += 1
rq = r - 1; r += 1
H(ws, r, ['Conta nova (< 30 dias) ao assinar?', 'Assinaturas', 'Em rajada', '% em rajada'], [HD, HD, HB, HB]); r += 1
for v in ['Sim', 'Não']:
    C(ws, r, 1, v, bold=True); C(ws, r, 2, f'=COUNTIFS({Rg("Base_Assinaturas","conta_nova_no_inicio")},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({Rg("Base_Assinaturas","conta_nova_no_inicio")},A{r},{ER},"Sim")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}', '0%'); r += 1
r += 1; note(ws, r, 'Leitura: a rajada é um fenômeno do 4º trimestre de 2024 e de contas recém-cadastradas: quase 3 em cada 10 assinaturas do 4T/24 entram em rajada, contra menos de 1 em 20 até o 2T/24.', merge='F')
ch = BarChart(); ch.type = 'col'; ch.title = '% das assinaturas iniciadas em rajada, por trimestre'; ch.height = 7; ch.width = 15; ch.legend = None
ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=rq), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rq)); ch.series[0].graphicalProperties.solidFill = '7C3AED'; ch.y_axis.number_format = '0%'; ws.add_chart(ch, 'F4')
for col_, w in zip('ABCD', [30, 12, 10, 10]): ws.column_dimensions[col_].width = w

# ---------- 4_Rajadas_por_segmento ----------
ws = r4; title(ws, 'A rajada se concentra em algum país, indústria ou canal?', 'Duas visões: (A) assinaturas iniciadas em nov/dez que entraram em rajada; (B) contas que tiveram rajada em nov/dez. "Confiável?" compara a categoria com todas as outras juntas.', merge='J')
PI = Rg('Base_Assinaturas', 'periodo_inicio'); TRN = Rg('Base_Contas', 'teve_rajada_novdez')
r = 4
SN_ = S[S.periodo_inicio == 'Nov–Dez/24']
for col, lab, colK in [('country', 'País', 'country'), ('industry', 'Indústria', 'industry'), ('referral_source', 'Canal de aquisição', 'referral_source')]:
    ws.cell(r, 1, lab).font = Fn(bold=True, size=12); r += 1
    H(ws, r, [lab, '(A) Assinaturas nov/dez', 'Em rajada', '% em rajada', 'Confiável?', '', '(B) Contas', 'Com rajada em nov/dez', '% com rajada', 'Confiável?'], [HD, HD, HB, HB, HX, HD, HD, HB, HB, HX]); r += 1
    for v in sorted(S[col].unique()):
        k1 = int(((SN_[col] == v) & (SN_.em_rajada == 'Sim')).sum()); n1 = int((SN_[col] == v).sum()); k2 = int((SN_.em_rajada == 'Sim').sum()) - k1; n2 = len(SN_) - n1
        kk1 = int(((K[colK] == v) & (K.teve_rajada_novdez == 'Sim')).sum()); nn1 = int((K[colK] == v).sum()); kk2 = int((K.teve_rajada_novdez == 'Sim').sum()) - kk1; nn2 = len(K) - nn1
        C(ws, r, 1, v, bold=True)
        C(ws, r, 2, f'=COUNTIFS({PI},"Nov–Dez/24",{Rg("Base_Assinaturas", col)},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({PI},"Nov–Dez/24",{Rg("Base_Assinaturas", col)},A{r},{ER},"Sim")', '#,##0'); C(ws, r, 4, f'=IF(B{r}>0,C{r}/B{r},"")', '0%')
        C(ws, r, 5, conf(k1, n1, k2, n2)); C(ws, r, 6, '')
        C(ws, r, 7, f'=COUNTIFS({Rg("Base_Contas", colK)},A{r})', '#,##0'); C(ws, r, 8, f'=COUNTIFS({Rg("Base_Contas", colK)},A{r},{TRN},"Sim")', '#,##0'); C(ws, r, 9, f'=IF(G{r}>0,H{r}/G{r},"")', '0%'); C(ws, r, 10, conf(kk1, nn1, kk2, nn2)); r += 1
    C(ws, r, 1, 'Todos', bold=True, fill=GRY); C(ws, r, 2, f'=COUNTIFS({PI},"Nov–Dez/24")', '#,##0', True, GRY); C(ws, r, 3, f'=COUNTIFS({PI},"Nov–Dez/24",{ER},"Sim")', '#,##0', True, GRY); C(ws, r, 4, f'=C{r}/B{r}', '0%', True, GRY)
    C(ws, r, 7, f'=COUNTA({Rg("Base_Contas","account_id")})', '#,##0', True, GRY); C(ws, r, 8, f'=COUNTIF({TRN},"Sim")', '#,##0', True, GRY); C(ws, r, 9, f'=H{r}/G{r}', '0%', True, GRY); r += 2
ws.conditional_formatting.add(f'E4:E{r}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL)); ws.conditional_formatting.add(f'J4:J{r}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
note(ws, r, 'Leitura: acima da média (diferença confiável): DevTools e canal orgânico. Abaixo: FinTech, canal partner e Alemanha (DE). Mesmo assim, a rajada aparece em todos os países, indústrias e canais: ela se concentra mais no TEMPO (4T/24) e em contas novas do que num segmento.', merge='J')
for col_, w in zip('ABCDEFGHIJ', [16, 12, 10, 10, 24, 2, 10, 12, 10, 24]): ws.column_dimensions[col_].width = w

# ---------- 5_Rajadas_e_cancelamento ----------
ws = r5; title(ws, 'A rajada tem a ver com os cancelamentos de nov/dez?', 'As 266 assinaturas encerradas em nov/dez, separadas por estarem ou não em rajada.')
VE, TPc, RJ = Rg(EB, 'velocidade'), Rg(EB, 'tipo_cancelamento'), Rg(EB, 'em_rajada')
H(ws, 4, ['Encerrada estava em rajada?', 'Encerradas', 'Rápidas (< 30 dias)', '%', 'Troca/ajuste', '%', 'Redução', '%'], [HD, HD, HB, HB, HU, HU, HDn, HDn]); r = 5
for v in ['Sim', 'Não']:
    C(ws, r, 1, v, bold=True); C(ws, r, 2, f'=COUNTIFS({RJ},A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({RJ},A{r},{VE},"Rápida (< {LIMITE} dias)")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}', '0%')
    C(ws, r, 5, f'=COUNTIFS({RJ},A{r},{TPc},"Troca/ajuste")', '#,##0'); C(ws, r, 6, f'=E{r}/B{r}', '0%'); C(ws, r, 7, f'=COUNTIFS({RJ},A{r},{TPc},"Redução")', '#,##0'); C(ws, r, 8, f'=G{r}/B{r}', '0%'); r += 1
kr = int(((E.em_rajada == 'Sim') & E.velocidade.str.startswith('Rápida')).sum()); nr = int((E.em_rajada == 'Sim').sum()); knr = int(((E.em_rajada == 'Não') & E.velocidade.str.startswith('Rápida')).sum()); nnr = int((E.em_rajada == 'Não').sum())
r += 1; C(ws, r, 1, 'Rápida é mais comum dentro da rajada? Confiável?', bold=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); C(ws, r, 4, conf(kr, nr, knr, nnr)); ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
r += 2; note(ws, r, 'Leitura: assinaturas compradas em rajada são canceladas rápido com mais frequência e quase sempre como troca/ajuste: a conta compra várias de uma vez e logo desfaz uma parte. Isso aponta para o processo de contratação (pedido duplicado, configuração errada, teste de planos) mais do que para insatisfação.')
for col_, w in zip('ABCDEFGH', [26, 11, 12, 8, 11, 8, 11, 8]): ws.column_dimensions[col_].width = w

# ---------- LEIA-ME ----------
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
pct = lambda k, n: f'{k/n*100:.0f}%' if n else '-'
seg_sim = []
for col, lab in [('country', 'País'), ('industry', 'Indústria'), ('referral_source', 'Canal')]:
    for v in sorted(S[col].unique()):
        k1 = int(((SN_[col] == v) & (SN_.em_rajada == 'Sim')).sum()); n1 = int((SN_[col] == v).sum()); k2 = int((SN_.em_rajada == 'Sim').sum()) - k1; n2 = len(SN_) - n1
        if conf(k1, n1, k2, n2) == 'Sim': seg_sim.append(f'{lab} {v}: {pct(k1, n1)} das assinaturas de nov/dez em rajada (média {pct(int((SN_.em_rajada=="Sim").sum()), len(SN_))})')
Lm = [('Upgrade/downgrade e compras em rajada', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. Só quantidades.', dict(italic=True, color=GRAY)),
      ('', {}), ('PERGUNTAS', dict(bold=True, size=12)),
      ('1. As marcações de upgrade/downgrade batem com o que aconteceu com o plano e/ou os seats?  → abas 1 e 2', {}),
      (f'2. As compras em rajada ({RAJ_N}+ assinaturas da mesma conta em {RAJ_D} dias) são de algum país, indústria ou canal?  → abas 3, 4 e 5', {}),
      ('', {}), ('RESPOSTAS', dict(bold=True, size=12)),
      (f'1. NÃO batem. Entre as assinaturas marcadas como upgrade, o plano e/ou os seats subiram em relação à assinatura anterior em {pct(ku, nu)} dos casos; entre as SEM marcação, em {pct(kn, nn)}. Downgrade: desceram em {pct(kd, nd)}, contra {pct(kn2, nn)} sem marcação. Ou seja, a marcação se comporta como se fosse aleatória, e há upgrades/downgrades marcados na primeira assinatura da conta (sem anterior para comparar).', {}),
      ('   Nas trocas de nov/dez, a maioria não tem marcação, mesmo quando plano e/ou seats mudaram. Para medir upgrade/downgrade, use a comparação de plano e seats entre assinaturas (aba Base_Assinaturas), não as flags.', {}),
      ('   Ressalva: a tabela guarda um único plano por assinatura. Se a marcação se refere a uma mudança dentro da própria assinatura, isso não aparece nos dados e não dá para conferir.', dict(italic=True, color=GRAY)),
      (f'2. A rajada é sobretudo um fenômeno de TEMPO e de CONTA NOVA: cresce forte no 4º trimestre de 2024 e acontece em contas recém-cadastradas. Ela aparece em todos os países, indústrias e canais.', {}),
      ('   Destaques com diferença confiável em nov/dez: ' + ('; '.join(seg_sim) if seg_sim else 'nenhum') + '.', {}),
      ('3. Assinaturas compradas em rajada são canceladas rápido com mais frequência e quase sempre como troca/ajuste. Isso reforça a leitura de que o "churn rápido" de nov/dez nasce no processo de contratação (pedidos duplicados, configuração, teste de planos) e não em insatisfação.', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      (f'• Base fictícia. "Rajada" = {RAJ_N} ou mais assinaturas em {RAJ_D} dias: é uma definição escolhida para a análise. "Anterior" = assinatura anterior da mesma conta pela data de início.', {})]
for i, (tx, st) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st) if st else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT)
print('upgrade', ku, nu, 'nenhuma', kn, nn, conf(ku, nu, kn, nn), '| downgrade', kd, nd, kn2, conf(kd, nd, kn2, nn))
print('primeira com marcacao', S[(S.plano_vs_anterior == 'Primeira da conta') & (S.marcacao != 'Nenhuma')].shape[0])
tt = E[E.tipo_cancelamento == 'Troca/ajuste']; print('trocas por marcacao nova', tt.marcacao_da_nova.value_counts().to_dict())
print(pd.crosstab(tt.marcacao_da_nova, tt.par_subiu))
print('seg sim', seg_sim)
print('rajada x rapida', kr, nr, knr, nnr, conf(kr, nr, knr, nnr), E.groupby('em_rajada').tipo_cancelamento.value_counts().to_dict())
print(S.groupby('conta_nova_no_inicio').em_rajada.apply(lambda x: (x == 'Sim').mean()).round(3).to_dict())
print(K.groupby('country').teve_rajada_novdez.apply(lambda x: (x == 'Sim').mean()).round(2).to_dict())
