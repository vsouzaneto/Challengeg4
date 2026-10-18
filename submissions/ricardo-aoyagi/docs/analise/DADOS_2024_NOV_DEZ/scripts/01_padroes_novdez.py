# Nov/Dez 2024: existe um caminho Perfil -> Reclamação -> Feature (com/sem erro) -> Cancelamento rápido?
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava uma planilha apartada. Só quantidades (sem US$).
import sys
from pathlib import Path
from math import lgamma, exp
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '01_Churn_NovDez_Padroes.xlsx'
INI, FIM = pd.Timestamp('2024-11-01'), pd.Timestamp('2024-12-31'); LIMITE = 30   # "rápido" = menos de 30 dias

a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv', parse_dates=['usage_date'])
t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv', parse_dates=['submitted_at','closed_at'])
c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])

def fisher_p(a_, b_, c_, d_):
    n1, n2, m1 = a_ + b_, c_ + d_, a_ + c_; N = n1 + n2
    if N == 0 or n1 == 0 or n2 == 0: return 1.0
    lc = lambda n, k: lgamma(n + 1) - lgamma(k + 1) - lgamma(n - k + 1)
    pmf = lambda x: exp(lc(n1, x) + lc(n2, m1 - x) - lc(N, m1))
    p0 = pmf(a_); lo, hi = max(0, m1 - n2), min(m1, n1)
    return min(1.0, sum(pmf(x) for x in range(lo, hi + 1) if pmf(x) <= p0 * (1 + 1e-7)))
CONF = lambda p: 'Sim' if p < 0.05 else 'Não (pode ser coincidência)'
SN = lambda b: np.where(b, 'Sim', 'Não')

# ===================== FONTE A: assinaturas =====================
S = s.merge(a[['account_id', 'industry', 'country', 'referral_source']], on='account_id')
S['dias_de_vida'] = (S.end_date - S.start_date).dt.days
enc_nd = S.end_date.between(INI, FIM)
S['grupo'] = np.select([enc_nd & (S.dias_de_vida < LIMITE), enc_nd, S.end_date.isna()],
                       [f'Rápida (< {LIMITE} dias)', f'Lenta ({LIMITE}+ dias)', 'Ativa em 31/12'], 'Fora do recorte')
S = S[S.grupo != 'Fora do recorte'].copy()
GR, GL, GA = f'Rápida (< {LIMITE} dias)', f'Lenta ({LIMITE}+ dias)', 'Ativa em 31/12'
S['trial'] = SN(S.is_trial)
# tickets da conta em nov/dez (mesma janela para todos os grupos)
tnd = t[t.submitted_at.between(INI, FIM + pd.Timedelta(days=1))]
S['tickets_conta_novdez'] = S.account_id.map(tnd.groupby('account_id').size()).fillna(0).astype(int)
S['teve_ticket_novdez'] = SN(S.tickets_conta_novdez > 0)
# uso ligado à assinatura (qualquer data: as datas de uso não são confiáveis)
fs = f.merge(S[['subscription_id']], on='subscription_id')
S['registros_uso'] = S.subscription_id.map(fs.groupby('subscription_id').size()).fillna(0).astype(int)
S['registros_com_erro'] = S.subscription_id.map(fs[fs.error_count > 0].groupby('subscription_id').size()).fillna(0).astype(int)
S['usou_feature_com_erro'] = SN(S.registros_com_erro > 0)
S['usou_beta'] = SN(S.subscription_id.map(fs[fs.is_beta_feature].groupby('subscription_id').size()).fillna(0) > 0)
# evento de churn declarado na conta em nov/dez
cnd = c[c.churn_date.between(INI, FIM)]
S['conta_tem_evento_novdez'] = SN(S.account_id.isin(cnd.account_id))

# ===================== FONTE B: eventos de churn =====================
B = cnd.merge(a[['account_id', 'industry', 'country', 'referral_source', 'plan_tier', 'signup_date']], on='account_id')
B['dias_de_conta'] = (B.churn_date - B.signup_date).dt.days
B['grupo'] = np.where(B.dias_de_conta < LIMITE, f'Conta nova (< {LIMITE} dias)', f'Conta antiga ({LIMITE}+ dias)')
BN, BV = f'Conta nova (< {LIMITE} dias)', f'Conta antiga ({LIMITE}+ dias)'
B['feedback_text'] = B.feedback_text.fillna('(sem comentário)')
B['tickets_conta_novdez'] = B.account_id.map(tnd.groupby('account_id').size()).fillna(0).astype(int)
B['teve_ticket_novdez'] = SN(B.tickets_conta_novdez > 0)
B['tickets_30d_antes'] = [int(((t.account_id == r.account_id) & (t.submitted_at <= r.churn_date) & (t.submitted_at > r.churn_date - pd.Timedelta(days=30))).sum()) for r in B.itertuples()]
B['conta_tem_assin_rapida_novdez'] = SN(B.account_id.isin(S[S.grupo == GR].account_id))

# ===================== detalhe das assinaturas rápidas =====================
R = S[S.grupo == GR].copy().sort_values('end_date')
def jan(r):
    x = t[(t.account_id == r.account_id) & (t.submitted_at <= r.end_date) & (t.submitted_at > r.end_date - pd.Timedelta(days=30))]
    pr = {'urgent': 4, 'high': 3, 'medium': 2, 'low': 1}
    return pd.Series({'tickets_30d_antes': len(x), 'maior_prioridade_30d': (x.priority.map(pr).idxmax() if False else (max(x.priority, key=lambda p: pr[p]) if len(x) else '')),
                      'escalado_30d': 'Sim' if x.escalation_flag.any() else ('Não' if len(x) else ''), 'menor_satisfacao_30d': x.satisfaction_score.min() if len(x) else None})
R = pd.concat([R, R.apply(jan, axis=1)], axis=1)
def feats(sid, only_err=False):
    x = fs[fs.subscription_id == sid]
    if only_err: x = x[x.error_count > 0]
    return ', '.join(x.feature_name.value_counts().index[:3]) if len(x) else ''
R['features_mais_usadas'] = R.subscription_id.map(lambda z: feats(z)); R['features_com_erro'] = R.subscription_id.map(lambda z: feats(z, True))
ev_last = cnd.sort_values('churn_date').groupby('account_id').last()
R['motivo_evento_novdez'] = R.account_id.map(ev_last.reason_code).fillna(''); R['comentario_evento_novdez'] = R.account_id.map(ev_last.feedback_text).fillna('')

# ===================== uso por feature x assinatura =====================
FS = fs.merge(S[['subscription_id', 'grupo', 'industry']], on='subscription_id').groupby(['subscription_id', 'grupo', 'industry', 'feature_name']).agg(
    registros=('usage_id', 'size'), registros_com_erro=('error_count', lambda x: int((x > 0).sum())), beta=('is_beta_feature', 'max')).reset_index()
FS['beta'] = SN(FS.beta); FS['feature_num'] = FS.feature_name.str.replace('feature_', '').astype(int)

# ===================== testes (para "diferença confiável?") =====================
nR, nA = (S.grupo == GR).sum(), (S.grupo == GA).sum()
feat_rows = []
for fe in sorted(FS.feature_name.unique(), key=lambda z: int(z.split('_')[1])):
    x = FS[FS.feature_name == fe]
    uR = x[x.grupo == GR].subscription_id.nunique(); uA = x[x.grupo == GA].subscription_id.nunique()
    rR, eR = x[x.grupo == GR].registros.sum(), x[x.grupo == GR].registros_com_erro.sum()
    rA, eA = x[x.grupo == GA].registros.sum(), x[x.grupo == GA].registros_com_erro.sum()
    feat_rows.append(dict(feature=fe, conf_uso=CONF(fisher_p(uR, nR - uR, uA, nA - uA)), conf_erro=CONF(fisher_p(eR, rR - eR, eA, rA - eA))))
FT = pd.DataFrame(feat_rows)
PERFIL = [('industry', 'Indústria'), ('plan_tier', 'Plano da assinatura'), ('billing_frequency', 'Cobrança'), ('trial', 'Trial'), ('country', 'País'), ('referral_source', 'Canal de aquisição')]
perf_conf = {}
for col, lab in PERFIL:
    for v in sorted(S[col].unique()):
        k1 = ((S.grupo == GR) & (S[col] == v)).sum(); k2 = ((S.grupo == GA) & (S[col] == v)).sum()
        perf_conf[(col, v)] = CONF(fisher_p(k1, nR - k1, k2, nA - k2))
PERFIL_B = [('industry', 'Indústria'), ('plan_tier', 'Plano inicial da conta'), ('country', 'País'), ('referral_source', 'Canal de aquisição'), ('reason_code', 'Motivo registrado'), ('feedback_text', 'Comentário do cliente')]
nBN, nBV = (B.grupo == BN).sum(), (B.grupo == BV).sum(); perfB_conf = {}
for col, lab in PERFIL_B:
    for v in sorted(B[col].unique()):
        k1 = ((B.grupo == BN) & (B[col] == v)).sum(); k2 = ((B.grupo == BV) & (B[col] == v)).sum()
        perfB_conf[(col, v)] = CONF(fisher_p(k1, nBN - k1, k2, nBV - k2))
def conf2(col, val='Sim', base=S, g1=GR, g2=GA):
    k1 = ((base.grupo == g1) & (base[col] == val)).sum(); n1 = (base.grupo == g1).sum(); k2 = ((base.grupo == g2) & (base[col] == val)).sum(); n2 = (base.grupo == g2).sum()
    return CONF(fisher_p(k1, n1 - k1, k2, n2 - k2))

# ===================== planilha =====================
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
HD = PatternFill('solid', fgColor='1F2937'); HRp = PatternFill('solid', fgColor='C2410C'); HLt = PatternFill('solid', fgColor='9A3412'); HAt = PatternFill('solid', fgColor='1D4ED8'); HX = PatternFill('solid', fgColor='374151')
YEL = PatternFill('solid', fgColor='FFF3CD'); GRY = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); gr = S_('Grupos'); pf = S_('1_Perfil'); rc = S_('2_Reclamacoes'); ft = S_('3_Features'); jn = S_('4_Jornada'); lr = S_('5_Lista_Rapidas'); le = S_('6_Eventos_NovDez')
BS = {}
def py(v):
    if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
    if isinstance(v, pd.Timestamp): return v.to_pydatetime()
    if isinstance(v, (np.floating,)): return round(float(v), 4)
    return v.item() if hasattr(v, 'item') else v
def base(nome, df, cols, fmts):
    ws = wb.create_sheet(nome); ws.append(cols)
    for row in df[cols].itertuples(index=False): ws.append([py(v) for v in row])
    n = len(df) + 1; BS[nome] = (cols, n)
    for i, cn in enumerate(cols, 1):
        h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(26, len(cn) + 2))
        if cn in fmts:
            for r in range(2, n + 1): ws.cell(r, i).number_format = fmts[cn]
    tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
    return ws
def Rg(b, col): cols, n = BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
D = 'dd/mm/yyyy'
colsS = ['subscription_id', 'account_id', 'grupo', 'industry', 'country', 'referral_source', 'plan_tier', 'billing_frequency', 'trial', 'start_date', 'end_date', 'dias_de_vida',
         'tickets_conta_novdez', 'teve_ticket_novdez', 'registros_uso', 'registros_com_erro', 'usou_feature_com_erro', 'usou_beta', 'conta_tem_evento_novdez']
base('Base_Assinaturas', S, colsS, {'start_date': D, 'end_date': D})
base('Base_Uso_Feature', FS, ['subscription_id', 'grupo', 'industry', 'feature_name', 'feature_num', 'registros', 'registros_com_erro', 'beta'], {})
colsB = ['churn_event_id', 'account_id', 'grupo', 'industry', 'country', 'referral_source', 'plan_tier', 'signup_date', 'churn_date', 'dias_de_conta', 'reason_code', 'feedback_text',
         'tickets_conta_novdez', 'teve_ticket_novdez', 'tickets_30d_antes', 'conta_tem_assin_rapida_novdez']
TK = tnd.merge(a[['account_id']], on='account_id'); TK['grupo_conta'] = np.select([TK.account_id.isin(S[S.grupo == GR].account_id), TK.account_id.isin(S[S.grupo == GL].account_id)], ['Conta com assinatura rápida', 'Conta com assinatura lenta'], 'Conta sem cancelamento em nov/dez')
TK['nota_3'] = SN(TK.satisfaction_score == 3); TK['respondeu_satisfacao'] = SN(TK.satisfaction_score.notna()); TK['escalado'] = SN(TK.escalation_flag)
def H(ws, row, labels, fills, h=42):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fill=None):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x
def title(ws, t1, t2):
    ws['A1'] = t1; ws['A1'].font = Fn(bold=True, size=15); ws['A2'] = t2; ws['A2'].font = Fn(italic=True, size=9, color=GRAY)
gS = Rg('Base_Assinaturas', 'grupo'); gB = None

# ---------- Grupos ----------
ws = gr; title(ws, 'Quem estamos olhando: cancelamentos de nov e dez/2024, lado a lado', 'Fonte A = assinaturas encerradas (fonte confiável). Fonte B = eventos de churn (têm motivo e comentário, mas não batem com as assinaturas). Contagens por fórmula.')
H(ws, 4, ['Fonte A: assinaturas', 'Assinaturas', 'Contas', 'O que é'], [HD, HD, HD, HD])
gruposA = [(GR, f'Encerrada em nov/dez com menos de {LIMITE} dias de vida. FOCO DA ANÁLISE.'), (GL, f'Encerrada em nov/dez com {LIMITE} dias ou mais de vida.'), (GA, 'Assinatura que segue ativa no fim de 2024. Serve de referência do "normal".')]
for i, (g_, desc) in enumerate(gruposA):
    r = 5 + i; fill = {GR: HRp, GL: HLt, GA: HAt}[g_]
    C(ws, r, 1, g_, bold=True); ws.cell(r, 1).font = Fn(size=9, bold=True, color='FFFFFF'); ws.cell(r, 1).fill = fill
    C(ws, r, 2, f'=COUNTIFS({gS},A{r})', '#,##0'); C(ws, r, 3, S[S.grupo == g_].account_id.nunique(), '#,##0'); C(ws, r, 4, desc)
ws.cell(8, 1, 'Contas (coluna C) contadas no script: a mesma conta pode ter mais de uma assinatura no grupo.').font = Fn(italic=True, size=8, color=GRAY)
H(ws, 10, ['Fonte B: eventos de churn', 'Eventos', 'Contas', 'O que é'], [HD, HD, HD, HD])
for i, (g_, desc) in enumerate([(BN, f'Evento em nov/dez de conta cadastrada há menos de {LIMITE} dias. FOCO DA ANÁLISE.'), (BV, f'Evento em nov/dez de conta com {LIMITE} dias ou mais.')]):
    r = 11 + i; C(ws, r, 1, g_, bold=True); ws.cell(r, 1).font = Fn(size=9, bold=True, color='FFFFFF'); ws.cell(r, 1).fill = HRp if i == 0 else HLt
    C(ws, r, 2, f'=COUNTIFS(Base_Eventos!$C$2:$C${len(B)+1},A{r})', '#,##0'); C(ws, r, 3, B[B.grupo == g_].account_id.nunique(), '#,##0'); C(ws, r, 4, desc)
H(ws, 15, ['Cruzamento das duas fontes', 'Contas', '', 'Leitura'], [HD, HD, HD, HD])
accR = set(S[S.grupo == GR].account_id); accBN = set(B[B.grupo == BN].account_id); accB = set(B.account_id)
for i, (lab, n, txt) in enumerate([('Contas com assinatura rápida (A)', len(accR), ''), ('… que também têm evento de churn em nov/dez (B)', len(accR & accB), 'As duas fontes concordam que houve churn'),
                                   ('… e cujo evento é de conta nova (< 30 dias)', len(accR & accBN), 'Concordam também que foi rápido'),
                                   ('Contas com evento de conta nova (B)', len(accBN), ''), ('… sem nenhuma assinatura encerrada em nov/dez', len(accBN - set(S[S.grupo.isin([GR, GL])].account_id)), 'Evento sem cancelamento real na tabela de assinaturas')]):
    r = 16 + i; C(ws, r, 1, lab, bold=(txt == '')); C(ws, r, 2, n, '#,##0'); C(ws, r, 4, txt)
ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 12; ws.column_dimensions['C'].width = 10; ws.column_dimensions['D'].width = 80

# ---------- 1_Perfil ----------
ws = pf; title(ws, 'Etapa 1 · Perfil / Área: quem cancela rápido é diferente de quem fica?', 'Distribuição (% do grupo). "Diferença confiável?" compara Rápida vs Ativa: "Não" = a diferença é pequena demais para não ser coincidência.')
H(ws, 4, ['Fonte A · Característica', 'Categoria', 'Rápida: qtd', 'Rápida: %', 'Lenta: qtd', 'Lenta: %', 'Ativa: qtd', 'Ativa: %', 'Rápida − Ativa (pontos %)', 'Diferença confiável?'],
  [HD, HD, HRp, HRp, HLt, HLt, HAt, HAt, HX, HX])
r = 5
for col, lab in PERFIL:
    fc = Rg('Base_Assinaturas', col if col != 'trial' else 'trial'); first = r
    for v in sorted(S[col].unique()):
        C(ws, r, 1, lab if r == first else '', bold=True); C(ws, r, 2, v, bold=True)
        for j, g_ in [(3, GR), (5, GL), (7, GA)]:
            C(ws, r, j, f'=COUNTIFS({gS},"{g_}",{fc},"{v}")', '#,##0'); C(ws, r, j + 1, f'={L(j)}{r}/COUNTIFS({gS},"{g_}")', '0%')
        C(ws, r, 9, f'=(D{r}-H{r})*100', '+0;-0;0'); C(ws, r, 10, perf_conf[(col, v)]); r += 1
    r += 1
la = r - 1
ws.conditional_formatting.add(f'J5:J{la}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
r += 1; ws.cell(r, 1, 'Fonte B · eventos de churn de nov/dez').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Fonte B · Característica', 'Categoria', 'Conta nova: qtd', 'Conta nova: %', 'Conta antiga: qtd', 'Conta antiga: %', '', '', 'Nova − Antiga (pontos %)', 'Diferença confiável?'], [HD, HD, HRp, HRp, HLt, HLt, HD, HD, HX, HX]); r += 1
rb0 = r; gBr = f'Base_Eventos!$C$2:$C${len(B)+1}'
colsBmap = {c_: i for i, c_ in enumerate(colsB)}
def RB(col): i = colsBmap[col] + 1; return f'Base_Eventos!${L(i)}$2:${L(i)}${len(B)+1}'
for col, lab in PERFIL_B:
    first = r
    for v in sorted(B[col].unique()):
        C(ws, r, 1, lab if r == first else '', bold=True); C(ws, r, 2, v, bold=True)
        for j, g_ in [(3, BN), (5, BV)]:
            C(ws, r, j, f'=COUNTIFS({gBr},"{g_}",{RB(col)},"{v}")', '#,##0'); C(ws, r, j + 1, f'={L(j)}{r}/COUNTIFS({gBr},"{g_}")', '0%')
        C(ws, r, 9, f'=(D{r}-F{r})*100', '+0;-0;0'); C(ws, r, 10, perfB_conf[(col, v)]); r += 1
    r += 1
ws.conditional_formatting.add(f'J{rb0}:J{r}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
for col_, w in zip('ABCDEFGHIJ', [26, 22, 9, 9, 9, 9, 9, 9, 12, 24]): ws.column_dimensions[col_].width = w
ws.freeze_panes = 'C5'

# ---------- 2_Reclamacoes ----------
ws = rc; title(ws, 'Etapa 2 · Reclamação: houve tickets antes do cancelamento?', 'Tickets da conta em nov/dez (mesma janela para todos os grupos). Atenção: 54% dos tickets da base têm data anterior ao cadastro da conta; use as datas com cautela.')
H(ws, 4, ['Indicador (por assinatura, olhando a conta)', 'Rápida', 'Lenta', 'Ativa', 'Diferença confiável? (Rápida vs Ativa)'], [HD, HRp, HLt, HAt, HX])
ind = [('Assinaturas no grupo', lambda g_: f'=COUNTIFS({gS},"{g_}")', '#,##0', ''),
       ('% cuja conta abriu ticket em nov/dez', lambda g_: f'=COUNTIFS({gS},"{g_}",{Rg("Base_Assinaturas","teve_ticket_novdez")},"Sim")/COUNTIFS({gS},"{g_}")', '0%', conf2('teve_ticket_novdez')),
       ('Tickets da conta em nov/dez (média)', lambda g_: f'=AVERAGEIFS({Rg("Base_Assinaturas","tickets_conta_novdez")},{gS},"{g_}")', '0.00', '')]
for i, (lab, fn, fm, cf_) in enumerate(ind):
    r = 5 + i; C(ws, r, 1, lab, bold=True)
    for j, g_ in [(2, GR), (3, GL), (4, GA)]: C(ws, r, j, fn(g_), fm)
    C(ws, r, 5, cf_)
r = 9; ws.cell(r, 1, 'Como foram os tickets de nov/dez (por grupo da conta)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Indicador (por ticket)', 'Conta com assinatura rápida', 'Conta com assinatura lenta', 'Conta sem cancelamento em nov/dez'], [HD, HRp, HLt, HAt]); r += 1
base('Base_Tickets_NovDez', TK, ['ticket_id', 'account_id', 'grupo_conta', 'submitted_at', 'priority', 'escalado', 'satisfaction_score', 'respondeu_satisfacao', 'nota_3', 'resolution_time_hours', 'first_response_time_minutes'], {'submitted_at': D})
gT = Rg('Base_Tickets_NovDez', 'grupo_conta'); tg = ['Conta com assinatura rápida', 'Conta com assinatura lenta', 'Conta sem cancelamento em nov/dez']
tind = [('Tickets', lambda g_: f'=COUNTIFS({gT},"{g_}")', '#,##0'),
        ('% urgentes', lambda g_: f'=COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","priority")},"urgent")/{{X}}', '0%'),
        ('% alta prioridade (high)', lambda g_: f'=COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","priority")},"high")/{{X}}', '0%'),
        ('% escalados', lambda g_: f'=COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","escalado")},"Sim")/{{X}}', '0%'),
        ('Satisfação média (1 a 5)', lambda g_: f'=IFERROR(AVERAGEIFS({Rg("Base_Tickets_NovDez","satisfaction_score")},{gT},"{g_}"),"")', '0.00'),
        ('% sem resposta de satisfação', lambda g_: f'=IFERROR(COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","respondeu_satisfacao")},"Não")/{{X}},"")', '0%'),
        ('% com nota 3 (a mais baixa que existe na base)', lambda g_: f'=IFERROR(COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","nota_3")},"Sim")/COUNTIFS({gT},"{g_}",{Rg("Base_Tickets_NovDez","respondeu_satisfacao")},"Sim"),"")', '0%'),
        ('Tempo médio de resolução (h)', lambda g_: f'=IFERROR(AVERAGEIFS({Rg("Base_Tickets_NovDez","resolution_time_hours")},{gT},"{g_}"),"")', '0.0')]
r_tickets = r
for i, (lab, fn, fm) in enumerate(tind):
    rr = r + i; C(ws, rr, 1, lab, bold=True)
    for j, g_ in enumerate(tg, 2):
        fx = fn(g_).replace('{X}', f'{L(j)}${r_tickets}'); C(ws, rr, j, fx, fm)
r = r + len(tind) + 1
ws.cell(r, 1, 'Assinaturas rápidas: tickets nos 30 dias ANTES do cancelamento (pela data registrada)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Indicador', 'Qtd', '% das assinaturas rápidas'], [HD, HRp, HRp]); r += 1
lrR = len(R) + 1
for lab, fx in [('Assinaturas rápidas', f"=COUNTA('5_Lista_Rapidas'!$A$5:$A${lrR+3})"), ('… com ticket nos 30 dias antes', f"=COUNTIF('5_Lista_Rapidas'!$N$5:$N${lrR+3},\">0\")"),
                ('… com ticket urgente nos 30 dias antes', f"=COUNTIF('5_Lista_Rapidas'!$O$5:$O${lrR+3},\"urgent\")"), ('… com ticket escalado nos 30 dias antes', f"=COUNTIF('5_Lista_Rapidas'!$P$5:$P${lrR+3},\"Sim\")"),
                ('… com nota de satisfação 3 nos 30 dias antes', f"=COUNTIF('5_Lista_Rapidas'!$Q$5:$Q${lrR+3},3)"), ('… SEM nenhum ticket nos 30 dias antes', f"=COUNTIF('5_Lista_Rapidas'!$N$5:$N${lrR+3},0)")]:
    C(ws, r, 1, lab, bold=True); C(ws, r, 2, fx, '#,##0')
    if lab == 'Assinaturas rápidas': r0x = r
    C(ws, r, 3, f'=B{r}/B${r0x}', '0%'); r += 1
r += 1; ws.cell(r, 1, 'Fonte B: o que os clientes alegaram nos eventos de nov/dez').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Indicador', 'Conta nova', 'Conta antiga', 'Diferença confiável?'], [HD, HRp, HLt, HX]); r += 1
for lab, col, val, fm in [('Eventos', None, None, '#,##0'), ('% cuja conta abriu ticket em nov/dez', 'teve_ticket_novdez', 'Sim', '0%'), ('% sem comentário do cliente', 'feedback_text', '(sem comentário)', '0%')]:
    C(ws, r, 1, lab, bold=True)
    for j, g_ in [(2, BN), (3, BV)]:
        C(ws, r, j, f'=COUNTIFS({gBr},"{g_}")' if col is None else f'=COUNTIFS({gBr},"{g_}",{RB(col)},"{val}")/COUNTIFS({gBr},"{g_}")', fm)
    C(ws, r, 4, '' if col is None else conf2(col, val, B, BN, BV)); r += 1
ws.cell(r, 1, 'Motivos e comentários por grupo estão na aba 1_Perfil (parte Fonte B).').font = Fn(italic=True, size=9, color=GRAY)
ws.cell(r + 1, 1, 'Atenção: em toda a base as notas de satisfação vão só de 3 a 5 (nenhum 1 ou 2) e ~41% dos tickets ficam sem nota. A pesquisa de satisfação não consegue mostrar insatisfação.').font = Fn(italic=True, size=9, color='B45309')
ws.column_dimensions['A'].width = 48
for c_ in 'BCDE': ws.column_dimensions[c_].width = 20

# ---------- 3_Features ----------
ws = ft; title(ws, 'Etapa 3 · Features: o que as assinaturas rápidas usaram, e com que erro?', 'Uso ligado à assinatura (as datas de uso não são confiáveis, então não dá para separar "antes" e "depois"). Contagens por fórmula sobre Base_Uso_Feature.')
gF = Rg('Base_Uso_Feature', 'grupo'); fF = Rg('Base_Uso_Feature', 'feature_name')
H(ws, 4, ['Resumo', 'Rápida', 'Lenta', 'Ativa', 'Diferença confiável? (Rápida vs Ativa)'], [HD, HRp, HLt, HAt, HX])
summ = [('Registros de uso por assinatura (média)', lambda g_: f'=AVERAGEIFS({Rg("Base_Assinaturas","registros_uso")},{gS},"{g_}")', '0.0', ''),
        ('% de registros de uso com erro', lambda g_: f'=SUMIFS({Rg("Base_Uso_Feature","registros_com_erro")},{gF},"{g_}")/SUMIFS({Rg("Base_Uso_Feature","registros")},{gF},"{g_}")', '0.0%',
         CONF(fisher_p(int(FS[FS.grupo == GR].registros_com_erro.sum()), int(FS[FS.grupo == GR].registros.sum() - FS[FS.grupo == GR].registros_com_erro.sum()), int(FS[FS.grupo == GA].registros_com_erro.sum()), int(FS[FS.grupo == GA].registros.sum() - FS[FS.grupo == GA].registros_com_erro.sum())))),
        ('% de assinaturas que tiveram algum erro', lambda g_: f'=COUNTIFS({gS},"{g_}",{Rg("Base_Assinaturas","usou_feature_com_erro")},"Sim")/COUNTIFS({gS},"{g_}")', '0%', conf2('usou_feature_com_erro')),
        ('% de assinaturas que usaram feature beta', lambda g_: f'=COUNTIFS({gS},"{g_}",{Rg("Base_Assinaturas","usou_beta")},"Sim")/COUNTIFS({gS},"{g_}")', '0%', conf2('usou_beta'))]
for i, (lab, fn, fm, cf_) in enumerate(summ):
    r = 5 + i; C(ws, r, 1, lab, bold=True)
    for j, g_ in [(2, GR), (3, GL), (4, GA)]: C(ws, r, j, fn(g_), fm)
    C(ws, r, 5, cf_)
r = 10; ws.cell(r, 1, 'Por feature (ordene pelas colunas para achar destaques)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Feature', 'Rápidas que usaram', '% das rápidas', '% das ativas', 'Uso: rápidas ÷ ativas', 'Uso: diferença confiável?',
          'Registros (rápidas)', 'Com erro (rápidas)', '% com erro (rápidas)', '% com erro (ativas)', 'Erro: rápidas ÷ ativas', 'Erro: diferença confiável?'],
  [HD, HRp, HRp, HAt, HX, HX, HRp, HRp, HRp, HAt, HX, HX]); hr = r; r += 1
f0 = r
for _, x in FT.iterrows():
    fe = f'"{x.feature}"'
    C(ws, r, 1, x.feature, bold=True)
    C(ws, r, 2, f'=COUNTIFS({gF},"{GR}",{fF},{fe})', '#,##0'); C(ws, r, 3, f'=B{r}/COUNTIFS({gS},"{GR}")', '0%'); C(ws, r, 4, f'=COUNTIFS({gF},"{GA}",{fF},{fe})/COUNTIFS({gS},"{GA}")', '0%')
    C(ws, r, 5, f'=IF(D{r}>0,C{r}/D{r},"")', '0.00"x"'); C(ws, r, 6, x.conf_uso)
    C(ws, r, 7, f'=SUMIFS({Rg("Base_Uso_Feature","registros")},{gF},"{GR}",{fF},{fe})', '#,##0'); C(ws, r, 8, f'=SUMIFS({Rg("Base_Uso_Feature","registros_com_erro")},{gF},"{GR}",{fF},{fe})', '#,##0')
    C(ws, r, 9, f'=IF(G{r}>0,H{r}/G{r},"")', '0%')
    C(ws, r, 10, f'=SUMIFS({Rg("Base_Uso_Feature","registros_com_erro")},{gF},"{GA}",{fF},{fe})/SUMIFS({Rg("Base_Uso_Feature","registros")},{gF},"{GA}",{fF},{fe})', '0%')
    C(ws, r, 11, f'=IF(AND(ISNUMBER(I{r}),J{r}>0),I{r}/J{r},"")', '0.00"x"'); C(ws, r, 12, x.conf_erro); r += 1
fl = r - 1
for colx in 'FL': ws.conditional_formatting.add(f'{colx}{f0}:{colx}{fl}', CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
ws.auto_filter.ref = f'A{hr}:L{fl}'
ws.cell(r + 1, 1, 'Leitura: "1,30x" = as rápidas usam (ou erram) 30% mais que as ativas naquela feature. Com ~160 assinaturas rápidas, cada feature tem poucas dezenas de casos; por isso a coluna "diferença confiável?" é importante.').font = Fn(italic=True, size=9, color=GRAY)
for col_, w in zip('ABCDEFGHIJKL', [14, 11, 10, 10, 11, 22, 11, 11, 11, 11, 11, 22]): ws.column_dimensions[col_].width = w
ws.freeze_panes = f'B{hr+1}'
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'; ch.title = '% de registros com erro por feature: rápidas vs ativas'; ch.height = 22; ch.width = 16
ch.add_data(Reference(ws, min_col=9, min_row=hr, max_row=fl), titles_from_data=True); ch.add_data(Reference(ws, min_col=10, min_row=hr, max_row=fl), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=f0, max_row=fl)); ch.series[0].graphicalProperties.solidFill = 'C2410C'; ch.series[1].graphicalProperties.solidFill = '2A78D6'; ch.legend.position = 'b'; ch.x_axis.number_format = '0%'
ws.add_chart(ch, 'N4')

# ---------- 4_Jornada ----------
ws = jn; title(ws, 'Etapa 4 · A jornada: Perfil → Reclamação → Feature com erro → Cancelamento', 'Cada linha é um caminho. Compara quantas assinaturas rápidas seguiram cada caminho com quantas ativas seguiram o mesmo caminho. Se o caminho levasse ao cancelamento, ele seria bem mais comum entre as rápidas.')
H(ws, 4, ['Indústria (perfil)', 'Conta abriu ticket em nov/dez? (reclamação)', 'Assinatura teve erro em feature?', 'Rápidas: qtd', 'Rápidas: % do grupo', 'Ativas: qtd', 'Ativas: % do grupo', 'Rápidas − Ativas (pontos %)'],
  [HD, HD, HD, HRp, HRp, HAt, HAt, HX], 54)
r = 5; ii = Rg('Base_Assinaturas', 'industry'); tt = Rg('Base_Assinaturas', 'teve_ticket_novdez'); ee = Rg('Base_Assinaturas', 'usou_feature_com_erro')
for ind_ in sorted(S.industry.unique()):
    for tk in ['Sim', 'Não']:
        for er in ['Sim', 'Não']:
            C(ws, r, 1, ind_, bold=True); C(ws, r, 2, tk); C(ws, r, 3, er)
            C(ws, r, 4, f'=COUNTIFS({gS},"{GR}",{ii},A{r},{tt},B{r},{ee},C{r})', '#,##0'); C(ws, r, 5, f'=D{r}/COUNTIFS({gS},"{GR}")', '0.0%')
            C(ws, r, 6, f'=COUNTIFS({gS},"{GA}",{ii},A{r},{tt},B{r},{ee},C{r})', '#,##0'); C(ws, r, 7, f'=F{r}/COUNTIFS({gS},"{GA}")', '0.0%')
            C(ws, r, 8, f'=(E{r}-G{r})*100', '+0.0;-0.0;0.0'); r += 1
jl = r - 1
C(ws, r, 1, 'Total', bold=True, fill=GRY); C(ws, r, 4, f'=SUM(D5:D{jl})', '#,##0', True, GRY); C(ws, r, 5, f'=SUM(E5:E{jl})', '0%', True, GRY); C(ws, r, 6, f'=SUM(F5:F{jl})', '#,##0', True, GRY); C(ws, r, 7, f'=SUM(G5:G{jl})', '0%', True, GRY)
ws.conditional_formatting.add(f'H5:H{jl}', CellIsRule(operator='greaterThan', formula=['3'], fill=YEL))
r += 2; ws.cell(r, 1, 'Resumo do caminho completo (sem separar por indústria)').font = Fn(bold=True, size=12); r += 1
H(ws, r, ['Caminho', '', '', 'Rápidas: qtd', 'Rápidas: %', 'Ativas: qtd', 'Ativas: %', 'Rápidas − Ativas (pontos %)'], [HD, HD, HD, HRp, HRp, HAt, HAt, HX]); r += 1
for lab, tk, er in [('Reclamou e teve erro', 'Sim', 'Sim'), ('Reclamou, sem erro', 'Sim', 'Não'), ('Não reclamou, teve erro', 'Não', 'Sim'), ('Não reclamou, sem erro', 'Não', 'Não')]:
    C(ws, r, 1, lab, bold=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    C(ws, r, 4, f'=COUNTIFS({gS},"{GR}",{tt},"{tk}",{ee},"{er}")', '#,##0'); C(ws, r, 5, f'=D{r}/COUNTIFS({gS},"{GR}")', '0%')
    C(ws, r, 6, f'=COUNTIFS({gS},"{GA}",{tt},"{tk}",{ee},"{er}")', '#,##0'); C(ws, r, 7, f'=F{r}/COUNTIFS({gS},"{GA}")', '0%'); C(ws, r, 8, f'=(E{r}-G{r})*100', '+0;-0;0'); r += 1
ws.cell(r + 1, 1, 'Destaque amarelo = caminho pelo menos 3 pontos % mais comum entre as rápidas. Com poucas assinaturas por linha, diferenças pequenas podem ser coincidência.').font = Fn(italic=True, size=9, color=GRAY)
for col_, w in zip('ABCDEFGH', [16, 18, 16, 11, 11, 11, 11, 13]): ws.column_dimensions[col_].width = w

# ---------- 5_Lista_Rapidas ----------
ws = lr; title(ws, f'Lista: as {len(R)} assinaturas encerradas em nov/dez com menos de {LIMITE} dias', 'Uma linha por assinatura, com o que aconteceu antes (tickets) e o que foi usado. Use o filtro dos cabeçalhos.')
colsR = ['subscription_id', 'account_id', 'industry', 'country', 'referral_source', 'plan_tier', 'billing_frequency', 'trial', 'start_date', 'end_date', 'dias_de_vida', 'tickets_conta_novdez',
         'teve_ticket_novdez', 'tickets_30d_antes', 'maior_prioridade_30d', 'escalado_30d', 'menor_satisfacao_30d', 'registros_uso', 'registros_com_erro', 'features_mais_usadas', 'features_com_erro',
         'conta_tem_evento_novdez', 'motivo_evento_novdez', 'comentario_evento_novdez']
nomes = ['Assinatura', 'Conta', 'Indústria', 'País', 'Canal', 'Plano', 'Cobrança', 'Trial', 'Início', 'Fim', 'Dias de vida', 'Tickets da conta nov/dez', 'Ticket em nov/dez?', 'Tickets 30 dias antes',
         'Maior prioridade (30d)', 'Escalado (30d)', 'Menor satisfação (30d)', 'Registros de uso', 'Registros com erro', 'Features mais usadas', 'Features com erro', 'Evento de churn na conta (nov/dez)?', 'Motivo do evento', 'Comentário do evento']
H(ws, 4, nomes, [HD] * len(nomes), 54)
for i, row in enumerate(R[colsR].itertuples(index=False), 5):
    for j, v in enumerate(row, 1):
        x = ws.cell(i, j, py(v)); x.font = Fn(size=9); x.border = BOX
        if colsR[j - 1] in ('start_date', 'end_date'): x.number_format = D
ws.auto_filter.ref = f'A4:{L(len(colsR))}{4+len(R)}'; ws.freeze_panes = 'C5'
for j, w in enumerate([11, 10, 13, 6, 9, 10, 9, 6, 10, 10, 7, 8, 8, 8, 9, 8, 8, 8, 8, 30, 30, 10, 11, 20], 1): ws.column_dimensions[L(j)].width = w

# ---------- 6_Eventos_NovDez ----------
ws = base('Base_Eventos', B, colsB, {'signup_date': D, 'churn_date': D})
ws2 = le; title(ws2, 'Eventos de churn de nov/dez (Fonte B)', 'Os 185 eventos estão na aba Base_Eventos (uma linha por evento, com motivo, comentário e tickets). Resumo abaixo.')
H(ws2, 4, ['Motivo registrado', 'Conta nova', 'Conta antiga', 'Total'], [HD, HRp, HLt, HD])
r = 5
for v in sorted(B.reason_code.unique()):
    C(ws2, r, 1, v, bold=True)
    for j, g_ in [(2, BN), (3, BV)]: C(ws2, r, j, f'=COUNTIFS({gBr},"{g_}",{RB("reason_code")},A{r})', '#,##0')
    C(ws2, r, 4, f'=B{r}+C{r}', '#,##0'); r += 1
C(ws2, r, 1, 'Total', bold=True, fill=GRY)
for j in (2, 3, 4): C(ws2, r, j, f'=SUM({L(j)}5:{L(j)}{r-1})', '#,##0', True, GRY)
r += 2; H(ws2, r, ['Motivo \\ Comentário'] + sorted(B.feedback_text.unique()), [HD] * (1 + B.feedback_text.nunique())); r += 1
for v in sorted(B.reason_code.unique()):
    C(ws2, r, 1, v, bold=True)
    for j, fb in enumerate(sorted(B.feedback_text.unique()), 2): C(ws2, r, j, f'=COUNTIFS({RB("reason_code")},$A{r},{RB("feedback_text")},"{fb}")', '#,##0')
    r += 1
ws2.cell(r + 1, 1, 'Mesmo em nov/dez, o motivo registrado não conversa com o comentário do cliente (ex.: "competitor" com "too expensive").').font = Fn(italic=True, size=9, color=GRAY)
ws2.column_dimensions['A'].width = 22
for j in range(2, 7): ws2.column_dimensions[L(j)].width = 18

# ---------- LEIA-ME ----------
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
Lm = [('Churn de nov e dez/2024: existe um caminho Perfil → Reclamação → Feature → Cancelamento rápido?', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. Só quantidades (valores em US$ ficam para depois).', dict(italic=True, color=GRAY)),
      ('', {}), ('QUEM ESTAMOS OLHANDO (aba Grupos)', dict(bold=True, size=12)),
      (f'• Fonte A (assinaturas, a fonte confiável): {int((S.grupo==GR).sum())} assinaturas encerradas em nov/dez com menos de {LIMITE} dias de vida ("rápidas"), comparadas com {int((S.grupo==GL).sum())} encerradas com {LIMITE}+ dias ("lentas") e com as {int((S.grupo==GA).sum())} ainda ativas (o "normal").', {}),
      (f'• Fonte B (eventos de churn): {int((B.grupo==BN).sum())} eventos de contas com menos de {LIMITE} dias de cadastro, comparados com {int((B.grupo==BV).sum())} de contas mais antigas. Mostramos lado a lado, sabendo que os eventos não batem com as assinaturas.', {}),
      ('', {}), ('COMO LER', dict(bold=True, size=12)),
      ('• Abas 1 a 4 seguem o caminho: Perfil → Reclamação → Feature → Jornada completa. A aba 5 lista cada assinatura rápida; a aba 6 resume os eventos.', {}),
      ('• "Diferença confiável?": "Sim" quando a diferença entre os grupos é grande o bastante para não ser coincidência; "Não" quando pode ser só acaso (comum com poucos casos).', {}),
      ('• Datas de uso e de tickets têm problemas (ver DADOS_RECONCILIA). Por isso o uso é ligado à assinatura sem olhar data, e os tickets são vistos na janela de nov/dez da conta.', {}),
      ('', {}), ('O QUE OS NÚMEROS MOSTRAM', dict(bold=True, size=12)),
      ('• Perfil: as rápidas se parecem com as ativas em indústria, plano, cobrança, trial e país. A única diferença confiável é o canal orgânico (32% das rápidas vs 23% das ativas). Nos eventos (Fonte B), o destaque é outro canal (ads), então as duas fontes não confirmam a mesma coisa. Atenção: em cerca de 70 comparações, 3 ou 4 "Sim" podem aparecer só por acaso.', {}),
      ('• Reclamação: a conta de quem cancela rápido abre tickets na mesma proporção de quem fica (~3 em cada 10 em nov/dez) e com satisfação parecida. Quase 9 em cada 10 assinaturas rápidas cancelam SEM ticket nos 30 dias anteriores.', {}),
      ('• Achado extra: a pesquisa de satisfação da base só tem notas 3, 4 e 5 (nenhum 1 ou 2) e ~41% dos tickets ficam sem nota. Do jeito que está, ela não consegue mostrar cliente insatisfeito.', {}),
      ('• Features: rápidas e ativas usam as mesmas features e erram na mesma proporção (~31% dos registros com erro nos dois grupos). A única feature com diferença confiável (feature_6) aponta para o lado oposto: as rápidas erraram MENOS nela. Não há feature "culpada".', {}),
      ('• Jornada: os 4 caminhos (reclamou ou não × teve erro ou não) aparecem quase na mesma proporção entre rápidas e ativas (diferenças de até 2 pontos %). O caminho mais comum nos dois grupos é "não reclamou, teve erro" (~55%). O maior destaque por indústria (Cybersecurity sem reclamação e com erro: 15% das rápidas vs 11% das ativas) é pequeno e pode ser coincidência.', {}),
      ('• Conclusão: nestes dados não há uma trilha de reclamação ou erro antes do cancelamento rápido. É um cancelamento SILENCIOSO: quem sai rápido não avisa pelo suporte nem mostra mais erro no uso. Isso explica por que o CS não percebe.', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      ('• Base fictícia; datas de uso e tickets pouco confiáveis; grupos pequenos (cerca de 160 assinaturas rápidas).', {}),
      ('• "Rápida" = menos de 30 dias entre início e fim da assinatura. Para eventos, "conta nova" = menos de 30 dias entre o cadastro da conta e o evento.', {})]
for i, (tx, st) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st) if st else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT)
print(S.grupo.value_counts().to_dict(), B.grupo.value_counts().to_dict())
print('feat conf sim:', FT[(FT.conf_uso == 'Sim') | (FT.conf_erro == 'Sim')].to_dict('records'))
print('perfil conf sim:', {k: v for k, v in perf_conf.items() if v == 'Sim'}, {k: v for k, v in perfB_conf.items() if v == 'Sim'})
print('ticket conf', conf2('teve_ticket_novdez'), 'erro conf', conf2('usou_feature_com_erro'), 'beta', conf2('usou_beta'))
