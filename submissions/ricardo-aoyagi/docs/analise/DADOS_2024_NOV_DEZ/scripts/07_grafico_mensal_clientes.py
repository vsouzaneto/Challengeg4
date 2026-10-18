# Mês a mês (jan/23–dez/24): clientes ingressantes, cancelados de verdade, em redução, migrados e normais — em quantidade e em US$.
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada.
import sys
from pathlib import Path
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_2024_NOV_DEZ' / '07_Grafico_Mensal_Clientes.xlsx'
JP = 7
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
# 1) cada assinatura encerrada: troca/ajuste (migração), redução ou saída (mesma regra das planilhas 04 e 06)
def tipo(r):
    o = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
    ativas = o[(o.start_date <= r.end_date) & (o.end_date.isna() | (o.end_date > r.end_date))]
    perto = o[(o.start_date - r.end_date).abs().dt.days <= JP]; depois = o[(o.start_date > r.end_date) & (o.start_date <= r.end_date + pd.Timedelta(days=JP))]
    if len(ativas) == 0 and len(depois) == 0: return 'Saída'
    return 'Migração' if len(perto) else 'Redução'
s['tipo_fim'] = ''
enc = s.end_date.notna(); s.loc[enc, 'tipo_fim'] = s[enc].apply(tipo, axis=1)
primeira = s.groupby('account_id').start_date.min()
# 1b) para cada assinatura encerrada em migração: a assinatura nova mais próxima (até 7 dias) e o MRR reposto (cada nova contada uma vez só)
par = []
for r in s[s.tipo_fim == 'Migração'].sort_values('end_date').itertuples():
    o_ = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
    perto = o_[(o_.start_date - r.end_date).abs().dt.days <= JP]
    p = perto.iloc[(perto.start_date - r.end_date).abs().argsort()].iloc[0]
    par.append(dict(subscription_id=r.subscription_id, nova=p.subscription_id, mrr_nova=int(p.mrr_amount),
                    nova_em_ingressante=p.start_date.to_period('M') == primeira[r.account_id].to_period('M')))
PAR = pd.DataFrame(par); PAR['ja_contada'] = PAR.nova.duplicated()
PAR['mrr_reposto'] = np.where(PAR.ja_contada, 0, PAR.mrr_nova); PAR['mrr_reposto_em_ingr'] = np.where(PAR.nova_em_ingressante, PAR.mrr_reposto, 0)
s = s.merge(PAR[['subscription_id', 'mrr_reposto', 'mrr_reposto_em_ingr']], on='subscription_id', how='left').fillna({'mrr_reposto': 0, 'mrr_reposto_em_ingr': 0})
meses = pd.date_range('2023-01-01', '2024-12-01', freq='MS')
PRIOR = ['Cancelado', 'Redução', 'Migrado', 'Ingressante', 'Normal']
rows = []
for ms in meses:
    me = ms + pd.offsets.MonthEnd(0)
    ativo_no_mes = s[(s.start_date <= me) & (s.end_date.isna() | (s.end_date >= ms))]
    for ac, g in ativo_no_mes.groupby('account_id'):
        fim_mes = g[(g.start_date <= me) & (g.end_date.isna() | (g.end_date > me))]
        enc_mes = g[g.end_date.between(ms, me)]
        ingr = ms <= primeira[ac] <= me
        tipos = set(enc_mes.tipo_fim)
        status = 'Cancelado' if 'Saída' in tipos else 'Redução' if 'Redução' in tipos else 'Migrado' if 'Migração' in tipos else 'Ingressante' if ingr else 'Normal'
        rows.append(dict(mes=ms, conta=ac, status=status, ingressou_no_mes='Sim' if ingr else 'Não',
            mrr_fim_do_mes=int(fim_mes.mrr_amount.sum()),
            mrr_novo_de_ingressante=int(g[g.start_date.between(ms, me)].mrr_amount.sum()) if ingr else 0,
            mrr_cancelado_saida=int(enc_mes[enc_mes.tipo_fim == 'Saída'].mrr_amount.sum()),
            mrr_cancelado_reducao=int(enc_mes[enc_mes.tipo_fim == 'Redução'].mrr_amount.sum()),
            mrr_cancelado_migracao=int(enc_mes[enc_mes.tipo_fim == 'Migração'].mrr_amount.sum()),
            mrr_reposto_migracao=int(enc_mes.mrr_reposto.sum()), mrr_reposto_ja_em_ingressantes=int(enc_mes.mrr_reposto_em_ingr.sum()),
            mrr_cancelado_total=int(enc_mes.mrr_amount.sum()), assinaturas_encerradas=len(enc_mes)))
B = pd.DataFrame(rows)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, Reference
FONT = 'Arial'; GRAY = '52514E'; Fn = lambda **k: Font(name=FONT, **k)
COR = {'Ingressante': '2A78D6', 'Normal': 'C3C2B7', 'Migrado': '1BAF7A', 'Redução': 'EDA100', 'Cancelado': 'E34948'}
HD = PatternFill('solid', fgColor='1F2937'); fillc = lambda k: PatternFill('solid', fgColor=COR[k])
GRY = PatternFill('solid', fgColor='EDEDED'); thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '"US$ "#,##0;-"US$ "#,##0;"-"'
wb = Workbook(); wb.remove(wb.active)
def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
leia = S_('LEIA-ME'); g1 = S_('Grafico_1_Quantidade'); g2 = S_('Grafico_2_Valores')
cols = list(B.columns)
ws = wb.create_sheet('Base_Conta_Mes'); ws.append(cols)
for row in B.itertuples(index=False): ws.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else (v.item() if hasattr(v, 'item') else v) for v in row])
n = len(B) + 1
for i, cn in enumerate(cols, 1):
    h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(12, min(22, len(cn) + 2))
    for r in range(2, n + 1):
        x = ws.cell(r, i)
        if cn == 'mes': x.number_format = 'mmm/yy'
        elif cn.startswith('mrr'): x.number_format = USD
tb = Table(displayName='tBaseContaMes', ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
def R(col): i = cols.index(col) + 1; return f'Base_Conta_Mes!${L(i)}$2:${L(i)}${n}'
MESL = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
def H(ws, row, labels, fills, h=48):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF' if fl not in (fillc('Normal'),) else '1F2937', size=9); c_.fill = fl; c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fill=None):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    return x
def style_chart(ch, cats_ref, series_cols, ws, hdr_row, last, colors, stacked=True, fmt='#,##0'):
    for col_, cor in zip(series_cols, colors):
        ch.add_data(Reference(ws, min_col=col_, min_row=hdr_row, max_row=last), titles_from_data=True); ch.series[-1].graphicalProperties.solidFill = cor; ch.series[-1].graphicalProperties.line.solidFill = cor
    ch.set_categories(cats_ref)
    if ch.legend is not None: ch.legend.position = 'b'
    ch.y_axis.number_format = fmt; ch.y_axis.majorGridlines = ch.y_axis.majorGridlines
    ch.x_axis.delete = False; ch.y_axis.delete = False

# ---------- Gráfico 1: quantidade ----------
ws = g1
ws['A1'] = 'Clientes por situação no mês (quantidade de contas)'; ws['A1'].font = Fn(bold=True, size=15)
ws['A2'] = ('Cada conta que foi cliente no mês aparece UMA vez, na situação de maior prioridade: Cancelado > Redução > Migrado > Ingressante > Normal. '
            'Ex.: uma conta nova que já trocou de assinatura no mesmo mês conta como Migrado (a coluna H mostra todos os ingressantes do mês).')
ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:K2'); ws.row_dimensions[2].height = 30
H(ws, 4, ['Mês', 'Ingressantes', 'Normais', 'Migrados (troca/ajuste — NÃO é churn)', 'Redução (cliente ficou, encerrou assinatura sem repor)', 'Cancelados de verdade (cliente saiu)', 'Total de clientes no mês',
          'Ingressantes (todos, inclusive os que já migraram)', '"Churn" como é contado hoje (contas com assinatura encerrada)', 'Cancelamento real (cancelados)', 'Diferença (o que não era cancelamento)'],
  [HD, fillc('Ingressante'), fillc('Normal'), fillc('Migrado'), fillc('Redução'), fillc('Cancelado'), HD, HD, HD, HD, HD], 60)
st = R('status'); mc = R('mes')
for i, ms in enumerate(meses):
    r = 5 + i; ws.cell(r, 13, ms.to_pydatetime()).number_format = 'mmm/yy'
    C(ws, r, 1, f'{MESL[ms.month-1]}/{str(ms.year)[2:]}', bold=True)
    for j, k in zip((2, 3, 4, 5, 6), ['Ingressante', 'Normal', 'Migrado', 'Redução', 'Cancelado']): C(ws, r, j, f'=COUNTIFS({mc},$M{r},{st},"{k}")', '#,##0')
    C(ws, r, 7, f'=SUM(B{r}:F{r})', '#,##0', True); C(ws, r, 8, f'=COUNTIFS({mc},$M{r},{R("ingressou_no_mes")},"Sim")', '#,##0')
    C(ws, r, 9, f'=COUNTIFS({mc},$M{r},{R("assinaturas_encerradas")},">0")', '#,##0'); C(ws, r, 10, f'=F{r}', '#,##0'); C(ws, r, 11, f'=I{r}-J{r}', '#,##0')
last = 5 + len(meses) - 1; rt = last + 1
C(ws, rt, 1, 'Total 24 meses', bold=True, fill=GRY)
for j in (2, 4, 5, 6, 8, 9, 10, 11): C(ws, rt, j, f'=SUM({L(j)}5:{L(j)}{last})', '#,##0', True, GRY)
for j in (3, 7): C(ws, rt, j, '—', fill=GRY)
ws.column_dimensions['M'].hidden = True
ws.cell(rt + 1, 1, 'Coluna M (oculta) guarda a data do mês usada nas fórmulas. "Normais" e "Total" não se somam entre meses (são as mesmas contas mês após mês).').font = Fn(italic=True, size=8, color=GRAY)
for col_, w in zip('ABCDEFGHIJK', [9, 11, 9, 14, 16, 13, 11, 14, 16, 12, 13]): ws.column_dimensions[col_].width = w
ws.freeze_panes = 'B5'
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 40
ch.title = 'Gráfico 1 · Clientes por situação, mês a mês (quantidade)'; ch.height = 11; ch.width = 30
style_chart(ch, Reference(ws, min_col=1, min_row=5, max_row=last), [2, 3, 4, 5, 6], ws, 4, last, [COR[k] for k in ['Ingressante', 'Normal', 'Migrado', 'Redução', 'Cancelado']])
ws.add_chart(ch, f'A{rt+3}')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 40
ch.title = 'O "churn" de hoje aberto: migração (não é churn) + redução + cancelamento real (contas por mês)'; ch.height = 8; ch.width = 30
style_chart(ch, Reference(ws, min_col=1, min_row=5, max_row=last), [4, 5, 6], ws, 4, last, [COR['Migrado'], COR['Redução'], COR['Cancelado']])
ws.add_chart(ch, f'A{rt+26}')

# ---------- Gráfico 2: valores ----------
ws = g2
ws['A1'] = 'Movimento de MRR por situação, mês a mês (US$ por mês)'; ws['A1'].font = Fn(bold=True, size=15)
ws['A2'] = ('Ingressantes = MRR das assinaturas abertas no mês por contas novas. Migrados: "encerrado" = MRR da assinatura que saiu; "reposto" = MRR da assinatura nova que a conta abriu no lugar (até 7 dias; cada nova contada uma vez); '
            '"saldo" = reposto − encerrado. Redução = MRR encerrado sem reposição, com a conta seguindo cliente. Cancelados = MRR encerrado quando a conta saiu. '
            'PERDA REAL = redução + cancelados − saldo das migrações (o saldo negativo é perda; o positivo compensa). Normais = MRR ativo no fim do mês das contas normais (gráfico separado, pela escala).')
ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:O2'); ws.row_dimensions[2].height = 52
HM = PatternFill('solid', fgColor='0E7C66'); HMg = PatternFill('solid', fgColor='9CA3AF'); HDf = PatternFill('solid', fgColor='EB6834'); HY = PatternFill('solid', fgColor='FFF3CD')
H(ws, 4, ['Mês', 'Ingressantes: MRR novo', 'Migrados: MRR encerrado', 'Migrados: MRR REPOSTO na assinatura nova (não é perda)', 'Migrados: saldo (reposto − encerrado)', 'Redução: MRR perdido', 'Cancelados: MRR perdido',
          'Normais: MRR ativo no fim do mês', '"Churn" de receita como contado hoje', 'PERDA REAL (redução + cancelados − saldo das migrações)', 'O que NÃO era perda ("churn" de hoje − perda real)', 'MRR total ativo no fim do mês',
          'Do reposto, quanto já está em Ingressantes', 'Migração: parte reposta (NÃO é perda)', 'Migração: diferença perdida na troca'],
  [HD, fillc('Ingressante'), HMg, HM, HM, fillc('Redução'), fillc('Cancelado'), fillc('Normal'), HD, fillc('Cancelado'), HM, HD, HD, HD, HD], 72)
for i, ms in enumerate(meses):
    r = 5 + i; ws.cell(r, 17, ms.to_pydatetime()).number_format = 'mmm/yy'
    C(ws, r, 1, f'{MESL[ms.month-1]}/{str(ms.year)[2:]}', bold=True)
    C(ws, r, 2, f'=SUMIFS({R("mrr_novo_de_ingressante")},{mc},$Q{r})', USD); C(ws, r, 3, f'=SUMIFS({R("mrr_cancelado_migracao")},{mc},$Q{r})', USD)
    C(ws, r, 4, f'=SUMIFS({R("mrr_reposto_migracao")},{mc},$Q{r})', USD); C(ws, r, 5, f'=D{r}-C{r}', USD)
    C(ws, r, 6, f'=SUMIFS({R("mrr_cancelado_reducao")},{mc},$Q{r})', USD); C(ws, r, 7, f'=SUMIFS({R("mrr_cancelado_saida")},{mc},$Q{r})', USD)
    C(ws, r, 8, f'=SUMIFS({R("mrr_fim_do_mes")},{mc},$Q{r},{st},"Normal")', USD)
    C(ws, r, 9, f'=SUMIFS({R("mrr_cancelado_total")},{mc},$Q{r})', USD); C(ws, r, 10, f'=F{r}+G{r}-E{r}', USD, True, HY); C(ws, r, 11, f'=I{r}-J{r}', USD)
    C(ws, r, 12, f'=SUMIFS({R("mrr_fim_do_mes")},{mc},$Q{r})', USD); C(ws, r, 13, f'=SUMIFS({R("mrr_reposto_ja_em_ingressantes")},{mc},$Q{r})', USD)
    C(ws, r, 14, f'=MIN(C{r},D{r})', USD); C(ws, r, 15, f'=MAX(C{r}-D{r},0)', USD)
C(ws, rt, 1, 'Total 24 meses', bold=True, fill=GRY)
for j in (2, 3, 4, 5, 6, 7, 9, 10, 11, 13, 14, 15): C(ws, rt, j, f'=SUM({L(j)}5:{L(j)}{last})', USD, True, GRY)
for j in (8, 12): C(ws, rt, j, '—', fill=GRY)
ws.cell(rt + 1, 1, 'Colunas H e L são saldos (não se somam entre meses). Coluna M: parte do reposto que já aparece em Ingressantes (contas novas que trocaram no próprio mês de entrada); o restante do reposto são assinaturas novas de clientes antigos, que não entram em Ingressantes. Colunas N e O só servem ao gráfico de baixo.').font = Fn(italic=True, size=8, color=GRAY)
rr = rt + 3; ws.cell(rr, 1, 'Resumo dos 24 meses').font = Fn(bold=True, size=12); rr += 1
for j_, t_ in ((1, ''), (4, 'MRR (US$)'), (5, 'Leitura')):
    c_ = ws.cell(rr, j_, t_); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = HD
ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3); ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10); rr += 1
for lab, fx, txt, fl in [('"Churn" de receita como contado hoje', f'=I{rt}', 'Tudo o que foi encerrado.', None),
                         ('− reposto nas migrações', f'=-D{rt}', 'Voltou em assinatura nova da mesma conta: NÃO é perda.', None),
                         ('PERDA REAL (resultado)', f'=J{rt}', 'Redução + cancelados + o que as migrações perderam no saldo.', HY),
                         ('   … redução', f'=F{rt}', '', None), ('   … cancelamento real', f'=G{rt}', '', None), ('   … saldo das migrações (troca para assinatura mais barata)', f'=-E{rt}', '', None)]:
    C(ws, rr, 1, lab, bold=True, fill=fl); ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3); C(ws, rr, 4, fx, USD, True, fl)
    C(ws, rr, 5, txt, fill=fl); ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10); rr += 1
ws.column_dimensions['Q'].hidden = True
for col_, w in zip('ABCDEFGHIJKLMNO', [9, 13, 13, 16, 13, 12, 12, 15, 14, 16, 15, 15, 14, 12, 12]): ws.column_dimensions[col_].width = w
ws.column_dimensions['C'].width = 13; ws.freeze_panes = 'B5'
cr = rr + 2
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.gapWidth = 40
ch.title = 'Gráfico 2 · Movimento de MRR por situação, mês a mês (US$)'; ch.height = 11; ch.width = 32
style_chart(ch, Reference(ws, min_col=1, min_row=5, max_row=last), [2, 3, 4, 6, 7], ws, 4, last, [COR['Ingressante'], '9CA3AF', '0E7C66', COR['Redução'], COR['Cancelado']], stacked=False)
ws.add_chart(ch, f'A{cr}')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 40
ch.title = 'O "churn" de receita de hoje aberto: reposto na migração (NÃO é perda) + perdas reais (US$ por mês)'; ch.height = 9; ch.width = 32
style_chart(ch, Reference(ws, min_col=1, min_row=5, max_row=last), [14, 15, 6, 7], ws, 4, last, ['0E7C66', 'EB6834', COR['Redução'], COR['Cancelado']])
ws.add_chart(ch, f'A{cr+23}')
ch = BarChart(); ch.type = 'col'; ch.title = 'Normais: MRR ativo no fim do mês (US$)'; ch.height = 7; ch.width = 32
style_chart(ch, Reference(ws, min_col=1, min_row=5, max_row=last), [8], ws, 4, last, [COR['Normal']]); ch.legend = None
ws.add_chart(ch, f'A{cr+42}')

# ---------- LEIA-ME ----------
tot = B.groupby('status').size(); val = B[['mrr_novo_de_ingressante', 'mrr_cancelado_migracao', 'mrr_reposto_migracao', 'mrr_reposto_ja_em_ingressantes', 'mrr_cancelado_reducao', 'mrr_cancelado_saida', 'mrr_cancelado_total']].sum()
real = val.mrr_cancelado_reducao + val.mrr_cancelado_saida + (val.mrr_cancelado_migracao - val.mrr_reposto_migracao)
fmt = lambda v: f'US$ {v:,.0f}'.replace(',', '.')
ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
Lm = [('Mês a mês (jan/23–dez/24): o que é cancelamento de verdade e o que é migração', dict(bold=True, size=15)),
      ('Planilha apartada. Lê os CSVs originais sem alterá-los. Usa só a tabela de assinaturas (a confiável).', dict(italic=True, color=GRAY)),
      ('', {}), ('SITUAÇÕES DE CADA CLIENTE (CONTA) NO MÊS', dict(bold=True, size=12)),
      ('• Ingressante: a primeira assinatura da conta começou no mês.', {}),
      ('• Migrado: encerrou assinatura no mês, mas abriu outra até 7 dias antes ou depois (troca/ajuste). Aparece no "churn" de hoje, mas NÃO é churn.', {}),
      ('• Redução: encerrou assinatura no mês sem repor e continuou cliente com as demais. Não perdeu o cliente, mas perdeu receita.', {}),
      ('• Cancelado: encerrou assinatura e ficou sem nenhuma ativa (sem reabrir em 7 dias). É a saída real do cliente.', {}),
      ('• Normal: era cliente e não teve nenhum dos movimentos acima.', {}),
      ('Prioridade quando a conta tem mais de um movimento no mês: Cancelado > Redução > Migrado > Ingressante > Normal.', dict(italic=True, color=GRAY)),
      ('Incluímos "Redução" como situação própria: ela não é migração (houve perda de receita) nem cancelamento (o cliente ficou).', dict(italic=True, color=GRAY)),
      ('', {}), ('RESULTADO NOS 24 MESES', dict(bold=True, size=12)),
      (f'• Cancelamentos de verdade (cliente saiu): {int(tot.get("Cancelado", 0))} casos em 24 meses (contas que depois voltaram a assinar). Nenhum em nov/dez/24.', {}),
      (f'• Contas-mês em migração: {int(tot.get("Migrado", 0))}; em redução: {int(tot.get("Redução", 0))}.', {}),
      (f'• Em valores: "churn" de receita como contado hoje = {fmt(val.mrr_cancelado_total)}. Desse total, {fmt(val.mrr_cancelado_migracao)} foram encerrados em migrações, {fmt(val.mrr_cancelado_reducao)} em reduções e {fmt(val.mrr_cancelado_saida)} em cancelamentos reais.', {}),
      (f'• Nas migrações, as assinaturas novas abertas no lugar somam {fmt(val.mrr_reposto_migracao)} de MRR: esse valor NÃO é perda. O saldo das migrações é {fmt(val.mrr_reposto_migracao - val.mrr_cancelado_migracao)} (trocas, na média, para assinaturas um pouco mais baratas).', {}),
      (f'• PERDA REAL em 24 meses = redução + cancelados − saldo das migrações = {fmt(real)} ({real / val.mrr_cancelado_total * 100:.0f}% do "churn" de hoje).', {}),
      (f'• Do MRR reposto, {fmt(val.mrr_reposto_ja_em_ingressantes)} já aparecem em Ingressantes (contas novas que trocaram no mês de entrada); o resto são assinaturas novas de clientes antigos, que não estão em Ingressantes. Não há dupla contagem na perda real.', {}),
      (f'• No mesmo período entraram {fmt(val.mrr_novo_de_ingressante)} de MRR novo só de clientes ingressantes (sem contar novas assinaturas de quem já era cliente).', {}),
      ('', {}), ('ABAS', dict(bold=True, size=12)),
      ('• Grafico_1_Quantidade: tabela + gráfico empilhado das situações + o "churn" de hoje aberto em migração, redução e cancelamento real.', {}),
      ('• Grafico_2_Valores: tabela (com MRR encerrado, REPOSTO e saldo das migrações, e a perda real corrigida) + resumo dos 24 meses + gráfico do movimento de MRR + o "churn" de receita de hoje aberto em "reposto na migração (não é perda)" e perdas reais + MRR dos normais (separado, por ser muito maior).', {}),
      ('• Base_Conta_Mes: uma linha por conta e mês, com a situação e os valores (use para tabelas dinâmicas).', {}),
      ('', {}), ('LIMITES', dict(bold=True, size=12)),
      ('• Base fictícia. "7 dias" é a janela escolhida para chamar de migração (ver sensibilidade na planilha 04). Parte das migrações pode ser compra independente.', {})]
for i, (tx, st_) in enumerate(Lm, 1):
    x = ws.cell(i, 2, tx); x.font = Fn(**st_) if st_ else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
wb.calculation.fullCalcOnLoad = True; wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); print('ok', OUT)
print(tot.to_dict()); print(val.to_dict()); print('perda real', real)
print(B[B.mes >= '2024-07-01'].groupby(['mes', 'status']).size().unstack(fill_value=0).to_string())
