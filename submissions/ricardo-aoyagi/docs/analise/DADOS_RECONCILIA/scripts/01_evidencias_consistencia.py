# Passo 1 da reconciliação: testes de consistência por tabela (gerado com Claude/IA)
# Lê os CSVs originais (só leitura) e cria uma planilha APARTADA com cópias dos dados
# + colunas de teste em FÓRMULAS do Excel (verificáveis linha a linha).
import sys
from pathlib import Path
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[3]                     # G4_DESAFIO
DADOS = RAIZ / 'dados'
OUT = RAIZ / 'analise' / 'DADOS_RECONCILIA' / '01_Evidencias_Consistencia.xlsx'
if len(sys.argv) > 2: DADOS, OUT = Path(sys.argv[1]), Path(sys.argv[2])

a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv', parse_dates=['usage_date'])
t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv', parse_dates=['submitted_at','closed_at'])
c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])

FONT='Arial'; GRAY='52514E'
F=lambda **k: Font(name=FONT,**k)
H_ORIG=PatternFill('solid',fgColor='1F2937'); H_AUX=PatternFill('solid',fgColor='4B5563'); H_TEST=PatternFill('solid',fgColor='B45309')
RED=PatternFill('solid',fgColor='F8D7DA'); GREEN=PatternFill('solid',fgColor='D1E7DD'); TOT=PatternFill('solid',fgColor='EDEDED')
thin=Side(style='thin',color='D0D0D0'); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
wb=Workbook(); wb.remove(wb.active)
def S_(n): ws=wb.create_sheet(n); ws.sheet_view.showGridLines=False; return ws
leia=S_('LEIA-ME'); sc=S_('Scorecard'); cz=S_('Cruzamentos'); ex=S_('Exemplos')

SHEETS={}   # nome -> (colunas, n_linhas)
def py(v):
    if v is None or v is pd.NaT: return None
    if isinstance(v,float) and np.isnan(v): return None
    if isinstance(v,pd.Timestamp): return v.to_pydatetime()
    return v.item() if hasattr(v,'item') else v
def data_sheet(name, df, aux, tests, datefmt):
    """aux/tests: lista de (coluna, descrição, função(r, col)->fórmula)"""
    ws=wb.create_sheet(name); cols=list(df.columns)+[x[0] for x in aux]+[x[0] for x in tests]
    n=len(df)+1; SHEETS[name]=(cols,n)
    ws.append(cols)
    for row in df.itertuples(index=False): ws.append([py(v) for v in row])
    return ws
def R(sheet,col):
    cols,n=SHEETS[sheet]; i=cols.index(col)+1; return f"{sheet}!${L(i)}$2:${L(i)}${n}"
def fill_formulas(ws, name, df, aux, tests, datefmt, widths=14):
    cols,n=SHEETS[name]; base=len(df.columns)
    ref=lambda col,r: f"{L(cols.index(col)+1)}{r}"
    for j,(col,desc,fn) in enumerate(aux+tests, base+1):
        for r in range(2,n+1): ws.cell(r,j,'='+fn(lambda cc: ref(cc,r), r))
    for i,cn in enumerate(cols,1):
        h=ws.cell(1,i); h.font=F(bold=True,color='FFFFFF',size=9); h.alignment=Alignment(wrap_text=True,vertical='center')
        h.fill=H_ORIG if i<=base else (H_AUX if i<=base+len(aux) else H_TEST)
        ws.column_dimensions[L(i)].width=max(12,min(24,len(cn)+2))
        if cn in datefmt:
            for r in range(2,n+1): ws.cell(r,i).number_format=datefmt[cn]
    ws.row_dimensions[1].height=42; ws.freeze_panes='B2'
    tb=Table(displayName='t_'+name,ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo=TableStyleInfo(name='TableStyleLight1',showRowStripes=True); ws.add_table(tb)
    if tests:
        rng=f'{L(base+len(aux)+1)}2:{L(len(cols))}{n}'
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"FALHA"'],fill=RED))
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"OK"'],fill=GREEN))
    for i,(col,desc,fn) in enumerate(aux+tests, base+1):
        from openpyxl.comments import Comment
        ws.cell(1,i).comment=Comment(desc,'Claude')
OKF=lambda cond: f'IF({cond},"OK","FALHA")'
D='dd/mm/yyyy'; DT='dd/mm/yyyy hh:mm'

# ---------- ordem de criação: as abas referenciam umas às outras, então primeiro escrevemos todas ----------
S_aux=[('aux_ativa','Sim se end_date vazio (assinatura ativa em 31/12/2024)',lambda c_,r: f'IF({c_("end_date")}="","Sim","Não")'),
       ('aux_fim_calc','end_date, ou 31/12/2099 se ativa (facilita comparar datas)',lambda c_,r: f'IF({c_("end_date")}="",DATE(2099,12,31),{c_("end_date")})'),
       ('aux_signup_conta','signup_date da conta (accounts)',lambda c_,r: f'INDEX({R("accounts","signup_date")},MATCH({c_("account_id")},{R("accounts","account_id")},0))'),
       ('aux_preco_seat','Preço por seat do plano: Basic 19, Pro 49, Enterprise 199 (US$)',lambda c_,r: f'IF({c_("plan_tier")}="Basic",19,IF({c_("plan_tier")}="Pro",49,IF({c_("plan_tier")}="Enterprise",199,0)))'),
       ('aux_primeira_assin','Sim se é a assinatura mais antiga da conta (empates contam como Sim)',lambda c_,r: f'IF({c_("start_date")}=_xlfn.MINIFS({R("subscriptions","start_date")},{R("subscriptions","account_id")},{c_("account_id")}),"Sim","Não")')]
S_tests=[('S01_churnflag_x_enddate','churn_flag = TRUE exatamente quando end_date está preenchido',lambda c_,r: OKF(f'{c_("churn_flag")}=({c_("end_date")}<>"")')),
       ('S02_fim_apos_inicio','end_date >= start_date (quando preenchido)',lambda c_,r: OKF(f'OR({c_("end_date")}="",{c_("end_date")}>={c_("start_date")})')),
       ('S03_mrr_seats_preco','mrr_amount = seats × preço do plano (trial = 0)',lambda c_,r: OKF(f'{c_("mrr_amount")}=IF({c_("is_trial")},0,{c_("seats")}*{c_("aux_preco_seat")})')),
       ('S04_trial_x_mrr0','is_trial = TRUE exatamente quando mrr_amount = 0',lambda c_,r: OKF(f'{c_("is_trial")}=({c_("mrr_amount")}=0)')),
       ('S05_arr_12x_mrr','arr_amount = 12 × mrr_amount',lambda c_,r: OKF(f'{c_("arr_amount")}=12*{c_("mrr_amount")}')),
       ('S06_inicio_apos_cadastro','start_date >= signup_date da conta',lambda c_,r: OKF(f'{c_("start_date")}>={c_("aux_signup_conta")}')),
       ('S07_conta_existe','account_id existe em accounts',lambda c_,r: OKF(f'COUNTIF({R("accounts","account_id")},{c_("account_id")})=1')),
       ('S08_id_unico','subscription_id aparece uma única vez',lambda c_,r: OKF(f'COUNTIF({R("subscriptions","subscription_id")},{c_("subscription_id")})=1')),
       ('S09_up_e_down_juntos','Não tem upgrade_flag e downgrade_flag TRUE ao mesmo tempo',lambda c_,r: OKF(f'NOT(AND({c_("upgrade_flag")},{c_("downgrade_flag")}))'))]
A_aux=[('aux_qtd_eventos','Qtd. de eventos em churn_events',lambda c_,r: f'COUNTIF({R("churn_events","account_id")},{c_("account_id")})'),
       ('aux_tem_evento','Sim se tem ao menos 1 evento de churn',lambda c_,r: f'IF({c_("aux_qtd_eventos")}>0,"Sim","Não")'),
       ('aux_assin_ativas','Qtd. de assinaturas ativas em 31/12/2024',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","aux_ativa")},"Sim")'),
       ('aux_assin_encerradas','Qtd. de assinaturas encerradas',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","aux_ativa")},"Não")'),
       ('aux_tem_assin_encerrada','Sim se tem ao menos 1 assinatura encerrada',lambda c_,r: f'IF({c_("aux_assin_encerradas")}>0,"Sim","Não")'),
       ('aux_flag_txt','churn_flag como Sim/Não (para cruzamentos)',lambda c_,r: f'IF({c_("churn_flag")},"Sim","Não")')]
A_tests=[('A01_flag_x_evento','churn_flag = TRUE exatamente quando a conta tem evento em churn_events',lambda c_,r: OKF(f'{c_("churn_flag")}=({c_("aux_qtd_eventos")}>0)')),
       ('A02_flag_sem_assin_ativa','Conta com churn_flag = TRUE não deveria ter assinatura ativa (N/A se flag = FALSE)',lambda c_,r: f'IF({c_("churn_flag")},IF({c_("aux_assin_ativas")}=0,"OK","FALHA"),"N/A")'),
       ('A03_plano_x_1a_assin','plan_tier (plano inicial) = plano da primeira assinatura',lambda c_,r: OKF(f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","aux_primeira_assin")},"Sim",{R("subscriptions","plan_tier")},{c_("plan_tier")})>0')),
       ('A04_seats_x_1a_assin','seats = seats da primeira assinatura',lambda c_,r: OKF(f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","aux_primeira_assin")},"Sim",{R("subscriptions","seats")},{c_("seats")})>0'))]
C_aux=[('aux_signup_conta','signup_date da conta',lambda c_,r: f'INDEX({R("accounts","signup_date")},MATCH({c_("account_id")},{R("accounts","account_id")},0))'),
       ('aux_flag_conta','churn_flag da conta (accounts)',lambda c_,r: f'INDEX({R("accounts","churn_flag")},MATCH({c_("account_id")},{R("accounts","account_id")},0))'),
       ('aux_assin_fim_7d','Qtd. de assinaturas da conta com end_date a até 7 dias do churn_date',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","aux_fim_calc")},">="&({c_("churn_date")}-7),{R("subscriptions","aux_fim_calc")},"<="&({c_("churn_date")}+7))'),
       ('aux_assin_ativas_data','Qtd. de assinaturas da conta ativas no churn_date',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","start_date")},"<="&{c_("churn_date")},{R("subscriptions","aux_fim_calc")},">"&{c_("churn_date")})'),
       ('aux_eventos_anteriores','Qtd. de eventos da mesma conta com data anterior',lambda c_,r: f'COUNTIFS({R("churn_events","account_id")},{c_("account_id")},{R("churn_events","churn_date")},"<"&{c_("churn_date")})'),
       ('aux_downgrade_90d','Qtd. de assinaturas com downgrade_flag iniciadas nos 90 dias antes do churn',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","downgrade_flag")},TRUE,{R("subscriptions","start_date")},"<="&{c_("churn_date")},{R("subscriptions","start_date")},">="&({c_("churn_date")}-90))'),
       ('aux_upgrade_90d','Qtd. de assinaturas com upgrade_flag iniciadas nos 90 dias antes do churn',lambda c_,r: f'COUNTIFS({R("subscriptions","account_id")},{c_("account_id")},{R("subscriptions","upgrade_flag")},TRUE,{R("subscriptions","start_date")},"<="&{c_("churn_date")},{R("subscriptions","start_date")},">="&({c_("churn_date")}-90))'),
       ('aux_feedback','feedback_text, ou (vazio)',lambda c_,r: f'IF({c_("feedback_text")}="","(vazio)",{c_("feedback_text")})')]
C_tests=[('C01_fim_assin_7d','Alguma assinatura da conta termina a até 7 dias do churn_date',lambda c_,r: OKF(f'{c_("aux_assin_fim_7d")}>0')),
       ('C02_sem_assin_ativa','A conta não tem assinatura ativa no churn_date (saiu de fato)',lambda c_,r: OKF(f'{c_("aux_assin_ativas_data")}=0')),
       ('C03_flag_conta','A conta do evento tem churn_flag = TRUE em accounts',lambda c_,r: OKF(f'{c_("aux_flag_conta")}')),
       ('C04_reativacao','is_reactivation = TRUE exatamente quando existe evento anterior da conta',lambda c_,r: OKF(f'{c_("is_reactivation")}=({c_("aux_eventos_anteriores")}>0)')),
       ('C05_prec_downgrade','preceding_downgrade_flag = existe downgrade real nos 90 dias anteriores',lambda c_,r: OKF(f'{c_("preceding_downgrade_flag")}=({c_("aux_downgrade_90d")}>0)')),
       ('C06_prec_upgrade','preceding_upgrade_flag = existe upgrade real nos 90 dias anteriores',lambda c_,r: OKF(f'{c_("preceding_upgrade_flag")}=({c_("aux_upgrade_90d")}>0)')),
       ('C07_churn_apos_cadastro','churn_date >= signup_date da conta',lambda c_,r: OKF(f'{c_("churn_date")}>={c_("aux_signup_conta")}'))]
U_aux=[('aux_inicio_assin','start_date da assinatura',lambda c_,r: f'INDEX({R("subscriptions","start_date")},MATCH({c_("subscription_id")},{R("subscriptions","subscription_id")},0))'),
       ('aux_fim_assin','end_date da assinatura (31/12/2099 se ativa)',lambda c_,r: f'INDEX({R("subscriptions","aux_fim_calc")},MATCH({c_("subscription_id")},{R("subscriptions","subscription_id")},0))'),
       ('aux_conta','account_id da assinatura',lambda c_,r: f'INDEX({R("subscriptions","account_id")},MATCH({c_("subscription_id")},{R("subscriptions","subscription_id")},0))'),
       ('aux_signup_conta','signup_date da conta',lambda c_,r: f'INDEX({R("accounts","signup_date")},MATCH({c_("aux_conta")},{R("accounts","account_id")},0))')]
U_tests=[('U01_uso_apos_inicio','usage_date >= start_date da assinatura',lambda c_,r: OKF(f'{c_("usage_date")}>={c_("aux_inicio_assin")}')),
       ('U02_uso_antes_fim','usage_date <= end_date da assinatura (se encerrada)',lambda c_,r: OKF(f'{c_("usage_date")}<={c_("aux_fim_assin")}')),
       ('U03_uso_apos_cadastro','usage_date >= signup_date da conta',lambda c_,r: OKF(f'{c_("usage_date")}>={c_("aux_signup_conta")}')),
       ('U04_count0_dur0','Se usage_count = 0, usage_duration_secs = 0',lambda c_,r: OKF(f'OR({c_("usage_count")}>0,{c_("usage_duration_secs")}=0)'))]
T_aux=[('aux_signup_conta','signup_date da conta',lambda c_,r: f'INDEX({R("accounts","signup_date")},MATCH({c_("account_id")},{R("accounts","account_id")},0))')]
T_tests=[('T01_fecha_apos_abrir','closed_at >= submitted_at',lambda c_,r: OKF(f'{c_("closed_at")}>={c_("submitted_at")}')),
       ('T02_resolucao_calc','resolution_time_hours = (closed_at − submitted_at) em horas',lambda c_,r: OKF(f'ABS(({c_("closed_at")}-{c_("submitted_at")})*24-{c_("resolution_time_hours")})<0.01')),
       ('T03_ticket_apos_cadastro','submitted_at >= signup_date da conta',lambda c_,r: OKF(f'{c_("submitted_at")}>={c_("aux_signup_conta")}')),
       ('T04_satisfacao_1a5','satisfaction_score vazio ou entre 1 e 5',lambda c_,r: OKF(f'OR({c_("satisfaction_score")}="",AND({c_("satisfaction_score")}>=1,{c_("satisfaction_score")}<=5))'))]
c['feedback_text']=c.feedback_text.where(c.feedback_text.notna(),None)
spec=[('subscriptions',s,S_aux,S_tests,{'start_date':D,'end_date':D,'aux_fim_calc':D,'aux_signup_conta':D}),
      ('accounts',a,A_aux,A_tests,{'signup_date':D}),
      ('churn_events',c,C_aux,C_tests,{'churn_date':D,'aux_signup_conta':D}),
      ('feature_usage',f,U_aux,U_tests,{'usage_date':D,'aux_inicio_assin':D,'aux_fim_assin':D,'aux_signup_conta':D}),
      ('support_tickets',t,T_aux,T_tests,{'submitted_at':DT,'closed_at':DT,'aux_signup_conta':D})]
wss={n:data_sheet(n,df,ax,ts,dfm) for n,df,ax,ts,dfm in spec}
for n,df,ax,ts,dfm in spec: fill_formulas(wss[n],n,df,ax,ts,dfm)

# ---------- Scorecard ----------
ws=sc
ws['A1']='Scorecard de consistência: o que cada tabela promete vs o que entrega'; ws['A1'].font=F(bold=True,size=15)
ws['A2']='Cada linha é uma regra testada linha a linha na aba da tabela (colunas laranja). Contagens = fórmulas sobre essas colunas.'; ws['A2'].font=F(italic=True,size=9,color=GRAY)
hdr=['Teste','Tabela','Regra testada','Tipo','Registros testados','OK','FALHA','% OK','Referência: % esperado se fosse ao acaso','Veredito']
for j,h in enumerate(hdr,1):
    x=ws.cell(4,j,h); x.font=F(bold=True,color='FFFFFF',size=9); x.fill=H_ORIG; x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); x.border=BOX
ws.row_dimensions[4].height=40
TIPO={'S':'Interna','A':'Cruzada','C':'Cruzada','U':'Cruzada','T':'Interna'}
CROSS={'S06','S07','A01','A02','A03','A04','C01','C02','C03','C05','C06','C07','U01','U02','U03','T03'}
ACASO={'A01':'=ROUND(Cruzamentos!$B$22,3)','A03':'≈ 33% (3 planos)','C03':'=ROUND(COUNTIF('+R('accounts','churn_flag')+',TRUE)/COUNTA('+R('accounts','account_id')+'),3)'}
r=5
for n,df,ax,ts,dfm in spec:
    for col,desc,fn in ts:
        tid=col[:3]; rng=R(n,col)
        vals=[tid,n,desc,'Cruzada (com outra tabela)' if tid in CROSS else 'Interna (a própria tabela)',
              f'=COUNTIF({rng},"OK")+COUNTIF({rng},"FALHA")',f'=COUNTIF({rng},"OK")',f'=COUNTIF({rng},"FALHA")',f'=IF(E{r}>0,F{r}/E{r},"")',
              ACASO.get(tid,''),f'=IF(E{r}=0,"",IF(H{r}>=0.99,"Consistente",IF(H{r}>=0.9,"Pequenas exceções","Inconsistente")))']
        for j,v in enumerate(vals,1):
            x=ws.cell(r,j,v); x.font=F(size=9,bold=(j==1)); x.border=BOX; x.alignment=Alignment(wrap_text=(j==3),vertical='center')
            if j in (5,6,7): x.number_format='#,##0'
            if j in (8,9): x.number_format='0.0%'
        r+=1
last=r-1
ws.conditional_formatting.add(f'J5:J{last}',CellIsRule(operator='equal',formula=['"Inconsistente"'],fill=RED))
ws.conditional_formatting.add(f'J5:J{last}',CellIsRule(operator='equal',formula=['"Consistente"'],fill=GREEN))
ws.conditional_formatting.add(f'J5:J{last}',CellIsRule(operator='equal',formula=['"Pequenas exceções"'],fill=PatternFill('solid',fgColor='FFF3CD')))
r+=1; ws.cell(r,1,'Resumo por tabela').font=F(bold=True,size=12); r+=1
for j,h in enumerate(['Tabela','Testes','Testes "Consistente"','Testes "Inconsistente"','% médio OK'],1):
    x=ws.cell(r,j,h); x.font=F(bold=True,color='FFFFFF',size=9); x.fill=H_ORIG; x.border=BOX
r+=1
for n,*_ in spec:
    ws.cell(r,1,n).font=F(bold=True,size=9)
    ws.cell(r,2,f'=COUNTIF($B$5:$B${last},A{r})'); ws.cell(r,3,f'=COUNTIFS($B$5:$B${last},A{r},$J$5:$J${last},"Consistente")')
    ws.cell(r,4,f'=COUNTIFS($B$5:$B${last},A{r},$J$5:$J${last},"Inconsistente")'); ws.cell(r,5,f'=AVERAGEIF($B$5:$B${last},A{r},$H$5:$H${last})').number_format='0.0%'
    for j in range(1,6): ws.cell(r,j).border=BOX; ws.cell(r,j).font=F(size=9,bold=(j==1))
    r+=1
r+=1
notes=['Veredito: ≥ 99% OK = Consistente; 90–99% = Pequenas exceções; < 90% = Inconsistente.',
 'Referência ao acaso: taxa de acerto que se teria se as duas informações fossem sorteadas de forma independente. % OK perto dela = as informações não têm relação.',
 'Testes "Interna" usam só a própria tabela. Testes "Cruzada" comparam com outra tabela (ex.: data de uso vs início da assinatura).',
 'Passar num teste prova coerência, não verdade: a tabela de assinaturas é coerente consigo mesma e com o cadastro, mas isso não garante que reflita a realidade.']
for tx in notes: ws.cell(r,1,tx).font=F(italic=True,size=9,color=GRAY); r+=1
for col,w in zip('ABCDEFGHIJ',[14,16,58,22,11,9,9,9,16,18]): ws.column_dimensions[col].width=w
ws.freeze_panes='A5'

# ---------- Cruzamentos ----------
ws=cz
ws['A1']='Cruzamentos: as marcações de churn têm relação entre si?'; ws['A1'].font=F(bold=True,size=15)
ws['A2']='Tabelas 2×2 por conta (500 contas) e teste qui-quadrado (CHITEST). p-valor alto (> 0,05) = não há evidência de relação: as marcações se comportam como independentes.'; ws['A2'].font=F(italic=True,size=9,color=GRAY)
def cross(r0,title,colx,lx,coly,ly):
    ws.cell(r0,1,title).font=F(bold=True,size=11)
    ws.cell(r0+1,1,f'{lx} \\ {ly}').font=F(bold=True,size=9)
    for j,v in enumerate(['Sim','Não','Total'],2): ws.cell(r0+1,j,v).font=F(bold=True,size=9)
    ws.cell(r0+1,6,'Esperado se independente').font=F(italic=True,size=9,color=GRAY)
    for i,u in enumerate(['Sim','Não'],r0+2):
        ws.cell(i,1,u).font=F(bold=True,size=9)
        for j,v in enumerate(['Sim','Não'],2):
            ws.cell(i,j,f'=COUNTIFS({R("accounts",colx)},"{u}",{R("accounts",coly)},"{v}")')
            ws.cell(i,j+4,f'={L(4)}{i}*{L(j)}{r0+4}/$D${r0+4}')
        ws.cell(i,4,f'=B{i}+C{i}')
    ws.cell(r0+4,1,'Total').font=F(bold=True,size=9)
    for j in (2,3,4): ws.cell(r0+4,j,f'={L(j)}{r0+2}+{L(j)}{r0+3}')
    ws.cell(r0+5,1,f'% {ly}=Sim quando {lx}=Sim'); ws.cell(r0+5,2,f'=B{r0+2}/D{r0+2}').number_format='0.0%'
    ws.cell(r0+6,1,f'% {ly}=Sim quando {lx}=Não'); ws.cell(r0+6,2,f'=B{r0+3}/D{r0+3}').number_format='0.0%'
    ws.cell(r0+7,1,'p-valor (qui-quadrado)'); ws.cell(r0+7,2,f'=CHITEST(B{r0+2}:C{r0+3},F{r0+2}:G{r0+3})').number_format='0.00'
    ws.cell(r0+7,3,f'=IF(B{r0+7}>0.05,"Sem relação detectável","Há relação")')
    for rr in range(r0+5,r0+8): ws.cell(rr,1).font=F(size=9)
cross(4,'1. churn_flag (accounts) × tem evento em churn_events','aux_flag_txt','churn_flag','aux_tem_evento','tem evento')
cross(13,'2. churn_flag (accounts) × tem assinatura encerrada','aux_flag_txt','churn_flag','aux_tem_assin_encerrada','tem assin. encerrada')
# linha 22: acerto esperado ao acaso para A01
ws['A22']='Acerto esperado ao acaso no teste A01 (flag = tem evento)'; ws['A22'].font=F(size=9)
ws['B22']='=(D6/D8)*(B8/D8)+(D7/D8)*(C8/D8)'; ws['B22'].number_format='0.0%'
cross(24,'3. tem evento em churn_events × tem assinatura encerrada','aux_tem_evento','tem evento','aux_tem_assin_encerrada','tem assin. encerrada')
r0=34; ws.cell(r0,1,'4. Motivo registrado (reason_code) × comentário do cliente (feedback_text), por evento').font=F(bold=True,size=11)
reasons=sorted(c.reason_code.unique()); fbs=['(vazio)']+sorted(c.feedback_text.dropna().unique())
for j,v in enumerate(fbs+['Total'],2): ws.cell(r0+1,j,v).font=F(bold=True,size=9)
for i,rc in enumerate(reasons,r0+2):
    ws.cell(i,1,rc).font=F(bold=True,size=9)
    for j,fb in enumerate(fbs,2): ws.cell(i,j,f'=COUNTIFS({R("churn_events","reason_code")},$A{i},{R("churn_events","aux_feedback")},{L(j)}${r0+1})')
    ws.cell(i,len(fbs)+2,f'=SUM(B{i}:{L(len(fbs)+1)}{i})')
rl=r0+1+len(reasons); rt=rl+1; ws.cell(rt,1,'Total').font=F(bold=True,size=9)
for j in range(2,len(fbs)+3): ws.cell(rt,j,f'=SUM({L(j)}{r0+2}:{L(j)}{rl})')
ec=len(fbs)+4; ws.cell(r0+1,ec,'Esperado se independente →').font=F(italic=True,size=9,color=GRAY)
for i in range(r0+2,rl+1):
    for j in range(len(fbs)):
        ws.cell(i,ec+1+j,f'=${L(len(fbs)+2)}{i}*{L(2+j)}${rt}/${L(len(fbs)+2)}${rt}').number_format='0.0'
ws.cell(rt+1,1,'p-valor (qui-quadrado)').font=F(size=9)
ws.cell(rt+1,2,f'=CHITEST(B{r0+2}:{L(len(fbs)+1)}{rl},{L(ec+1)}{r0+2}:{L(ec+len(fbs))}{rl})').number_format='0.00'
ws.cell(rt+1,3,f'=IF(B{rt+1}>0.05,"Sem relação detectável","Há relação")')
ws.cell(rt+2,1,'Ex.: eventos com motivo "competitor" têm o comentário "too expensive" tanto quanto "switched to competitor".').font=F(italic=True,size=9,color=GRAY)
ws.column_dimensions['A'].width=44
for j in range(2,16): ws.column_dimensions[L(j)].width=14

# ---------- Exemplos (cópias literais de linhas dos CSVs) ----------
ws=ex
ws['A1']='Exemplos concretos (linhas copiadas dos CSVs originais, sem alteração)'; ws['A1'].font=F(bold=True,size=15)
r=3
def block(title, why, df, datecols=()):
    global r
    ws.cell(r,1,title).font=F(bold=True,size=11); r+=1
    ws.cell(r,1,why).font=F(italic=True,size=9,color=GRAY); r+=1
    for j,h in enumerate(df.columns,1):
        x=ws.cell(r,j,h); x.font=F(bold=True,color='FFFFFF',size=9); x.fill=H_ORIG
    r+=1
    for row in df.itertuples(index=False):
        for j,v in enumerate(row,1):
            x=ws.cell(r,j,py(v)); x.font=F(size=9)
            if df.columns[j-1] in datecols: x.number_format=D
        r+=1
    r+=2
ex1='A-0a282f'
block(f'Exemplo 1: conta {ex1} — churn_flag = False, mas tem 2 eventos de churn','accounts:',a[a.account_id==ex1],['signup_date'])
block('','churn_events: eventos em 06/10/2024 (features) e 31/12/2024 (support)',c[c.account_id==ex1][['churn_event_id','account_id','churn_date','reason_code','is_reactivation','feedback_text']],['churn_date'])
block('','subscriptions: no dia do 1º churn (06/10) a conta ABRE a S-f32440 e segue abrindo assinaturas; nenhuma termina perto das datas de churn',s[s.account_id==ex1].sort_values('start_date')[['subscription_id','start_date','end_date','plan_tier','seats','mrr_amount','churn_flag']],['start_date','end_date'])
ex2='A-038089'
block(f'Exemplo 2: conta {ex2} — churn_flag = True e 2 eventos de churn, mas 3 assinaturas ativas (uma aberta em 30/12/2024, depois do último churn)','accounts:',a[a.account_id==ex2],['signup_date'])
block('','churn_events:',c[c.account_id==ex2][['churn_event_id','account_id','churn_date','reason_code','is_reactivation','feedback_text']],['churn_date'])
block('','subscriptions:',s[s.account_id==ex2].sort_values('start_date')[['subscription_id','start_date','end_date','plan_tier','seats','mrr_amount','churn_flag']],['start_date','end_date'])
ex3='A-10b8da'
block(f'Exemplo 3: conta {ex3} — churn_flag = True, nenhum evento de churn, 5 assinaturas ativas','accounts (sem linhas em churn_events):',a[a.account_id==ex3],['signup_date'])
block('','subscriptions:',s[s.account_id==ex3].sort_values('start_date')[['subscription_id','start_date','end_date','plan_tier','seats','mrr_amount','churn_flag']],['start_date','end_date'])
ex4='S-0896f4'; acc4=s[s.subscription_id==ex4].account_id.iloc[0]
block(f'Exemplo 4: assinatura {ex4} (conta {acc4}, cadastrada em 04/09/2024, assinatura iniciada em 08/12/2024) tem uso registrado desde 08/02/2023','subscriptions:',s[s.subscription_id==ex4][['subscription_id','account_id','start_date','end_date','plan_tier','mrr_amount']],['start_date','end_date'])
block('','feature_usage:',f[f.subscription_id==ex4].sort_values('usage_date'),['usage_date'])
ex5=s[s.account_id==ex1].iloc[0].subscription_id
block('Exemplo 5 (o contraponto): a tabela de assinaturas fecha a conta','Cada assinatura: MRR = seats × preço do plano (Basic 19, Pro 49, Enterprise 199), ARR = 12 × MRR, churn_flag = end_date preenchido. Ex.: conta A-0a282f:',
      s[s.account_id==ex1].sort_values('start_date').head(6)[['subscription_id','plan_tier','seats','mrr_amount','arr_amount','is_trial','end_date','churn_flag']],['end_date'])
for j in range(1,15): ws.column_dimensions[L(j)].width=15

# ---------- LEIA-ME ----------
ws=leia; ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=120
L_=[('Passo 1 da reconciliação: qual tabela é confiável o bastante para ser a base?',dict(bold=True,size=15)),
('Planilha APARTADA. As abas de dados são cópias dos CSVs originais (G4_DESAFIO\\dados), que não foram alterados. Nada aqui altera os arquivos de DADOS_ORIGINAIS.',dict(italic=True,color=GRAY)),
('',{}),('COMO LER',dict(bold=True,size=12)),
('• Scorecard: um teste por linha, com quantos registros passam e o veredito. Comece por aqui.',{}),
('• Abas subscriptions, accounts, churn_events, feature_usage, support_tickets: dados copiados (cabeçalho escuro) + colunas auxiliares (cinza) + colunas de teste (laranja, OK/FALHA). Passe o mouse no cabeçalho para ver a regra. Use o filtro da coluna para ver só as FALHAS.',{}),
('• Cruzamentos: tabelas 2×2 e qui-quadrado mostrando se as marcações de churn têm relação entre si.',{}),
('• Exemplos: casos concretos, com as linhas copiadas como estão nos CSVs.',{}),
('',{}),('DECISÃO PROPOSTA',dict(bold=True,size=12)),
('• A tabela subscriptions passa em todos os testes internos e no cruzamento com o cadastro (start_date ≥ signup_date). Ela vira a BASE das análises: o desfecho principal passa a ser o churn de receita (assinaturas encerradas, com MRR).',{}),
('• churn_events e accounts.churn_flag viram desfechos SECUNDÁRIOS: só servem para confirmar (triangular) o que aparecer na base principal, nunca como prova sozinhos.',{}),
('• Campos reprovados não sustentam conclusão causal: datas de uso (feature_usage), datas de ticket vs cadastro, reason_code vs feedback, preceding_*_flag, is_reactivation, plan_tier/seats do cadastro.',{}),
('',{}),('LIMITES',dict(bold=True,size=12)),
('• Passar nos testes prova coerência, não verdade. A base é sintética (README dos dados).',{}),
('• 25 assinaturas (0,5%) têm upgrade_flag e downgrade_flag ao mesmo tempo (S09): exceção pequena, registrada.',{}),
('• O teste A03/A04 usa a assinatura mais antiga da conta; empates de data contam a favor (critério brando).',{}),
]
for i,(tx,st) in enumerate(L_,1):
    x=ws.cell(i,2,tx); x.font=F(**st) if st else F(); x.alignment=Alignment(wrap_text=True,vertical='top')
wb.calculation.fullCalcOnLoad=True
wb.active=1
OUT.parent.mkdir(parents=True,exist_ok=True); wb.save(OUT); print('ok',OUT)
