# Passo 2: triangulação (gerado com Claude/IA). monta bases por conta e por assinatura, taxas por exposição e testes de permutação.
import sys
from pathlib import Path
import pandas as pd, numpy as np
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
FIM = pd.Timestamp('2024-12-31'); MES = 30.4375
N_PERM = 2000; SEED = 42

def carregar():
    a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
    s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date','end_date'])
    f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv')
    t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv')
    c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])
    return a, s, f, t, c

def quartil(x, labels):
    return pd.qcut(x.rank(method='first'), 4, labels=labels).astype(str)

def bases():
    a, s, f, t, c = carregar()
    # ---- assinaturas (unidade do desfecho principal) ----
    S = s.copy()
    S['encerrada'] = S.end_date.notna().astype(int)
    S['exposicao_meses'] = ((S.end_date.fillna(FIM).clip(upper=FIM) - S.start_date).dt.days.clip(lower=1)) / MES
    S['mrr_x_meses'] = S.mrr_amount * S.exposicao_meses
    S['mrr_encerrado'] = S.mrr_amount * S.encerrada
    S['f_plano_assin'] = S.plan_tier
    S['f_cobranca'] = S.billing_frequency
    S['f_seats_assin'] = pd.cut(S.seats, [0, 5, 15, 40, 10**6], labels=['1–5', '6–15', '16–40', '41+']).astype(str)
    S['f_trial'] = np.where(S.is_trial, 'Trial', 'Pago')
    S['f_auto_renew'] = np.where(S.auto_renew_flag, 'Sim', 'Não')
    S['f_upgrade_assin'] = np.where(S.upgrade_flag, 'Sim', 'Não')
    S['f_downgrade_assin'] = np.where(S.downgrade_flag, 'Sim', 'Não')
    # ---- contas ----
    K = a[['account_id', 'industry', 'country', 'referral_source', 'signup_date', 'churn_flag']].copy()
    K['exposicao_meses_conta'] = ((FIM - K.signup_date).dt.days.clip(lower=1)) / MES
    K['eventos_churn'] = K.account_id.map(c.account_id.value_counts()).fillna(0).astype(int)
    K['flag_churn'] = K.churn_flag.astype(int)
    g = S.groupby('account_id')
    mrr_plano = S.pivot_table(index='account_id', columns='plan_tier', values='mrr_amount', aggfunc='sum').fillna(0)
    K['f_plano_predominante'] = K.account_id.map(mrr_plano.idxmax(axis=1))
    mrr_cob = S.pivot_table(index='account_id', columns='billing_frequency', values='mrr_amount', aggfunc='sum').fillna(0)
    K['f_cobranca_predominante'] = K.account_id.map(mrr_cob.idxmax(axis=1))
    K['seats_total'] = K.account_id.map(g.seats.sum())
    K['f_seats_conta'] = quartil(K.seats_total, ['Q1 (menos seats)', 'Q2', 'Q3', 'Q4 (mais seats)'])
    K['f_teve_upgrade'] = np.where(K.account_id.map(g.upgrade_flag.any()), 'Sim', 'Não')
    K['f_teve_downgrade'] = np.where(K.account_id.map(g.downgrade_flag.any()), 'Sim', 'Não')
    fu = f.merge(s[['subscription_id', 'account_id']], on='subscription_id')
    gu = fu.groupby('account_id')
    K['uso_total'] = K.account_id.map(gu.usage_count.sum()).fillna(0)
    K['erros_100'] = (K.account_id.map(gu.error_count.sum()).fillna(0) / K.uso_total.replace(0, np.nan) * 100).fillna(0)
    K['f_uso_conta'] = quartil(K.uso_total, ['Q1 (menos uso)', 'Q2', 'Q3', 'Q4 (mais uso)'])
    K['f_erros_conta'] = quartil(K.erros_100, ['Q1 (menos erros)', 'Q2', 'Q3', 'Q4 (mais erros)'])
    K['f_beta'] = np.where(K.account_id.map(fu[fu.is_beta_feature].groupby('account_id').size()).fillna(0) > 0, 'Usou beta', 'Não usou beta')
    gt = t.groupby('account_id')
    K['tickets'] = K.account_id.map(gt.size()).fillna(0)
    K['f_tickets_conta'] = pd.cut(K.tickets, [-1, 2, 4, 6, 10**6], labels=['0–2', '3–4', '5–6', '7+']).astype(str)
    sat = K.account_id.map(gt.satisfaction_score.mean())
    K['f_satisfacao'] = np.select([sat.isna(), sat < 3, sat < 4], ['Sem resposta', '< 3', '3 a 3,99'], '≥ 4')
    K['f_escalacao'] = np.where(K.account_id.map(gt.escalation_flag.any()).fillna(False), 'Sim', 'Não')
    urg = K.account_id.map(t[t.priority == 'urgent'].groupby('account_id').size()).fillna(0)
    K['f_ticket_urgente'] = np.where(urg > 0, 'Sim', 'Não')
    K = K.rename(columns={'industry': 'f_industria', 'country': 'f_pais', 'referral_source': 'f_canal'})
    # atributos de conta replicados nas assinaturas (permutação é feita por conta)
    acc_f = [c_ for c_ in K.columns if c_.startswith('f_')]
    S = S.merge(K[['account_id'] + acc_f], on='account_id')
    return K, S, acc_f

FATORES = [  # (coluna, rótulo, nível: 'conta' ou 'assinatura')
 ('f_industria', 'Indústria', 'conta'), ('f_pais', 'País', 'conta'), ('f_canal', 'Canal de aquisição', 'conta'),
 ('f_plano_predominante', 'Plano predominante da conta (maior MRR)', 'conta'), ('f_cobranca_predominante', 'Cobrança predominante da conta', 'conta'),
 ('f_seats_conta', 'Seats totais da conta (quartis)', 'conta'), ('f_teve_upgrade', 'Conta teve upgrade', 'conta'), ('f_teve_downgrade', 'Conta teve downgrade', 'conta'),
 ('f_uso_conta', 'Uso total da conta (quartis)', 'conta'), ('f_erros_conta', 'Erros por 100 usos (quartis)', 'conta'), ('f_beta', 'Usou feature beta', 'conta'),
 ('f_tickets_conta', 'Tickets por conta', 'conta'), ('f_satisfacao', 'Satisfação média nos tickets', 'conta'), ('f_escalacao', 'Teve ticket escalado', 'conta'), ('f_ticket_urgente', 'Teve ticket urgente', 'conta'),
 ('f_plano_assin', 'Plano da assinatura', 'assinatura'), ('f_cobranca', 'Cobrança da assinatura', 'assinatura'), ('f_seats_assin', 'Seats da assinatura', 'assinatura'),
 ('f_trial', 'Assinatura trial', 'assinatura'), ('f_auto_renew', 'Renovação automática', 'assinatura'), ('f_upgrade_assin', 'Assinatura com upgrade', 'assinatura'), ('f_downgrade_assin', 'Assinatura com downgrade', 'assinatura'),
]
METRICAS = {  # nome: (base, numerador, denominador, escala)
 'P1_cancel': ('S', 'encerrada', 'exposicao_meses', 100),      # assinaturas encerradas por 100 assinaturas-mês
 'P2_mrr': ('S', 'mrr_encerrado', 'mrr_x_meses', 100),         # % do MRR cancelado por mês de exposição
 'S1_eventos': ('K', 'eventos_churn', 'exposicao_meses_conta', 100),  # eventos por 100 contas-mês
 'S2_flag': ('K', 'flag_churn', None, 100),                     # % contas com churn_flag
}

def rates(codes, num, den, nlev):
    n = np.bincount(codes, weights=num, minlength=nlev)
    d = np.bincount(codes, weights=den, minlength=nlev)
    with np.errstate(divide='ignore', invalid='ignore'): return n / d

def bh(p):
    p = np.asarray(p, float); ok = ~np.isnan(p); q = np.full_like(p, np.nan)
    pv = p[ok]; o = np.argsort(pv); m = len(pv)
    qq = pv[o] * m / np.arange(1, m + 1); qq = np.minimum.accumulate(qq[::-1])[::-1]
    tmp = np.empty(m); tmp[o] = np.minimum(qq, 1); q[ok] = tmp; return q

def analisar():
    K, S, acc_f = bases()
    rng = np.random.default_rng(SEED)
    acc_idx = {a: i for i, a in enumerate(K.account_id)}
    S_acc = S.account_id.map(acc_idx).values
    linhas = []; globais = []
    for col, lab, nivel in FATORES:
        niveis = sorted(pd.concat([S[col]]).unique()) if nivel == 'assinatura' else sorted(K[col].unique())
        lev = {v: i for i, v in enumerate(niveis)}; L_ = len(niveis)
        res = {}
        for m, (base, num, den, esc) in METRICAS.items():
            if base == 'K' and nivel == 'assinatura': continue
            if base == 'S':
                y = S[num].values.astype(float); d = S[den].values.astype(float)
                if nivel == 'conta':
                    acc_codes = K[col].map(lev).values
                    def codes_of(perm): return acc_codes[perm][S_acc]
                else:
                    sub_codes = S[col].map(lev).values
                    def codes_of(perm): return sub_codes[perm]
                N = len(K) if nivel == 'conta' else len(S)
            else:
                y = K[num].values.astype(float); d = (K[den].values.astype(float) if den else np.ones(len(K)))
                acc_codes = K[col].map(lev).values
                def codes_of(perm): return acc_codes[perm]
                N = len(K)
            ident = np.arange(N)
            obs = rates(codes_of(ident), y, d, L_); geral = y.sum() / d.sum()
            w = np.bincount(codes_of(ident), weights=d, minlength=L_)
            stat_obs = np.nansum(w * (obs - geral) ** 2)
            ext = np.zeros(L_); ext_g = 0
            for _ in range(N_PERM):
                pr = rates(codes_of(rng.permutation(N)), y, d, L_)
                ext += np.abs(pr - geral) >= np.abs(obs - geral) - 1e-12
            # estatística global (todas as categorias de uma vez), com permutações próprias
            cnt = 0
            for _ in range(N_PERM):
                cp = codes_of(rng.permutation(N)); pr = rates(cp, y, d, L_); wp = np.bincount(cp, weights=d, minlength=L_)
                cnt += np.nansum(wp * (pr - geral) ** 2) >= stat_obs - 1e-12
            pl = (ext + 1) / (N_PERM + 1); pl[np.isnan(obs)] = np.nan
            res[m] = dict(taxa=obs * esc, geral=geral * esc, p=pl, p_global=(cnt + 1) / (N_PERM + 1))
        n_assin = S.groupby(col).size(); n_contas = (K.groupby(col).size() if nivel == 'conta' else S.groupby(col).account_id.nunique())
        for v, i in lev.items():
            row = dict(fator=lab, coluna=col, nivel_fator=nivel, valor=v, n_assinaturas=int(n_assin.get(v, 0)), n_contas=int(n_contas.get(v, 0)))
            for m in METRICAS:
                if m in res: row[f'{m}_taxa'] = res[m]['taxa'][i]; row[f'{m}_geral'] = res[m]['geral']; row[f'{m}_p'] = res[m]['p'][i]
                else: row[f'{m}_taxa'] = row[f'{m}_geral'] = row[f'{m}_p'] = np.nan
            linhas.append(row)
        globais.append(dict(fator=lab, coluna=col, nivel_fator=nivel, n_niveis=L_, **{f'{m}_p_global': res[m]['p_global'] if m in res else np.nan for m in METRICAS}))
    M = pd.DataFrame(linhas); G = pd.DataFrame(globais)
    for m in METRICAS:
        M[f'{m}_q'] = bh(M[f'{m}_p'].values); G[f'{m}_q_global'] = bh(G[f'{m}_p_global'].values)
    return K, S, M, G

def _debug():
    K, S, M, G = analisar()
    pd.set_option('display.width', 250); pd.set_option('display.max_rows', 200)
    show = M[['fator', 'valor', 'n_contas', 'P1_cancel_taxa', 'P1_cancel_p', 'P1_cancel_q', 'P2_mrr_taxa', 'P2_mrr_q', 'S1_eventos_taxa', 'S1_eventos_q', 'S2_flag_taxa', 'S2_flag_q']]
    print(show.round(3).to_string()); print(G.round(3).to_string())
    print(M[['P1_cancel_geral', 'P2_mrr_geral', 'S1_eventos_geral', 'S2_flag_geral']].iloc[0])

# ======================= Fatores de TEMPO (variam ao longo da vida) =======================
from math import erf, sqrt, log
BANDAS = [0, 30, 90, 180, 365, 10**4]; LB = ['0–30 dias', '31–90 dias', '91–180 dias', '181–365 dias', '> 365 dias']

def p_poisson_razao(e1, t1, e0, t0):
    """p bilateral (aprox. normal no log da razão de taxas) comparando grupo 1 vs grupo 0."""
    if min(e1, e0) == 0: return np.nan
    z = abs(log((e1 / t1) / (e0 / t0))) / sqrt(1 / e1 + 1 / e0)
    return 2 * (1 - 0.5 * (1 + erf(z / sqrt(2))))

def p_poisson_obs_esp(o, lam):
    """p bilateral exato: o observado vem de uma Poisson com média lam (esperado)?"""
    from math import lgamma, exp, log as ln
    pmf = lambda k: exp(k * ln(lam) - lam - lgamma(k + 1))
    po = pmf(o); tot = 0.0
    for k in range(0, int(max(o, lam) * 3 + 50)):
        pk = pmf(k)
        if pk <= po * (1 + 1e-9): tot += pk
    return min(1.0, tot)

def tempo():
    a, s, f, t, c = carregar()
    s = s.copy(); s['stop'] = s.end_date.fillna(FIM); s['dur'] = (s.stop - s.start_date).dt.days
    # pessoa-tempo por mês-calendário x faixa de idade da assinatura
    recs = []
    for r in s.itertuples():
        for lo, hi, lb in zip(BANDAS[:-1], BANDAS[1:], LB):
            a0 = r.start_date + pd.Timedelta(days=lo); a1 = min(r.start_date + pd.Timedelta(days=hi), r.stop)
            if a1 <= a0: continue
            for m in pd.date_range(a0.to_period('M').to_timestamp(), a1, freq='MS'):
                me = m + pd.offsets.MonthEnd(0) + pd.Timedelta(days=1)
                x0, x1 = max(a0, m), min(a1, me)
                if x1 > x0: recs.append((m, lb, (x1 - x0).days / MES))
    pt = pd.DataFrame(recs, columns=['mes', 'faixa', 'assin_meses'])
    e = s[s.end_date.notna()].copy()
    e['faixa'] = pd.cut(e.dur, [-1, 30, 90, 180, 365, 10**5], labels=LB).astype(str)
    e['mes'] = e.end_date.dt.to_period('M').dt.to_timestamp()
    # (1) por idade da assinatura
    T1 = pd.DataFrame({'faixa': LB, 'encerradas': [int((e.faixa == lb).sum()) for lb in LB],
                       'assin_meses': [pt[pt.faixa == lb].assin_meses.sum() for lb in LB]})
    # eventos de churn por idade da conta (definição secundária S1)
    cc = c.merge(a[['account_id', 'signup_date']], on='account_id'); cc['ten'] = (cc.churn_date - cc.signup_date).dt.days
    cc['faixa'] = pd.cut(cc.ten, [-1, 30, 90, 180, 365, 10**5], labels=LB).astype(str)
    dura = (FIM - a.signup_date).dt.days
    T1['eventos'] = [int((cc.faixa == lb).sum()) for lb in LB]
    T1['contas_meses'] = [((dura.clip(upper=hi) - lo).clip(lower=0)).sum() / MES for lo, hi, lb in zip(BANDAS[:-1], BANDAS[1:], LB)]
    for col_e, col_t, nome in [('encerradas', 'assin_meses', 'p_P1'), ('eventos', 'contas_meses', 'p_S1')]:
        T1[nome] = [p_poisson_razao(r[col_e], r[col_t], T1[col_e].sum() - r[col_e], T1[col_t].sum() - r[col_t]) for _, r in T1.iterrows()]
    # (2) por mês-calendário: observado vs esperado pela idade das assinaturas
    haz = (T1.set_index('faixa').encerradas / T1.set_index('faixa').assin_meses)
    pt['esperado'] = pt.faixa.map(haz) * pt.assin_meses
    meses = pd.date_range('2023-06-01', '2024-12-01', freq='MS')
    T2 = pd.DataFrame({'mes': meses})
    T2['encerradas'] = T2.mes.map(e.groupby('mes').size()).fillna(0).astype(int)
    T2['assin_meses'] = T2.mes.map(pt.groupby('mes').assin_meses.sum())
    T2['esperado_pela_idade'] = T2.mes.map(pt.groupby('mes').esperado.sum())
    T2['p_obs_vs_esp'] = [p_poisson_obs_esp(int(o), x) for o, x in zip(T2.encerradas, T2.esperado_pela_idade)]
    cm = cc.churn_date.dt.to_period('M').dt.to_timestamp()
    T2['eventos'] = T2.mes.map(cm.value_counts()).fillna(0).astype(int)
    def contas_meses(m):
        fim_m = min(m + pd.offsets.MonthEnd(0) + pd.Timedelta(days=1), FIM + pd.Timedelta(days=1))
        return (fim_m - a.signup_date.clip(lower=m)).dt.days.clip(lower=0).sum() / MES
    T2['contas_meses'] = [contas_meses(m) for m in T2.mes]
    # (3) coorte de início: % de assinaturas encerradas em até 30 dias (só quem teve 30 dias para observar)
    ok = s[s.start_date <= FIM - pd.Timedelta(days=30)].copy(); ok['coorte'] = ok.start_date.dt.to_period('Q').astype(str)
    ok['enc30'] = (ok.end_date.notna() & (ok.dur <= 30)).astype(int)
    T3 = ok.groupby('coorte').agg(assinaturas=('enc30', 'size'), encerradas_30d=('enc30', 'sum')).reset_index()
    return T1, T2, T3

# ======================= Planilha =======================
def planilha(OUT):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter as L
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.chart import BarChart, LineChart, Reference
    K, S, M, G = analisar(); T1, T2, T3 = tempo()
    FONT = 'Arial'; GRAY = '52514E'; F = lambda **k: Font(name=FONT, **k)
    HD = PatternFill('solid', fgColor='1F2937'); HP = PatternFill('solid', fgColor='1D4ED8'); HS = PatternFill('solid', fgColor='B45309'); HV = PatternFill('solid', fgColor='374151')
    GREEN = PatternFill('solid', fgColor='D1E7DD'); YEL = PatternFill('solid', fgColor='FFF3CD'); RED = PatternFill('solid', fgColor='F8D7DA'); GR = PatternFill('solid', fgColor='EDEDED')
    thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
    wb = Workbook(); wb.remove(wb.active)
    def S_(n): ws = wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
    leia = S_('LEIA-ME'); mz = S_('Matriz'); rf = S_('Resumo_Fatores'); tp = S_('Tempo')
    # ---- bases ----
    colsS = ['subscription_id', 'account_id', 'encerrada', 'exposicao_meses', 'mrr_amount', 'mrr_encerrado', 'mrr_x_meses'] + [c_ for c_, _, _ in FATORES]
    colsK = ['account_id', 'exposicao_meses_conta', 'eventos_churn', 'flag_churn'] + [c_ for c_, _, n in FATORES if n == 'conta']
    BS = {}
    for nome, df, cols in [('Base_Assinaturas', S, colsS), ('Base_Contas', K, colsK)]:
        ws = wb.create_sheet(nome); ws.append(cols)
        for row in df[cols].itertuples(index=False):
            ws.append([round(float(v), 6) if isinstance(v, (float, np.floating)) else (v.item() if hasattr(v, 'item') else v) for v in row])
        n = len(df) + 1; BS[nome] = (cols, n)
        for i, cn in enumerate(cols, 1):
            h = ws.cell(1, i); h.font = F(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(26, len(cn) + 2))
        tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
    def R(b, col): cols, n = BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
    # ---- Matriz ----
    ws = mz
    ws['A1'] = 'Matriz de triangulação: cada fator testado nas 3 definições de churn'; ws['A1'].font = F(bold=True, size=15)
    ws['A2'] = ('P1/P2 = desfecho PRINCIPAL (assinaturas, base coerente). S1/S2 = desfechos SECUNDÁRIOS (churn_events e churn_flag). '
                'Taxas e lift = fórmulas sobre as abas Base_*. p e q = teste de permutação (2.000 sorteios, calculado no script) e correção de Benjamini-Hochberg.')
    ws['A2'].font = F(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True); ws.merge_cells('A2:X2'); ws.row_dimensions[2].height = 28
    grp = [('', 1, 5, HD), ('P1 · Assinaturas encerradas por 100 assinaturas-mês', 6, 9, HP), ('P2 · % do MRR cancelado por mês', 10, 13, HP),
           ('S1 · Eventos de churn por 100 contas-mês', 14, 17, HS), ('S2 · % de contas com churn_flag', 18, 21, HS), ('', 22, 22, HV)]
    for t_, c0, c1, fl in grp:
        ws.cell(4, c0, t_)
        for j in range(c0, c1 + 1): x = ws.cell(4, j); x.fill = fl; x.font = F(bold=True, color='FFFFFF', size=9); x.alignment = Alignment(horizontal='center', wrap_text=True)
        if c1 > c0: ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c1)
    hdr = ['Fator', 'Nível', 'Valor', 'Contas', 'Assinaturas'] + ['Taxa', 'Lift vs geral', 'p', 'q'] * 4 + ['Veredito']
    for j, h in enumerate(hdr, 1):
        x = ws.cell(5, j, h); x.font = F(bold=True, color='FFFFFF', size=9); x.fill = grp[[i for i, g in enumerate(grp) if g[1] <= j <= g[2]][0]][3]; x.alignment = Alignment(horizontal='center', wrap_text=True); x.border = BOX
    ws.row_dimensions[4].height = 30; ws.row_dimensions[5].height = 28
    # linha de referência geral
    r = 6
    ws.cell(r, 1, 'GERAL (toda a base)').font = F(bold=True, size=9)
    ws.cell(r, 6, f'=SUM({R("Base_Assinaturas","encerrada")})/SUM({R("Base_Assinaturas","exposicao_meses")})*100')
    ws.cell(r, 10, f'=SUM({R("Base_Assinaturas","mrr_encerrado")})/SUM({R("Base_Assinaturas","mrr_x_meses")})*100')
    ws.cell(r, 14, f'=SUM({R("Base_Contas","eventos_churn")})/SUM({R("Base_Contas","exposicao_meses_conta")})*100')
    ws.cell(r, 18, f'=SUM({R("Base_Contas","flag_churn")})/COUNT({R("Base_Contas","flag_churn")})*100')
    for j in range(1, 23): ws.cell(r, j).fill = GR; ws.cell(r, j).border = BOX
    for j in (6, 10, 14, 18): ws.cell(r, j).number_format = '0.00'; ws.cell(r, j).font = F(bold=True, size=9)
    r = 7; first = r
    for _, m in M.iterrows():
        col = m.coluna; v = m.valor; crit = f'"{v}"'
        vals = [m.fator, 'Conta' if m.nivel_fator == 'conta' else 'Assinatura', v, int(m.n_contas), int(m.n_assinaturas)]
        for j, x in enumerate(vals, 1): ws.cell(r, j, x)
        ws.cell(r, 6, f'=IFERROR(SUMIFS({R("Base_Assinaturas","encerrada")},{R("Base_Assinaturas",col)},{crit})/SUMIFS({R("Base_Assinaturas","exposicao_meses")},{R("Base_Assinaturas",col)},{crit})*100,"")')
        ws.cell(r, 10, f'=IFERROR(SUMIFS({R("Base_Assinaturas","mrr_encerrado")},{R("Base_Assinaturas",col)},{crit})/SUMIFS({R("Base_Assinaturas","mrr_x_meses")},{R("Base_Assinaturas",col)},{crit})*100,"")')
        if m.nivel_fator == 'conta':
            ws.cell(r, 14, f'=IFERROR(SUMIFS({R("Base_Contas","eventos_churn")},{R("Base_Contas",col)},{crit})/SUMIFS({R("Base_Contas","exposicao_meses_conta")},{R("Base_Contas",col)},{crit})*100,"")')
            ws.cell(r, 18, f'=IFERROR(AVERAGEIFS({R("Base_Contas","flag_churn")},{R("Base_Contas",col)},{crit})*100,"")')
        for base_c, key in [(6, 'P1_cancel'), (10, 'P2_mrr'), (14, 'S1_eventos'), (18, 'S2_flag')]:
            ws.cell(r, base_c + 1, f'=IF(AND(ISNUMBER({L(base_c)}{r}),{L(base_c)}$6>0),{L(base_c)}{r}/{L(base_c)}$6,"")')
            for k, suf in [(2, '_p'), (3, '_q')]:
                val = m[key + suf]; ws.cell(r, base_c + k, None if pd.isna(val) else round(float(val), 4))
        # veredito (fórmula): significativo = q < 0,05; secundária só conta se tiver a MESMA direção da principal (lift dos dois lados de 1)
        sigP = f'OR(AND(ISNUMBER(I{r}),I{r}<0.05),AND(ISNUMBER(M{r}),M{r}<0.05))'
        dirP = f'IF(AND(ISNUMBER(I{r}),I{r}<0.05),G{r},K{r})'
        sS1 = f'IF(AND(ISNUMBER(Q{r}),Q{r}<0.05),(O{r}-1)*({dirP}-1)>0,FALSE)'; sS2 = f'IF(AND(ISNUMBER(U{r}),U{r}<0.05),(S{r}-1)*({dirP}-1)>0,FALSE)'
        anyS = f'OR(AND(ISNUMBER(Q{r}),Q{r}<0.05),AND(ISNUMBER(U{r}),U{r}<0.05))'
        anyp = f'OR(AND(ISNUMBER(H{r}),H{r}<0.05),AND(ISNUMBER(L{r}),L{r}<0.05),AND(ISNUMBER(P{r}),P{r}<0.05),AND(ISNUMBER(T{r}),T{r}<0.05))'
        ws.cell(r, 22, f'=IF(AND({sigP},{sS1},{sS2}),"Robusto (3 definições)",IF(AND({sigP},OR({sS1},{sS2})),"Confirmado (receita + 1)",IF({sigP},"Só receita",IF({anyS},"Só definição secundária",IF({anyp},"Indício fraco (some na correção)","Sem sinal")))))')
        for j in range(1, 23):
            x = ws.cell(r, j); x.border = BOX; x.font = F(size=9, bold=(j == 22))
            if j in (6, 10, 14, 18): x.number_format = '0.00'
            if j in (7, 11, 15, 19): x.number_format = '0.00"x"'
            if j in (8, 9, 12, 13, 16, 17, 20, 21): x.number_format = '0.000'
        r += 1
    last = r - 1
    rngV = f'V{first}:V{last}'
    ws.conditional_formatting.add(rngV, CellIsRule(operator='equal', formula=['"Sem sinal"'], fill=GR))
    ws.conditional_formatting.add(rngV, CellIsRule(operator='equal', formula=['"Indício fraco (some na correção)"'], fill=YEL))
    for t_ in ['Robusto (3 definições)', 'Confirmado (receita + 1)']: ws.conditional_formatting.add(rngV, CellIsRule(operator='equal', formula=[f'"{t_}"'], fill=GREEN))
    for t_ in ['Só receita', 'Só definição secundária']: ws.conditional_formatting.add(rngV, CellIsRule(operator='equal', formula=[f'"{t_}"'], fill=RED))
    for cq in 'HILMPQTU': ws.conditional_formatting.add(f'{cq}{first}:{cq}{last}', CellIsRule(operator='lessThan', formula=['0.05'], fill=YEL))
    r = last + 2
    ws.cell(r, 1, 'Contagem de vereditos').font = F(bold=True, size=11); r += 1
    for t_ in ['Robusto (3 definições)', 'Confirmado (receita + 1)', 'Só receita', 'Só definição secundária', 'Indício fraco (some na correção)', 'Sem sinal']:
        ws.cell(r, 1, t_).font = F(size=9); ws.cell(r, 2, f'=COUNTIF({rngV},A{r})').font = F(bold=True, size=9); r += 1
    r += 1
    notas = ['Lift = taxa do grupo ÷ taxa geral (1,00x = igual à média). p = chance de uma diferença desse tamanho aparecer só por acaso (permutação). q = p corrigido por testar muitos grupos ao mesmo tempo (Benjamini-Hochberg).',
             'Fatores de nível "Conta" são sorteados por conta (as assinaturas da conta vão juntas), para não inflar a significância. Fatores de "Assinatura" não têm S1/S2 (eventos e flag são por conta).',
             'Veredito usa q < 0,05. Secundária só confirma se apontar na mesma direção da principal. "Indício fraco" = p < 0,05 em alguma definição, mas não sobrevive à correção.',
             'P2 de trials fica vazio (MRR = 0). Exposição = meses em que a assinatura (ou conta) esteve ativa até 31/12/2024: corrige o fato de assinaturas recentes terem tido menos tempo para cancelar.']
    for tx in notas: ws.cell(r, 1, tx).font = F(italic=True, size=9, color=GRAY); r += 1
    for col_, w in zip('ABCDEFGHIJKLMNOPQRSTUV', [34, 10, 16, 8, 10] + [8, 9, 7, 7] * 4 + [30]): ws.column_dimensions[col_].width = w
    ws.freeze_panes = 'D6'
    # ---- Resumo_Fatores ----
    ws = rf
    ws['A1'] = 'Resumo por fator: o fator como um todo muda o churn?'; ws['A1'].font = F(bold=True, size=15)
    ws['A2'] = 'Teste global por fator (todas as categorias juntas), por permutação. q = corrigido entre os 22 fatores.'; ws['A2'].font = F(italic=True, size=9, color=GRAY)
    hdr = ['Fator', 'Nível', 'Categorias', 'P1 p', 'P1 q', 'P2 p', 'P2 q', 'S1 p', 'S1 q', 'S2 p', 'S2 q', 'Leitura']
    for j, h in enumerate(hdr, 1): x = ws.cell(4, j, h); x.font = F(bold=True, color='FFFFFF', size=9); x.fill = HD; x.border = BOX
    r = 5
    for _, g in G.iterrows():
        vals = [g.fator, 'Conta' if g.nivel_fator == 'conta' else 'Assinatura', int(g.n_niveis)]
        for m in METRICAS: vals += [None if pd.isna(g[f'{m}_p_global']) else round(float(g[f'{m}_p_global']), 4), None if pd.isna(g[f'{m}_q_global']) else round(float(g[f'{m}_q_global']), 4)]
        for j, x in enumerate(vals, 1): c_ = ws.cell(r, j, x); c_.font = F(size=9); c_.border = BOX; c_.number_format = '0.000'
        qs = ','.join(f'{c}{r}' for c in 'EGIK'); ps = ','.join(f'{c}{r}' for c in 'DFHJ')
        ws.cell(r, 12, f'=IF(COUNTIF(E{r}:K{r},"<0.05")-COUNTIF(D{r},"<0.05")-COUNTIF(F{r},"<0.05")-COUNTIF(H{r},"<0.05")-COUNTIF(J{r},"<0.05")>0,"Significativo após correção",IF(COUNTIF(D{r}:K{r},"<0.05")>0,"Indício fraco (some na correção)","Sem sinal"))').font = F(size=9, bold=True)
        ws.cell(r, 12).border = BOX; r += 1
    for cq in 'DEFGHIJK': ws.conditional_formatting.add(f'{cq}5:{cq}{r-1}', CellIsRule(operator='lessThan', formula=['0.05'], fill=YEL))
    ws.column_dimensions['A'].width = 40; ws.column_dimensions['L'].width = 32
    # ---- Tempo ----
    ws = tp
    ws['A1'] = 'Tempo: quando o churn acontece (idade da assinatura, mês e coorte)'; ws['A1'].font = F(bold=True, size=15)
    ws['A2'] = 'Contagens e exposição calculadas no script a partir dos CSVs; taxas e razões são fórmulas. p = teste de Poisson (razão de taxas).'; ws['A2'].font = F(italic=True, size=9, color=GRAY)
    def hd(row, labels, fills):
        for j, (h, fl) in enumerate(zip(labels, fills), 1): x = ws.cell(row, j, h); x.font = F(bold=True, color='FFFFFF', size=9); x.fill = fl; x.alignment = Alignment(wrap_text=True, horizontal='center'); x.border = BOX
        ws.row_dimensions[row].height = 42
    ws['A4'] = '1. Por idade (tempo de casa): risco nos primeiros dias'; ws['A4'].font = F(bold=True, size=12)
    hd(5, ['Idade', 'P1: assinaturas encerradas', 'P1: assinaturas-mês', 'P1: por 100 assinaturas-mês', 'P1: razão vs > 180 dias', 'P1: p (vs demais)',
           'S1: eventos de churn', 'S1: contas-mês', 'S1: por 100 contas-mês', 'S1: razão vs > 180 dias', 'S1: p (vs demais)'], [HD] + [HP] * 5 + [HS] * 5)
    for i, row in T1.iterrows():
        r = 6 + i
        vals = [row.faixa, int(row.encerradas), round(row.assin_meses, 1), f'=B{r}/C{r}*100', f'=D{r}/(SUM(B$9:B$10)/SUM(C$9:C$10)*100)', round(row.p_P1, 6),
                int(row.eventos), round(row.contas_meses, 1), f'=G{r}/H{r}*100', f'=I{r}/(SUM(G$9:G$10)/SUM(H$9:H$10)*100)', round(row.p_S1, 6)]
        for j, x in enumerate(vals, 1):
            c_ = ws.cell(r, j, x); c_.font = F(size=9); c_.border = BOX
            c_.number_format = {4: '0.00', 9: '0.00', 5: '0.0"x"', 10: '0.0"x"', 6: '0.0000', 11: '0.0000', 3: '#,##0', 8: '#,##0'}.get(j, 'General')
    ws.cell(11, 1, 'Leitura: nos primeiros 30 dias o risco é várias vezes maior que depois de 180 dias, tanto nas assinaturas (P1) quanto nos eventos de churn (S1). churn_flag (S2) não tem data e não entra aqui.').font = F(italic=True, size=9, color=GRAY)
    ch = BarChart(); ch.type = 'col'; ch.title = 'P1: assinaturas encerradas por 100 assinaturas-mês, por idade'; ch.height = 7; ch.width = 15
    ch.add_data(Reference(ws, min_col=4, min_row=5, max_row=10), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=10)); ch.legend = None
    ch.series[0].graphicalProperties.solidFill = '2A78D6'; ws.add_chart(ch, 'M4')
    ch = BarChart(); ch.type = 'col'; ch.title = 'S1: eventos de churn por 100 contas-mês, por idade da conta'; ch.height = 7; ch.width = 15
    ch.add_data(Reference(ws, min_col=9, min_row=5, max_row=10), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=10)); ch.legend = None
    ch.series[0].graphicalProperties.solidFill = 'EB6834'; ws.add_chart(ch, 'V4')
    r0 = 22; ws.cell(r0, 1, '2. Por mês: o aumento recente é só porque há mais assinaturas novas?').font = F(bold=True, size=12)
    ws.cell(r0 + 1, 1, '"Esperado pela idade" = quantas assinaturas encerrariam no mês se cada faixa de idade tivesse o risco médio do período todo. Razão > 1 = churn acima do explicado pela quantidade de assinaturas novas.').font = F(italic=True, size=9, color=GRAY)
    hd(r0 + 2, ['Mês', 'P1: encerradas (observado)', 'P1: esperado pela idade', 'P1: razão observado ÷ esperado', 'P1: p', 'P1: assinaturas-mês', 'P1: por 100 assinaturas-mês',
                'S1: eventos de churn', 'S1: contas-mês', 'S1: por 100 contas-mês'], [HD] + [HP] * 6 + [HS] * 3)
    for i, row in T2.iterrows():
        r = r0 + 3 + i
        vals = [row.mes.to_pydatetime(), int(row.encerradas), round(row.esperado_pela_idade, 2), f'=IF(C{r}>0,B{r}/C{r},"")', None if pd.isna(row.p_obs_vs_esp) else round(row.p_obs_vs_esp, 6),
                round(row.assin_meses, 1), f'=B{r}/F{r}*100', int(row.eventos), round(row.contas_meses, 1), f'=H{r}/I{r}*100']
        for j, x in enumerate(vals, 1):
            c_ = ws.cell(r, j, x); c_.font = F(size=9); c_.border = BOX
            c_.number_format = {1: 'mmm/yy', 3: '0.0', 4: '0.00"x"', 5: '0.0000', 6: '#,##0', 7: '0.00', 9: '#,##0', 10: '0.00'}.get(j, 'General')
    r2l = r0 + 3 + len(T2) - 1
    ws.conditional_formatting.add(f'D{r0+3}:D{r2l}', CellIsRule(operator='greaterThan', formula=['1.5'], fill=RED))
    ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.title = 'P1: encerradas observado vs esperado pela idade'; ch.height = 8; ch.width = 22
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=r0 + 2, max_row=r2l), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 3, max_row=r2l))
    ch.series[0].graphicalProperties.solidFill = '2A78D6'; ch.series[1].graphicalProperties.solidFill = 'C3C2B7'; ch.x_axis.number_format = 'mmm/yy'; ch.legend.position = 'b'; ws.add_chart(ch, 'M22')
    ch = LineChart(); ch.title = 'Taxas mensais: P1 (assinaturas) e S1 (eventos)'; ch.height = 8; ch.width = 22
    ch.add_data(Reference(ws, min_col=7, min_row=r0 + 2, max_row=r2l), titles_from_data=True); ch.add_data(Reference(ws, min_col=10, min_row=r0 + 2, max_row=r2l), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 3, max_row=r2l)); ch.x_axis.number_format = 'mmm/yy'; ch.legend.position = 'b'
    ch.series[0].graphicalProperties.line.solidFill = '2A78D6'; ch.series[1].graphicalProperties.line.solidFill = 'EB6834'; ws.add_chart(ch, 'M40')
    r3 = r2l + 3; ws.cell(r3, 1, '3. Por coorte de início: assinaturas novas estão cancelando mais rápido?').font = F(bold=True, size=12)
    ws.cell(r3 + 1, 1, 'Só assinaturas iniciadas até 01/12/2024 (todas tiveram 30 dias para serem observadas).').font = F(italic=True, size=9, color=GRAY)
    hd(r3 + 2, ['Trimestre de início', 'Assinaturas', 'Encerradas em até 30 dias', '% encerradas em até 30 dias'], [HD, HP, HP, HP])
    for i, row in T3.iterrows():
        r = r3 + 3 + i
        for j, x in enumerate([row.coorte, int(row.assinaturas), int(row.encerradas_30d), f'=C{r}/B{r}'], 1):
            c_ = ws.cell(r, j, x); c_.font = F(size=9); c_.border = BOX; c_.number_format = '0.0%' if j == 4 else 'General'
    r3l = r3 + 3 + len(T3) - 1
    ch = BarChart(); ch.type = 'col'; ch.title = '% de assinaturas encerradas em até 30 dias, por trimestre de início'; ch.height = 7; ch.width = 15
    ch.add_data(Reference(ws, min_col=4, min_row=r3 + 2, max_row=r3l), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=r3 + 3, max_row=r3l)); ch.legend = None
    ch.y_axis.number_format = '0%'; ch.series[0].graphicalProperties.solidFill = '2A78D6'; ws.add_chart(ch, 'V22')
    ws.column_dimensions['A'].width = 22
    for j in range(2, 12): ws.column_dimensions[L(j)].width = 13
    # ---- LEIA-ME ----
    ws = leia; ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
    L_ = [('Passo 2 da reconciliação: triangulação das causas candidatas', dict(bold=True, size=15)),
          ('Planilha apartada. Lê os CSVs originais (G4_DESAFIO\\dados) sem alterá-los. Base principal = assinaturas (ver 01_Evidencias_Consistencia.xlsx).', dict(italic=True, color=GRAY)),
          ('', {}), ('PERGUNTA', dict(bold=True, size=12)),
          ('Algum fator (segmento, plano, uso, suporte, tempo) muda o churn de forma consistente nas três definições, e não só por acaso?', {}),
          ('', {}), ('AS QUATRO MÉTRICAS', dict(bold=True, size=12)),
          ('• P1 (principal): assinaturas encerradas por 100 assinaturas-mês de exposição. Exposição corrige o viés de assinaturas recentes terem tido menos tempo para cancelar.', {}),
          ('• P2 (principal, pesada por receita): MRR encerrado ÷ MRR × meses de exposição. Responde "nem todo churn pesa igual".', {}),
          ('• S1 (secundária): eventos de churn (churn_events) por 100 contas-mês desde o cadastro.', {}),
          ('• S2 (secundária): % de contas com churn_flag = TRUE (accounts). Não tem data, por isso não entra na aba Tempo.', {}),
          ('', {}), ('COMO DECIDIMOS SE É SINAL OU ACASO', dict(bold=True, size=12)),
          ('• Teste de permutação: embaralhamos 2.000 vezes os rótulos do fator entre as contas (ou assinaturas) e vemos com que frequência surge uma diferença tão grande quanto a real (p).', {}),
          ('• Correção por testes múltiplos (Benjamini-Hochberg, q): testamos 68 grupos. Com p < 0,05, uns 3 dariam "positivo" só por acaso. q < 0,05 controla isso.', {}),
          ('• Triangulação: um fator só é "Confirmado" se for significativo no desfecho principal E em ao menos uma secundária, na mesma direção.', {}),
          ('', {}), ('O QUE ENCONTRAMOS', dict(bold=True, size=12)),
          ('• Fatores fixos (aba Matriz e Resumo_Fatores): nenhum dos 22 fatores sobrevive à correção em nenhuma definição. Indústria, país, canal, plano, cobrança, seats, uso, erros, tickets, satisfação e escalação NÃO explicam o churn de forma robusta.', {}),
          ('• Indícios fracos (p < 0,05 antes da correção): contas que tiveram upgrade cancelam menos assinaturas; contas com plano predominante Basic (só 11) cancelam mais. Não aparecem nas secundárias: não sustentam conclusão.', {}),
          ('• Tempo (aba Tempo): é aqui que está o sinal, e ele aparece em duas definições (P1 e S1): (1) os primeiros 30 dias concentram o risco; (2) dez/2024 é anômalo: 193 assinaturas encerradas contra ~75 esperadas pela idade (2,6x); out–nov ficam perto do esperado; (3) as assinaturas iniciadas em 2024-T3 e T4 cancelam muito mais nos primeiros 30 dias (2,4% e 6,0%, contra ≤ 1% nas coortes anteriores).', {}),
          ('', {}), ('LIMITES', dict(bold=True, size=12)),
          ('• Idade e calendário se misturam numa janela de 2 anos (as assinaturas novas são as mais recentes). A aba Tempo separa os dois com a razão observado ÷ esperado, mas é uma aproximação.', {}),
          ('• Ausência de sinal num fator não prova que ele não importa; prova que estes dados não mostram efeito. Base sintética.', {}),
          ('• p e q são números calculados no script (permutação não é fórmula de Excel). Taxas, lifts e vereditos são fórmulas sobre as abas Base_* e podem ser conferidos.', {})]
    for i, (tx, st) in enumerate(L_, 1):
        x = ws.cell(i, 2, tx); x.font = F(**st) if st else F(); x.alignment = Alignment(wrap_text=True, vertical='top')
    wb.calculation.fullCalcOnLoad = True; wb.active = 1
    Path(OUT).parent.mkdir(parents=True, exist_ok=True); wb.save(OUT); return M, G, T1, T2, T3

if __name__ == '__main__':
    OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_RECONCILIA' / '02_Triangulacao.xlsx'
    M, G, T1, T2, T3 = planilha(OUT); print('ok', OUT)
    print(T1.round(4).to_string()); print(T2.round(3).tail(8).to_string()); print(T3.to_string())
