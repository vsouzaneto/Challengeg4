# Datas de tickets e de uso: dá para recuperar a data real? E o evento de churn: dá para ligá-lo ao cancelamento real (assinaturas)?
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada.
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from comum import *
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_CONSOLIDA_V1' / '04_Eventos_x_Cancelamento_Real.xlsx'
a, s = ler(DADOS); s = classificar(s)
t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv', parse_dates=['submitted_at']); c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])
f = pd.read_csv(DADOS/'ravenstack_feature_usage.csv', parse_dates=['usage_date'])
RNG = np.random.default_rng(2024); FIM = pd.Timestamp('2024-12-31')
sem = lambda d: f'{d.year}-S{1 if d.month <= 6 else 2}'
sn = lambda b: np.where(b, 'Sim', 'Não')
# ---------- tickets ----------
T = t.merge(a[['account_id', 'signup_date']], on='account_id'); T['trimestre'] = T.submitted_at.apply(trimestre); T['ano_ticket'] = T.submitted_at.dt.year
T['semestre_cadastro_conta'] = T.signup_date.apply(sem); T['antes_do_cadastro'] = sn(T.submitted_at < T.signup_date)
NS = s.groupby('account_id').size(); a['assinaturas'] = a.account_id.map(NS).fillna(0).astype(int); a['tickets'] = a.account_id.map(t.groupby('account_id').size()).fillna(0).astype(int)
a['faixa_assinaturas'] = pd.cut(a.assinaturas, [0, 5, 10, 15, 100], labels=['1–5', '6–10', '11–15', '16+']).astype(str); a['semestre_cadastro'] = a.signup_date.apply(sem)
# ---------- uso ----------
U = f.merge(s[['subscription_id', 'account_id', 'start_date', 'end_date']], on='subscription_id'); U['trimestre'] = U.usage_date.apply(trimestre); U['ano_uso'] = U.usage_date.dt.year
U['semestre_inicio_assinatura'] = U.start_date.apply(sem); U['antes_do_inicio'] = sn(U.usage_date < U.start_date); U['depois_do_fim'] = sn(U.end_date.notna() & (U.usage_date > U.end_date))
# ---------- eventos x fim real de assinatura ----------
E = s[s.end_date.notna()]
def mais_proximo(acc, d):
    e = E[E.account_id == acc]
    if len(e) == 0: return None, None, None
    x = (e.end_date - d).dt.days; i = x.abs().idxmin(); return e.loc[i, 'subscription_id'], int(x[i]), e.loc[i, 'situacao']
rows = []
for r in c.itertuples():
    sid, dist, sit = mais_proximo(r.account_id, r.churn_date)
    ativa = ((s.account_id == r.account_id) & (s.start_date <= r.churn_date) & (s.end_date.isna() | (s.end_date > r.churn_date))).any()
    rows.append(dict(churn_event_id=r.churn_event_id, account_id=r.account_id, churn_date=r.churn_date, trimestre=trimestre(r.churn_date), motivo=r.reason_code, comentario=r.feedback_text if isinstance(r.feedback_text, str) else '(vazio)',
                     conta_ativa_no_dia=('Sim' if ativa else 'Não'), fim_mais_proximo=sid or '', dias_ate_o_fim_mais_proximo=dist, situacao_do_fim_mais_proximo=sit or 'Conta sem nenhum fim de assinatura'))
EV = pd.DataFrame(rows)
EV['dias_ate_o_fim_mais_proximo'] = pd.to_numeric(EV.dias_ate_o_fim_mais_proximo)
EV['faixa_distancia'] = pd.cut(EV.dias_ate_o_fim_mais_proximo.abs(), [-1, 7, 30, 90, 10**5], labels=['Até 7 dias', '8 a 30 dias', '31 a 90 dias', 'Mais de 90 dias']).astype(object)
EV.loc[EV.dias_ate_o_fim_mais_proximo.isna(), 'faixa_distancia'] = 'Conta sem nenhum fim de assinatura'
EV['ligado_a_um_fim_em_7_dias'] = sn(EV.dias_ate_o_fim_mais_proximo.abs() <= 7); EV['ligado_a_um_fim_em_30_dias'] = sn(EV.dias_ate_o_fim_mais_proximo.abs() <= 30)
# referência: a mesma conta, numa data sorteada entre o cadastro e 31/12/2024
sig = a.set_index('account_id').signup_date
def perto(acc, d, w): e = E[E.account_id == acc]; return bool(((e.end_date - d).dt.days.abs() <= w).any())
REF = {}
for w in (7, 30):
    v = []
    for r in c.itertuples():
        e = E[E.account_id == r.account_id].end_date.values; d0 = sig[r.account_id]
        for _ in range(20):   # 20 datas sorteadas por evento
            d = np.datetime64(d0 + pd.Timedelta(days=int(RNG.integers(0, (FIM - d0).days + 1))))
            v.append(bool(len(e)) and bool((np.abs((e - d).astype('timedelta64[D]').astype(int)) <= w).any()))
    REF[w] = float(np.mean(v))
# fins reais x evento
EN = E.copy(); EN['trimestre'] = EN.end_date.apply(trimestre)
EN['tem_evento_em_7_dias'] = sn([((c.account_id == r.account_id) & ((c.churn_date - r.end_date).abs().dt.days <= 7)).any() for r in EN.itertuples()])
EN['tem_evento_em_30_dias'] = sn([((c.account_id == r.account_id) & ((c.churn_date - r.end_date).abs().dt.days <= 30)).any() for r in EN.itertuples()])
# contas: marcação churn_flag x perda real
perda_acc = set(E[E.situacao.isin(['Redução real', 'Saída do cliente'])].account_id); sai_acc = set(E[E.situacao == 'Saída do cliente'].account_id)
a['marcada_churn'] = sn(a.churn_flag); a['teve_perda_real'] = sn(a.account_id.isin(perda_acc)); a['teve_evento'] = sn(a.account_id.isin(c.account_id))
a['ativa_em_31_12_2024'] = sn(a.account_id.isin(s[s.end_date.isna()].account_id))

lv = Livro(); leia = lv.aba('LEIA-ME'); w1 = lv.aba('1_Datas_dos_tickets'); w2 = lv.aba('2_Datas_do_uso'); w3 = lv.aba('3_Evento_x_cancelamento'); w4 = lv.aba('4_Cancelamento_x_evento'); w5 = lv.aba('5_Regra_proposta')
lv.base('Base_Tickets', T, ['ticket_id', 'account_id', 'submitted_at', 'trimestre', 'ano_ticket', 'signup_date', 'semestre_cadastro_conta', 'antes_do_cadastro'], {'submitted_at': DT, 'signup_date': DT})
lv.base('Base_Uso', U, ['usage_id', 'subscription_id', 'usage_date', 'trimestre', 'ano_uso', 'start_date', 'end_date', 'semestre_inicio_assinatura', 'antes_do_inicio', 'depois_do_fim'], {'usage_date': DT, 'start_date': DT, 'end_date': DT})
lv.base('Base_Eventos', EV, list(EV.columns), {'churn_date': DT})
lv.base('Base_Encerradas', EN, ['subscription_id', 'account_id', 'plan_tier', 'mrr_amount', 'end_date', 'trimestre', 'situacao', 'tem_evento_em_7_dias', 'tem_evento_em_30_dias'], {'end_date': DT, 'mrr_amount': USD})
lv.base('Base_Assinaturas', s, ['subscription_id', 'account_id', 'start_date', 'end_date', 'situacao'], {'start_date': DT, 'end_date': DT})
lv.base('Base_Contas', a, ['account_id', 'signup_date', 'semestre_cadastro', 'assinaturas', 'faixa_assinaturas', 'tickets', 'marcada_churn', 'teve_evento', 'teve_perda_real', 'ativa_em_31_12_2024'], {'signup_date': DT})
HB2, HR_, HP, HN = fill('2A78D6'), fill('B45309'), fill('B91C1C'), fill('0E7C66')
SEMS = ['2023-S1', '2023-S2', '2024-S1', '2024-S2']
def R(b, col): return lv.R(b, col)

# ---------- 1_Datas_dos_tickets ----------
ws = w1; titulo(ws, 'A data dos tickets é real? Dá para recuperá-la?', 'Três testes: (A) volume por trimestre; (B) contas que se cadastraram em cada semestre x ano do ticket; (C) tickets por conta conforme o tamanho da conta.', 'H')
H(ws, 4, ['(A) Trimestre', 'Tickets', 'Contas cadastradas até o fim do trimestre'], [HD, HB2, HX], 34); r = 5
for q in TRIM:
    y, n = int(q[:4]), int(q[-1]); fim_q = f'DATE({y},{3 * n},{[31, 30, 30, 31][n - 1]})'
    C(ws, r, 1, q, bold=True); C(ws, r, 2, f'=COUNTIFS({R("Base_Tickets", "trimestre")},$A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({R("Base_Contas", "signup_date")},"<="&{fim_q})', '#,##0'); r += 1
r += 1; H(ws, r, ['(B) Contas cadastradas em', 'Tickets com data em 2023', 'Tickets com data em 2024', '% com data ANTES do cadastro da conta'], [HD, HB2, HB2, HP], 40); r += 1
for se in SEMS:
    C(ws, r, 1, se, bold=True)
    for j, y in ((2, 2023), (3, 2024)): C(ws, r, j, f'=COUNTIFS({R("Base_Tickets", "semestre_cadastro_conta")},$A{r},{R("Base_Tickets", "ano_ticket")},{y})', '#,##0')
    C(ws, r, 4, f'=COUNTIFS({R("Base_Tickets", "semestre_cadastro_conta")},$A{r},{R("Base_Tickets", "antes_do_cadastro")},"Sim")/(B{r}+C{r})', '0%'); r += 1
C(ws, r, 1, 'Todas', bold=True, fl=GRY); C(ws, r, 2, f'=SUM(B{r-4}:B{r-1})', '#,##0', True, GRY); C(ws, r, 3, f'=SUM(C{r-4}:C{r-1})', '#,##0', True, GRY); C(ws, r, 4, f'=COUNTIFS({R("Base_Tickets", "antes_do_cadastro")},"Sim")/(B{r}+C{r})', '0%', True, GRY); r += 2
H(ws, r, ['(C) Assinaturas da conta', 'Contas', 'Tickets por conta (média)'], [HD, HX, HB2], 34); r += 1
for fx in ['1–5', '6–10', '11–15', '16+']:
    C(ws, r, 1, fx, bold=True); C(ws, r, 2, f'=COUNTIFS({R("Base_Contas", "faixa_assinaturas")},$A{r})', '#,##0'); C(ws, r, 3, f'=AVERAGEIFS({R("Base_Contas", "tickets")},{R("Base_Contas", "faixa_assinaturas")},$A{r})', '0.00'); r += 1
nota(ws, r + 1, 'Leitura: (A) cerca de 250 tickets por trimestre do começo ao fim, mesmo com a base de contas crescendo muito. (B) Contas que só se cadastraram no 2º semestre de 2024 têm METADE dos tickets com data em 2023 — antes de existirem. (C) Conta com 16+ assinaturas abre tantos tickets quanto conta com 1–5. Conclusão: as datas dos tickets estão espalhadas ao acaso entre jan/23 e dez/24. Não há como recuperar a data real a partir destas tabelas: nenhuma regra (deslocamento fixo, data de cadastro + x dias) reconstrói o padrão. Só os tickets com data depois do cadastro da conta (≈46%) são datas possíveis, e mesmo assim não comprovadas.', 'H', 84)
larguras(ws, [22, 14, 14, 18])

# ---------- 2_Datas_do_uso ----------
ws = w2; titulo(ws, 'A data de uso é real?', '(A) Registros de uso por trimestre; (B) assinaturas iniciadas em cada semestre x ano do registro de uso; (C) registros fora da vida da própria assinatura.', 'H')
H(ws, 4, ['(A) Trimestre', 'Registros de uso', 'Assinaturas iniciadas até o fim do trimestre'], [HD, HB2, HX], 34); r = 5
for q in TRIM:
    y, n = int(q[:4]), int(q[-1]); fim_q = f'DATE({y},{3 * n},{[31, 30, 30, 31][n - 1]})'
    C(ws, r, 1, q, bold=True); C(ws, r, 2, f'=COUNTIFS({R("Base_Uso", "trimestre")},$A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({R("Base_Assinaturas", "start_date")},"<="&{fim_q})', '#,##0'); r += 1
r += 1; H(ws, r, ['(B) Assinatura iniciada em', 'Uso com data em 2023', 'Uso com data em 2024', '% com data ANTES do início da assinatura'], [HD, HB2, HB2, HP], 40); r += 1
for se in SEMS:
    C(ws, r, 1, se, bold=True)
    for j, y in ((2, 2023), (3, 2024)): C(ws, r, j, f'=COUNTIFS({R("Base_Uso", "semestre_inicio_assinatura")},$A{r},{R("Base_Uso", "ano_uso")},{y})', '#,##0')
    C(ws, r, 4, f'=COUNTIFS({R("Base_Uso", "semestre_inicio_assinatura")},$A{r},{R("Base_Uso", "antes_do_inicio")},"Sim")/(B{r}+C{r})', '0%'); r += 1
r += 1; H(ws, r, ['(C) Registros de uso', 'Quantidade', '%'], [HD, HB2, HB2], 30); r += 1; rt = r
C(ws, r, 1, 'Total', bold=True); C(ws, r, 2, f'=COUNTA({R("Base_Uso", "usage_id")})', '#,##0'); C(ws, r, 3, 1, '0%'); r += 1
for lab, col in (('Data ANTES do início da assinatura', 'antes_do_inicio'), ('Data DEPOIS do fim da assinatura', 'depois_do_fim')):
    C(ws, r, 1, lab, bold=True); C(ws, r, 2, f'=COUNTIFS({R("Base_Uso", col)},"Sim")', '#,##0'); C(ws, r, 3, f'=B{r}/B{rt}', '0%'); r += 1
nota(ws, r + 1, 'Leitura: o uso tem o mesmo volume em todos os trimestres e as assinaturas do 2º semestre de 2024 têm metade do uso com data em 2023 — antes de começarem. Como nos tickets, as datas de uso estão espalhadas ao acaso e não dá para recuperar a data real. A quantidade de uso por assinatura (≈5 registros cada) pode ser usada; a data, não.', 'H', 56)
larguras(ws, [30, 14, 14, 18])

# ---------- 3_Evento_x_cancelamento ----------
BE = 'Base_Eventos'
ws = w3; titulo(ws, 'Cada evento de churn está ligado a um cancelamento real?', 'Para cada um dos 600 eventos de churn, buscamos o fim de assinatura MAIS PRÓXIMO da mesma conta (tabela de assinaturas) e medimos a distância em dias. Para saber se a proximidade é real ou acaso, comparamos com uma data SORTEADA na vida da mesma conta.', 'G', 44)
H(ws, 4, ['Distância até o fim de assinatura mais próximo', 'Eventos', '%'], [HD, HP, HP], 40); r = 5
for fx in ['Até 7 dias', '8 a 30 dias', '31 a 90 dias', 'Mais de 90 dias', 'Conta sem nenhum fim de assinatura']:
    C(ws, r, 1, fx, bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, "faixa_distancia")},$A{r})', '#,##0'); C(ws, r, 3, f'=B{r}/COUNTA({R(BE, "churn_event_id")})', '0%'); r += 1
C(ws, r, 1, 'Total', bold=True, fl=GRY); C(ws, r, 2, f'=SUM(B5:B{r-1})', '#,##0', True, GRY); r += 2
H(ws, r, ['Teste', 'Eventos de churn', 'Data sorteada na vida da mesma conta', 'Leitura'], [HD, HP, HX, HX], 40); r += 1
for w in (7, 30):
    C(ws, r, 1, f'Há fim de assinatura da conta a até {w} dias?', bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, f"ligado_a_um_fim_em_{w}_dias")},"Sim")/COUNTA({R(BE, "churn_event_id")})', '0.0%')
    C(ws, r, 3, REF[w], '0.0%'); C(ws, r, 4, 'Igual ao acaso: a data do evento não "sabe" quando a conta cancelou.'); r += 1
C(ws, r, 1, 'A conta tinha assinatura ativa no dia do evento?', bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, "conta_ativa_no_dia")},"Sim")/COUNTA({R(BE, "churn_event_id")})', '0%'); C(ws, r, 3, '—'); C(ws, r, 4, 'O "churn" registrado acontece com o cliente ativo.'); r += 2
H(ws, r, ['Eventos a até 7 dias de um fim: que fim era?', 'Eventos'], [HD, HP], 34); r += 1
for k in ['Migração', 'Compra extra desfeita', 'Redução real', 'Saída do cliente']:
    C(ws, r, 1, k, bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, "ligado_a_um_fim_em_7_dias")},"Sim",{R(BE, "situacao_do_fim_mais_proximo")},$A{r})', '#,##0'); r += 1
r += 1; H(ws, r, ['Motivo registrado', 'Eventos', 'Ligados a um fim em até 7 dias', '%'], [HD, HP, HP, HP], 34); r += 1
for m in sorted(EV.motivo.unique()):
    C(ws, r, 1, m, bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, "motivo")},$A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({R(BE, "motivo")},$A{r},{R(BE, "ligado_a_um_fim_em_7_dias")},"Sim")', '#,##0'); C(ws, r, 4, f'=C{r}/B{r}', '0%'); r += 1
nota(ws, r + 1, 'Leitura: só 8% dos eventos têm um fim de assinatura da conta a até 7 dias — praticamente a mesma proporção de uma data sorteada na vida da conta. E 9 em cada 10 eventos acontecem com o cliente ativo. Não há como ligar o evento ao cancelamento real: a data do evento é independente das assinaturas. O motivo registrado também não muda essa proporção.', 'G', 52)
larguras(ws, [46, 14, 18, 52])

# ---------- 4_Cancelamento_x_evento ----------
BN = 'Base_Encerradas'
ws = w4; titulo(ws, 'Do outro lado: cada cancelamento real tem um evento de churn registrado?', 'Assinaturas encerradas (tabela de assinaturas), pela situação definida na planilha 02, e se a conta registrou algum evento de churn a até 7 ou 30 dias do fim.', 'G')
H(ws, 4, ['Situação do encerramento', 'Assinaturas encerradas', 'Com evento a até 7 dias', '%', 'Com evento a até 30 dias', '%'], [HD, HX, HP, HP, HP, HP], 40); r = 5
for k in ['Migração', 'Compra extra desfeita', 'Redução real', 'Saída do cliente']:
    C(ws, r, 1, k, bold=True); C(ws, r, 2, f'=COUNTIFS({R(BN, "situacao")},$A{r})', '#,##0')
    for j, col in ((3, 'tem_evento_em_7_dias'), (5, 'tem_evento_em_30_dias')): C(ws, r, j, f'=COUNTIFS({R(BN, "situacao")},$A{r},{R(BN, col)},"Sim")', '#,##0'); C(ws, r, j + 1, f'=IF(B{r}>0,{L(j)}{r}/B{r},"")', '0%')
    r += 1
C(ws, r, 1, 'Total', bold=True, fl=GRY)
for j in (2, 3, 5): C(ws, r, j, f'=SUM({L(j)}5:{L(j)}{r-1})', '#,##0', True, GRY)
for j in (4, 6): C(ws, r, j, f'={L(j-1)}{r}/B{r}', '0%', True, GRY)
r += 2; H(ws, r, ['Contas (500)', 'Teve perda real (redução real ou saída)', 'Não teve perda real'], [HD, HP, HN], 40); r += 1
for lab, col in (('Marcada como churn na tabela de contas', 'marcada_churn'), ('Teve evento de churn', 'teve_evento')):
    for v in ('Sim', 'Não'):
        C(ws, r, 1, f'{lab}: {v}', bold=True)
        for j, pr in ((2, 'Sim'), (3, 'Não')): C(ws, r, j, f'=COUNTIFS({R("Base_Contas", col)},"{v}",{R("Base_Contas", "teve_perda_real")},"{pr}")', '#,##0')
        r += 1
C(ws, r, 1, 'Contas ativas em 31/12/2024', bold=True, fl=GRY); C(ws, r, 2, f'=COUNTIFS({R("Base_Contas", "ativa_em_31_12_2024")},"Sim")', '#,##0', True, GRY); r += 1
nota(ws, r + 1, 'Leitura: só cerca de 6% das reduções reais têm um evento de churn a até 7 dias, e nenhuma das 4 saídas de cliente. Das 110 contas marcadas como churn, só 43 tiveram perda real — e outras 110 contas com perda real não estão marcadas. As marcações de churn (na conta e nos eventos) não servem para identificar quem perdeu receita. Nenhuma conta termina 2024 sem assinatura ativa.', 'F', 44)
larguras(ws, [42, 16, 14, 8, 14, 8])

# ---------- 5_Regra_proposta ----------
ws = w5; titulo(ws, 'Proposta: o churn oficial deve sair da tabela de assinaturas', 'Por trimestre: eventos de churn registrados hoje x fins de assinatura x perda real (redução real + saída). Os eventos não acompanham nem o "churn" de hoje nem a perda real.', 'H')
H(ws, 4, ['Trimestre', 'Eventos de churn registrados', 'Assinaturas encerradas ("churn" de hoje)', 'Perda real: redução real + saída (qtd)', 'Perda real (US$ de MRR, sem trocas)'], [HD, HX, fill('6B7280'), HP, HP], 46); r = 5
for q in TRIM:
    C(ws, r, 1, q, bold=True); C(ws, r, 2, f'=COUNTIFS({R(BE, "trimestre")},$A{r})', '#,##0'); C(ws, r, 3, f'=COUNTIFS({R(BN, "trimestre")},$A{r})', '#,##0')
    C(ws, r, 4, f'=COUNTIFS({R(BN, "trimestre")},$A{r},{R(BN, "situacao")},"Redução real")+COUNTIFS({R(BN, "trimestre")},$A{r},{R(BN, "situacao")},"Saída do cliente")', '#,##0')
    C(ws, r, 5, f'=SUMIFS({R(BN, "mrr_amount")},{R(BN, "trimestre")},$A{r},{R(BN, "situacao")},"Redução real")+SUMIFS({R(BN, "mrr_amount")},{R(BN, "trimestre")},$A{r},{R(BN, "situacao")},"Saída do cliente")', USD); r += 1
C(ws, r, 1, 'Total', bold=True, fl=GRY)
for j in (2, 3, 4): C(ws, r, j, f'=SUM({L(j)}5:{L(j)}{r-1})', '#,##0', True, GRY)
C(ws, r, 5, f'=SUM(E5:E{r-1})', USD, True, GRY)
ch = BarChart(); ch.type = 'col'; ch.title = 'Eventos de churn registrados x perda real (quantidade), por trimestre'; ch.height = 7.5; ch.width = 17; ch.gapWidth = 50
for j, nm, cor in ((2, 'Eventos de churn registrados', '6B7280'), (4, 'Perda real (redução real + saída)', 'B91C1C')):
    ch.add_data(Reference(ws, min_col=j, min_row=5, max_row=r - 1), titles_from_data=False); ch.series[-1].tx = SeriesLabel(v=nm); ch.series[-1].graphicalProperties.solidFill = cor; ch.series[-1].graphicalProperties.line.solidFill = cor
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=r - 1)); ch.legend.position = 'b'; ws.add_chart(ch, 'G4')
r += 2
for tx in ['REGRA PROPOSTA (para a empresa):',
           '1. Churn oficial = fim de assinatura na tabela de assinaturas, classificado em: troca (nova assinatura da conta em até 7 dias), compra extra desfeita (até 30 dias, já havia outra do mesmo plano), redução real (conta segue cliente) e saída (conta sem assinatura).',
           '2. Só redução real e saída contam como perda; nas trocas, conta só a diferença de valor.',
           '3. O evento de churn passa a ser criado AUTOMATICAMENTE quando uma assinatura termina como redução real ou saída — com o motivo perguntado ao cliente nesse momento. Os 600 eventos atuais não podem ser ligados aos cancelamentos e não devem ser usados para medir churn.',
           '4. Tickets e uso precisam gravar a data real do registro (hoje as datas estão espalhadas ao acaso). Até lá, usar só contagens, nunca datas, e descartar tickets com data anterior ao cadastro da conta.']:
    ws.cell(r, 1, tx).font = Fn(size=10, bold=tx.startswith('REGRA')); ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(f'A{r}:E{r}'); ws.row_dimensions[r].height = 18 if tx.startswith('REGRA') else 44; r += 1
larguras(ws, [14, 14, 16, 16, 18])

nev7 = int((EV.ligado_a_um_fim_em_7_dias == 'Sim').sum())
leiame(leia, [('DATAS DE TICKETS E DE USO · EVENTO DE CHURN x CANCELAMENTO REAL', dict(bold=True, size=15)),
    ('Planilha apartada. Lê os CSVs originais sem alterá-los. "Cancelamento real" = fim de assinatura na tabela de assinaturas (a única coerente), com as situações da planilha 02.', dict(italic=True, color=GRAY)),
    ('', {}), ('RESPOSTAS', dict(bold=True, size=12)),
    ('1. A data dos tickets é real? Dá para descobrir a data verdadeira?  → aba 1', dict(bold=True)),
    ('   Não e não. O volume é o mesmo em todo trimestre, contas cadastradas no fim de 2024 têm metade dos tickets em 2023, e o número de tickets não cresce com o tamanho da conta. As datas foram espalhadas ao acaso; nenhuma regra reconstrói a data real.', {}),
    ('2. E a data de uso?  → aba 2', dict(bold=True)),
    ('   Mesmo problema: volume fixo por trimestre e metade do uso de assinaturas novas datado antes de elas existirem. Dá para usar a quantidade de uso, não a data.', {}),
    ('3. Dá para ligar o evento de churn ao cancelamento real?  → abas 3 e 4', dict(bold=True)),
    (f'   Não. Só {nev7} dos 600 eventos ({nev7 / 600:.0%}) têm um fim de assinatura da conta a até 7 dias — praticamente a mesma proporção de uma data sorteada na vida da conta ({REF[7]:.0%}). 9 em cada 10 eventos acontecem com o cliente ativo. E do outro lado, só ~6% das reduções reais têm evento perto. Evento e cancelamento são independentes.', {}),
    ('4. O que fazer?  → aba 5', dict(bold=True)),
    ('   Usar a tabela de assinaturas como fonte oficial do churn (troca / compra extra desfeita / redução real / saída) e passar a gerar o evento de churn a partir do fim real da assinatura, perguntando o motivo nesse momento. Corrigir a gravação das datas de tickets e uso.', {}),
    ('', {}), ('LIMITES', dict(bold=True, size=12)),
    ('• Base fictícia (gerada por computador): as datas aleatórias podem ser efeito do gerador. Os valores "data sorteada" (aba 3) são calculados pelo programa com sorteio fixo, por isso aparecem como número e não como fórmula.', {})])
lv.salvar(OUT); print('ok', OUT); print('ev7', nev7, 'ref', REF)
