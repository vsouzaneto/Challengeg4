# Funções comuns às planilhas de DADOS_CONSOLIDA_V1 (classificação das assinaturas, testes e estilo).
# Gerado com Claude (IA). Só LÊ os CSVs originais.
from math import lgamma, exp
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo

JP = 7          # janela de troca (migração): nova assinatura da mesma conta até 7 dias antes/depois
EXTRA = 30      # compra extra desfeita: assinatura com até 30 dias de vida, havendo outra mais antiga do mesmo plano ativa na conta
RAJ_N, RAJ_D = 3, 7   # rajada: 3+ assinaturas da mesma conta em 7 dias corridos

def ler(DADOS):
    a = pd.read_csv(DADOS/'ravenstack_accounts.csv', parse_dates=['signup_date'])
    s = pd.read_csv(DADOS/'ravenstack_subscriptions.csv', parse_dates=['start_date', 'end_date'])
    return a, s

def classificar(s):
    """Para cada assinatura encerrada: Migração / Compra extra desfeita / Redução real / Saída do cliente (+ MRR reposto nas migrações)."""
    s = s.copy()
    def tipo(r):
        o = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
        at = o[(o.start_date <= r.end_date) & (o.end_date.isna() | (o.end_date > r.end_date))]
        pe = o[(o.start_date - r.end_date).abs().dt.days <= JP]; de = o[(o.start_date > r.end_date) & (o.start_date <= r.end_date + pd.Timedelta(days=JP))]
        if len(at) == 0 and len(de) == 0: return 'Saída do cliente'
        if len(pe): return 'Migração'
        vida = (r.end_date - r.start_date).days
        mais_antiga = at[(at.plan_tier == r.plan_tier) & (at.start_date < r.start_date)]
        return 'Compra extra desfeita' if (vida <= EXTRA and len(mais_antiga)) else 'Redução real'
    s['situacao'] = 'Ativa'; m = s.end_date.notna(); s.loc[m, 'situacao'] = s[m].apply(tipo, axis=1)
    # MRR reposto: assinatura nova mais próxima (até 7 dias) de cada migração; cada nova contada uma única vez, em ordem de data
    par = []
    for r in s[s.situacao == 'Migração'].sort_values('end_date').itertuples():
        o = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id)]
        pe = o[(o.start_date - r.end_date).abs().dt.days <= JP]
        p = pe.iloc[(pe.start_date - r.end_date).abs().argsort()].iloc[0]
        par.append(dict(subscription_id=r.subscription_id, assinatura_nova=p.subscription_id, mrr_nova=int(p.mrr_amount)))
    P = pd.DataFrame(par); P['ja_contada'] = P.assinatura_nova.duplicated()
    P['mrr_reposto'] = np.where(P.ja_contada, 0, P.mrr_nova)
    s = s.merge(P[['subscription_id', 'assinatura_nova', 'mrr_reposto']], on='subscription_id', how='left')
    s['mrr_reposto'] = s.mrr_reposto.fillna(0).astype(int); s['assinatura_nova'] = s.assinatura_nova.fillna('')
    s['dias_de_vida'] = (s.end_date - s.start_date).dt.days
    return s

def rajada(s):
    s = s.sort_values(['account_id', 'start_date'])
    def f(gr):
        d = gr.start_date.values; fl = np.zeros(len(gr), bool)
        for i in range(len(gr)):
            w = (d >= d[i]) & (d <= d[i] + np.timedelta64(RAJ_D - 1, 'D'))
            if w.sum() >= RAJ_N: fl |= w
        return pd.Series(fl, index=gr.index)
    return s.groupby('account_id', group_keys=False).apply(f).reindex(s.index)

def trimestre(d): return '' if pd.isna(d) else f'{d.year}-T{(d.month - 1) // 3 + 1}'
TRIM = [f'{y}-T{q}' for y in (2023, 2024) for q in (1, 2, 3, 4)]

# ---------- "diferença confiável?" ----------
def fisher_p(a_, b_, c_, d_):
    n1, n2, m1 = a_ + b_, c_ + d_, a_ + c_; N = n1 + n2
    if N == 0 or n1 == 0 or n2 == 0 or m1 == 0 or m1 == N: return 1.0
    lc = lambda n, k: lgamma(n + 1) - lgamma(k + 1) - lgamma(n - k + 1)
    pmf = lambda x: exp(lc(n1, x) + lc(n2, m1 - x) - lc(N, m1)); p0 = pmf(a_); lo, hi = max(0, m1 - n2), min(m1, n1)
    return min(1.0, sum(pmf(x) for x in range(lo, hi + 1) if pmf(x) <= p0 * (1 + 1e-7)))
def conf(k1, n1, k2, n2):
    k1, n1, k2, n2 = int(k1), int(n1), int(k2), int(n2)
    if k1 + k2 < 10: return 'Poucos casos para concluir'
    return 'Sim' if fisher_p(k1, n1 - k1, k2, n2 - k2) < 0.05 else 'Não (pode ser coincidência)'

# ---------- estilo ----------
FONT = 'Arial'; GRAY = '52514E'
def Fn(**k): return Font(name=FONT, **k)
def fill(c): return PatternFill('solid', fgColor=c)
HD, HX = fill('1F2937'), fill('374151'); YEL, GRY = fill('FFF3CD'), fill('EDEDED')
COR = {'Migração': '1BAF7A', 'Compra extra desfeita': '7FC8A9', 'Redução real': 'EDA100', 'Saída do cliente': 'E34948', 'Diferença na troca': 'B45309', 'Churn de hoje': '6B7280'}
thin = Side(style='thin', color='D0D0D0'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '"US$ "#,##0;-"US$ "#,##0;"-"'; DT = 'dd/mm/yyyy'

class Livro:
    def __init__(self):
        self.wb = Workbook(); self.wb.remove(self.wb.active); self.BS = {}
    def aba(self, n):
        ws = self.wb.create_sheet(n); ws.sheet_view.showGridLines = False; return ws
    @staticmethod
    def py(v):
        if v is None or v is pd.NaT or (isinstance(v, float) and np.isnan(v)): return None
        if isinstance(v, pd.Timestamp): return v.to_pydatetime()
        return v.item() if hasattr(v, 'item') else v
    def base(self, nome, df, cols, fmts={}):
        ws = self.wb.create_sheet(nome); ws.append(cols)
        for row in df[cols].itertuples(index=False): ws.append([self.py(v) for v in row])
        n = len(df) + 1; self.BS[nome] = (cols, n)
        for i, cn in enumerate(cols, 1):
            h = ws.cell(1, i); h.font = Fn(bold=True, color='FFFFFF', size=9); h.fill = HD; ws.column_dimensions[L(i)].width = max(11, min(24, len(cn) + 2))
            if cn in fmts:
                for r in range(2, n + 1): ws.cell(r, i).number_format = fmts[cn]
        tb = Table(displayName='t' + nome, ref=f'A1:{L(len(cols))}{n}'); tb.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(tb); ws.freeze_panes = 'B2'
    def R(self, b, col):
        cols, n = self.BS[b]; i = cols.index(col) + 1; return f"{b}!${L(i)}$2:${L(i)}${n}"
    def salvar(self, OUT, ativa=1):
        self.wb.calculation.fullCalcOnLoad = True; self.wb.active = ativa
        OUT.parent.mkdir(parents=True, exist_ok=True); self.wb.save(OUT)

def H(ws, row, labels, fills, h=46):
    for j, (x, fl) in enumerate(zip(labels, fills), 1):
        if x is None: continue
        c_ = ws.cell(row, j, x); c_.font = Fn(bold=True, color='FFFFFF', size=9); c_.fill = fl
        c_.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c_.border = BOX
    ws.row_dimensions[row].height = h
def C(ws, r, c_, v, fmt=None, bold=False, fl=None, wrap=True):
    x = ws.cell(r, c_, v); x.font = Fn(size=9, bold=bold); x.border = BOX; x.alignment = Alignment(vertical='center', wrap_text=wrap)
    if fmt: x.number_format = fmt
    if fl: x.fill = fl
    return x
def titulo(ws, t1, t2, merge='H', h=40):
    ws['A1'] = t1; ws['A1'].font = Fn(bold=True, size=15)
    ws['A2'] = t2; ws['A2'].font = Fn(italic=True, size=9, color=GRAY); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(f'A2:{merge}2'); ws.row_dimensions[2].height = h
def nota(ws, r, tx, merge='H', h=40):
    ws.cell(r, 1, tx).font = Fn(italic=True, size=9, color=GRAY); ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(f'A{r}:{merge}{r}'); ws.row_dimensions[r].height = h
def larguras(ws, ws_):
    for j, w in enumerate(ws_, 1): ws.column_dimensions[L(j)].width = w
def amarelo(ws, rng): ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Sim"'], fill=YEL))
def leiame(ws, linhas):
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 125
    for i, (tx, st_) in enumerate(linhas, 1):
        x = ws.cell(i, 2, tx); x.font = Fn(**st_) if st_ else Fn(); x.alignment = Alignment(wrap_text=True, vertical='top')
