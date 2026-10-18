# As 8 contas Enterprise que, ao reduzir, perderam mais da metade do que pagavam: o que aconteceu com cada uma (+ indústria, canal, país).
# Gerado com Claude (IA). Lê os CSVs originais sem alterá-los; grava planilha apartada.
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from comum import *
RAIZ = Path(__file__).resolve().parents[3]
DADOS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / 'dados'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / 'analise' / 'DADOS_CONSOLIDA_V1' / '01_Contas_Enterprise_Prioritarias.xlsx'
a, s = ler(DADOS); s = classificar(s)
t = pd.read_csv(DADOS/'ravenstack_support_tickets.csv', parse_dates=['submitted_at']); c = pd.read_csv(DADOS/'ravenstack_churn_events.csv', parse_dates=['churn_date'])
FIM24 = pd.Timestamp('2024-12-31')
# reduções Enterprise pagas (redução real ou compra extra desfeita) em que a conta seguiu com outra Enterprise
red = s[(s.plan_tier == 'Enterprise') & (~s.is_trial) & s.situacao.isin(['Redução real', 'Compra extra desfeita'])]
rows = []
for r in red.itertuples():
    o = s[(s.account_id == r.account_id) & (s.subscription_id != r.subscription_id) & (s.start_date <= r.end_date) & (s.end_date.isna() | (s.end_date > r.end_date))]
    oe = o[o.plan_tier == 'Enterprise']
    if len(oe) == 0: continue
    antes = r.mrr_amount + o.mrr_amount.sum()
    rows.append(dict(subscription_id=r.subscription_id, account_id=r.account_id, pct=r.mrr_amount / antes, antes=antes, depois=o.mrr_amount.sum(), o=o, oe=oe))
X = pd.DataFrame(rows); N_SEGUIU = len(X)
P = X[X.pct > 0.5].sort_values('pct', ascending=False).reset_index(drop=True)
CONTAS = list(P.account_id)
fmt = lambda v: f'US$ {v:,.0f}'.replace(',', '.'); dd = lambda d: d.strftime('%d/%m/%Y')
info = []
for i, p in P.iterrows():
    r = s[s.subscription_id == p.subscription_id].iloc[0]; ac = a[a.account_id == r.account_id].iloc[0]
    o, oe = p.o, p.oe
    if r.situacao == 'Compra extra desfeita': grupo = 'A. Compra extra desfeita'
    elif (oe.start_date > r.start_date).any(): grupo = 'B. Pacote antigo desligado (havia Enterprise mais nova)'
    else: grupo = 'C. Redução pura (sem Enterprise mais nova)'
    gemea = 'Sim' if (oe.seats == r.seats).any() else 'Não'
    fim = int(s[(s.account_id == r.account_id) & s.end_date.isna()].mrr_amount.sum())
    rec = 'Sim — hoje paga mais do que antes' if fim >= p.antes else ('Não — perda mantida' if fim >= p.depois else 'Não — continuou caindo')
    if grupo.startswith('A'): rec = 'Compra a mais desfeita — voltou ao valor de antes da compra'
    ev = c[c.account_id == r.account_id].sort_values('churn_date'); tk = t[t.account_id == r.account_id]
    vida = int(r.dias_de_vida); cob = 'anual' if r.billing_frequency == 'annual' else 'mensal'
    tx = [f'Encerrou uma Enterprise de {r.seats} seats ({cob}) com {vida} dias de vida, em {dd(r.end_date)}, e deixou de pagar {fmt(r.mrr_amount)}/mês — {p.pct:.0%} do que a conta pagava ({fmt(p.antes)}).']
    tx.append(f'Seguiu cliente com {len(o)} assinaturas ({len(oe)} Enterprise' + (', uma delas com os mesmos seats' if gemea == 'Sim' else '') + f'), que somavam {fmt(p.depois)}/mês.')
    nz = int((o.mrr_amount == 0).sum())
    if nz: tx.append((f'{nz} delas eram trial' if nz > 1 else '1 delas era trial') + ' (US$ 0), por isso a perda pesou tanto no total da conta.')
    tr = s[(s.account_id == r.account_id) & (s.plan_tier == 'Enterprise') & s.is_trial & ((s.start_date - r.end_date).abs().dt.days <= 30)]
    if len(tr): tx.append(f'Perto do cancelamento abriu uma Enterprise TRIAL de {int(tr.seats.iloc[0])} seats (US$ 0) em {dd(tr.start_date.iloc[0])}: pode ter trocado o contrato pago por um teste.')
    if r.billing_frequency == 'annual' and vida < 365: tx.append(f'Era contrato ANUAL encerrado com {vida} dias: cancelamento antecipado (verificar multa/reembolso).')
    if grupo.startswith('A'): tx.append('Já tinha outra Enterprise mais antiga: parece compra a mais desfeita, não perda de cliente.')
    tx.append(f'Em 31/12/2024 a conta paga {fmt(fim)}/mês: ' + ('recuperou e hoje paga mais do que antes da redução.' if rec.startswith('Sim') else 'perda mantida.' if rec.endswith('mantida') else 'continuou reduzindo depois.' if rec.endswith('caindo') else 'o mesmo que pagava antes da compra a mais.'))
    if len(ev):
        e_tx = '; '.join(f"{dd(e.churn_date)} ({e.reason_code}{', ' + e.feedback_text if isinstance(e.feedback_text, str) else ''})" for e in ev.itertuples())
        perto = ((ev.churn_date - r.end_date).abs().dt.days <= 30).any()
        tx.append(f'Eventos de churn registrados: {e_tx} — ' + ('um deles fica a até 30 dias do cancelamento.' if perto else 'nenhum perto do cancelamento (não explicam a redução).'))
    else: tx.append('Nenhum evento de churn registrado para a conta.')
    nt = len(tk); na = int((tk.submitted_at < ac.signup_date).sum())
    tx.append(f'Tickets: {nt}' + (f' ({na} com data anterior ao cadastro da conta — datas não confiáveis).' if na else '.'))
    info.append(dict(ordem=i + 1, account_id=r.account_id, industry=ac.industry, country=ac.country, referral_source=ac.referral_source, signup_date=ac.signup_date, subscription_id=r.subscription_id, seats=int(r.seats),
                     cobranca=cob, start_date=r.start_date, end_date=r.end_date, dias_de_vida=vida, mrr_perdido=int(r.mrr_amount), grupo=grupo, gemea=gemea, recuperou=rec,
                     eventos=len(ev), motivos=', '.join(ev.reason_code) if len(ev) else '—', tickets=nt, tickets_antes_cadastro=na, explicacao=' '.join(tx)))
I = pd.DataFrame(info)
S8 = s[s.account_id.isin(CONTAS)].copy().sort_values(['account_id', 'start_date'])
S8['fim_para_conta'] = S8.end_date.fillna(pd.Timestamp('2099-12-31')); S8['trial'] = np.where(S8.is_trial, 'Sim', 'Não'); S8['cobranca'] = np.where(S8.billing_frequency == 'annual', 'anual', 'mensal')
S8['eh_a_reducao_analisada'] = np.where(S8.subscription_id.isin(I.subscription_id), 'Sim', 'Não')
S8['situacao_txt'] = S8.situacao.replace({'Ativa': 'Ativa em 31/12/2024'})
lv = Livro(); leia = lv.aba('LEIA-ME'); w1 = lv.aba('1_As_8_Contas'); w2 = lv.aba('2_Linha_do_tempo'); w3 = lv.aba('3_Perfil')
B = 'Base_Assinaturas_8'
lv.base(B, S8, ['account_id', 'subscription_id', 'plan_tier', 'seats', 'mrr_amount', 'cobranca', 'trial', 'start_date', 'end_date', 'fim_para_conta', 'situacao_txt', 'eh_a_reducao_analisada'],
        {'start_date': DT, 'end_date': DT, 'fim_para_conta': DT, 'mrr_amount': USD})
lv.base('Base_Contas_8', I, ['ordem', 'account_id', 'industry', 'country', 'referral_source', 'signup_date', 'subscription_id', 'seats', 'cobranca', 'start_date', 'end_date', 'dias_de_vida', 'mrr_perdido', 'grupo', 'gemea', 'recuperou', 'eventos', 'motivos', 'tickets', 'tickets_antes_cadastro'],
        {'signup_date': DT, 'start_date': DT, 'end_date': DT, 'mrr_perdido': USD})
AC, ST, FM, MR = lv.R(B, 'account_id'), lv.R(B, 'start_date'), lv.R(B, 'fim_para_conta'), lv.R(B, 'mrr_amount')
HE, HR_, HP, HN = fill('0F766E'), fill('B45309'), fill('B91C1C'), fill('0E7C66')

# ---------- 1_As_8_Contas ----------
ws = w1; titulo(ws, 'As 8 contas Enterprise que perderam mais da metade do que pagavam numa redução',
    f'Entre as {N_SEGUIU} reduções Enterprise pagas em que a conta seguiu com outra Enterprise, estas 8 cortaram mais de 50% do valor mensal da conta de uma vez. Valores em MRR (US$ por mês). "Antes" = assinatura encerrada + outras ativas no dia; "depois" = outras ativas no dia.', 'V', 44)
hd = ['#', 'Conta', 'Indústria', 'País', 'Canal', 'Assinatura encerrada', 'Seats', 'Cobrança', 'Início', 'Fim', 'Dias de vida', 'MRR perdido', 'MRR da conta antes', 'MRR da conta depois', '% perdido', 'Grupo', 'Tinha gêmea (mesmos seats)?', 'MRR da conta em 31/12/2024', 'Recuperou?', 'Eventos de churn', 'Tickets', 'O que aconteceu']
H(ws, 4, hd, [HD, HD, HE, HE, HE, HX, HX, HX, HX, HX, HX, HR_, HX, HX, HP, HX, HX, HN, HN, HX, HX, HD], 52)
BC = 'Base_Contas_8'
def look(col, r): return f'=INDEX({lv.R(BC, col)},MATCH($B{r},{lv.R(BC, "account_id")},0))'
for i, row in I.iterrows():
    r = 5 + i
    C(ws, r, 1, int(row.ordem), bold=True); C(ws, r, 2, row.account_id, bold=True)
    for j, col, fm in [(3, 'industry', None), (4, 'country', None), (5, 'referral_source', None), (6, 'subscription_id', None), (7, 'seats', '#,##0'), (8, 'cobranca', None), (9, 'start_date', DT), (10, 'end_date', DT), (11, 'dias_de_vida', '#,##0'), (12, 'mrr_perdido', USD),
                       (16, 'grupo', None), (17, 'gemea', None), (19, 'recuperou', None), (20, 'eventos', '#,##0'), (21, 'tickets', '#,##0')]:
        C(ws, r, j, look(col, r), fm)
    C(ws, r, 13, f'=L{r}+SUMIFS({MR},{AC},$B{r},{ST},"<="&J{r},{FM},">"&J{r})', USD)
    C(ws, r, 14, f'=SUMIFS({MR},{AC},$B{r},{ST},"<="&J{r},{FM},">"&J{r})', USD)
    C(ws, r, 15, f'=L{r}/M{r}', '0%', True)
    C(ws, r, 18, f'=SUMIFS({MR},{AC},$B{r},{FM},">"&DATE(2024,12,31))', USD, True)
    C(ws, r, 22, row.explicacao); ws.row_dimensions[r].height = 150
rt = 5 + len(I); C(ws, rt, 2, 'Total', bold=True, fl=GRY)
for j in (12, 13, 14, 18): C(ws, rt, j, f'=SUM({L(j)}5:{L(j)}{rt-1})', USD, True, GRY)
C(ws, rt, 15, f'=L{rt}/M{rt}', '0%', True, GRY)
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add(f'S5:S{rt}', FormulaRule(formula=[f'LEFT(S5,3)<>"Não"'], fill=fill('DCFCE7')))
ws.conditional_formatting.add(f'S5:S{rt}', FormulaRule(formula=[f'LEFT(S5,3)="Não"'], fill=fill('FDE2E1')))
nota(ws, rt + 2, 'Leitura: das 8, só parte representa perda que ficou. Algumas contas recuperaram o valor (abriram outras assinaturas depois) e uma foi compra a mais desfeita em 3 dias. As que NÃO recuperaram são as prioritárias para conversa: a coluna "Recuperou?" separa os casos. Eventos de churn e tickets não explicam as reduções (datas fora do momento do cancelamento).', 'V', 44)
larguras(ws, [4, 11, 13, 7, 9, 11, 7, 8, 10, 10, 7, 11, 12, 12, 8, 22, 11, 12, 18, 8, 7, 90]); ws.freeze_panes = 'C5'

# ---------- 2_Linha_do_tempo ----------
ws = w2; titulo(ws, 'Linha do tempo das 8 contas: todas as assinaturas', 'Em destaque, a assinatura Enterprise cuja redução é analisada. "Situação" segue as regras da planilha 02 (Migração, Compra extra desfeita, Redução real, Saída do cliente, Ativa).', 'J')
hd = ['Conta', 'Assinatura', 'Plano', 'Seats', 'MRR', 'Cobrança', 'Trial', 'Início', 'Fim', 'Situação']
H(ws, 4, hd, [HD] * 10, 30); r = 5; HL = fill('FEF3C7')
for acct in CONTAS:
    x = S8[S8.account_id == acct]; ac = a[a.account_id == acct].iloc[0]
    ws.cell(r, 1, f'{acct} — {ac.industry} · {ac.country} · canal {ac.referral_source} · cadastro {dd(ac.signup_date)}').font = Fn(bold=True, size=10); r += 1
    for z in x.itertuples():
        fl = HL if z.eh_a_reducao_analisada == 'Sim' else None
        for j, v, fm in [(1, z.account_id, None), (2, z.subscription_id, None), (3, z.plan_tier, None), (4, int(z.seats), '#,##0'), (5, int(z.mrr_amount), USD), (6, z.cobranca, None), (7, z.trial, None),
                         (8, z.start_date.to_pydatetime(), DT), (9, None if pd.isna(z.end_date) else z.end_date.to_pydatetime(), DT), (10, z.situacao_txt, None)]:
            C(ws, r, j, v, fm, bold=(fl is not None), fl=fl)
        r += 1
    C(ws, r, 4, 'MRR em 31/12/2024:', bold=True); C(ws, r, 5, f'=SUMIFS({MR},{AC},"{acct}",{FM},">"&DATE(2024,12,31))', USD, True, GRY); r += 2
larguras(ws, [11, 11, 11, 7, 11, 9, 7, 11, 11, 22])

# ---------- 3_Perfil ----------
ws = w3; titulo(ws, 'Indústria, canal e país das 8 contas', f'Comparado com as {N_SEGUIU} reduções Enterprise pagas em que a conta seguiu com outra Enterprise e com todas as contas que tiveram Enterprise. Com 8 casos, nenhuma diferença pode ser considerada confiável.', 'G')
XA = X.merge(a[['account_id', 'industry', 'country', 'referral_source']], on='account_id')
ent_acc = a[a.account_id.isin(s[s.plan_tier == 'Enterprise'].account_id)]
r = 4
for col, lab in [('industry', 'Indústria'), ('referral_source', 'Canal'), ('country', 'País')]:
    H(ws, r, [lab, 'As 8 contas', 'MRR perdido pelas 8', f'% das 8', f'% nas {N_SEGUIU} reduções', '% das contas com Enterprise', 'Diferente? (8 x demais reduções)'], [HD, HP, HP, HP, HR_, HX, HX], 40); r += 1; r0 = r
    for v in sorted(ent_acc[col].unique()):
        k1 = int((I[col] == v).sum()); k2 = int((XA[col] == v).sum()) - k1
        C(ws, r, 1, v, bold=True); C(ws, r, 2, f'=COUNTIFS({lv.R(BC, col)},$A{r})', '#,##0'); C(ws, r, 3, f'=SUMIFS({lv.R(BC, "mrr_perdido")},{lv.R(BC, col)},$A{r})', USD)
        C(ws, r, 4, f'=B{r}/COUNTA({lv.R(BC, "account_id")})', '0%'); C(ws, r, 5, k1 + k2, None); ws.cell(r, 5).value = (k1 + k2) / N_SEGUIU; ws.cell(r, 5).number_format = '0%'
        C(ws, r, 6, float((ent_acc[col] == v).mean()), '0%'); C(ws, r, 7, conf(k1, len(I), k2, N_SEGUIU - len(I))); r += 1
    r += 1
nota(ws, r, 'Leitura: as 8 se concentram em FinTech (4) e DevTools (3), e no Reino Unido (4). É um sinal para olhar, não uma conclusão: com 8 casos qualquer concentração pode ser coincidência. As colunas E e F (valores calculados pelo programa) servem de comparação.', 'G', 40)
larguras(ws, [16, 10, 14, 9, 13, 14, 26])

# ---------- LEIA-ME ----------
rec = I.recuperou.str[:3].value_counts()
leiame(leia, [('AS 8 CONTAS ENTERPRISE PRIORITÁRIAS — perderam mais da metade do que pagavam numa redução', dict(bold=True, size=15)),
    ('Planilha apartada. Lê os CSVs originais sem alterá-los. Base: tabela de assinaturas (valores em MRR, US$ por mês).', dict(italic=True, color=GRAY)),
    ('', {}), ('COMO AS 8 FORAM ESCOLHIDAS', dict(bold=True, size=12)),
    (f'• Reduções de assinaturas Enterprise PAGAS (sem trial) em que a conta continuou com outra Enterprise: {N_SEGUIU}.', {}),
    ('• Para cada uma: % perdido = MRR da assinatura encerrada ÷ (ela + todas as outras assinaturas ativas da conta no dia). As 8 com mais de 50%.', {}),
    ('• Grupos: A = compra extra desfeita (até 30 dias, já havia Enterprise mais antiga); B = pacote antigo desligado havendo Enterprise mais nova; C = redução pura.', {}),
    ('', {}), ('O QUE MOSTRAM', dict(bold=True, size=12)),
    (f'• Total perdido de uma vez pelas 8: {fmt(int(I.mrr_perdido.sum()))}/mês.', {}),
    (f'• Recuperaram (pagam hoje mais do que antes da redução): {int(rec.get("Sim", 0))}. Compra a mais desfeita (não é perda): {int(rec.get("Com", 0))}. NÃO recuperaram: {int(rec.get("Não", 0))} — estas são a prioridade para conversa.', dict(bold=True)),
    ('• Três eram contratos ANUAIS encerrados antes de 12 meses (cancelamento antecipado).', {}),
    ('• Eventos de churn e tickets dessas contas não coincidem com o momento da redução (datas sem relação).', {}),
    ('• Perfil: FinTech e DevTools; Reino Unido e EUA; canais variados. Com 8 casos, nada disso é conclusivo.', {}),
    ('', {}), ('ABAS', dict(bold=True, size=12)),
    ('• 1_As_8_Contas: uma linha por conta, com indústria/canal/país, valores antes/depois/hoje (fórmulas), grupo, recuperação e o texto "O que aconteceu".', {}),
    ('• 2_Linha_do_tempo: todas as assinaturas de cada conta, com a reduzida em destaque.', {}),
    ('• 3_Perfil: indústria, canal e país das 8 contra as demais reduções.', {}),
    ('• Base_Assinaturas_8 e Base_Contas_8: dados usados pelas fórmulas.', {})])
lv.salvar(OUT); print('ok', OUT); print(I[['account_id', 'industry', 'country', 'referral_source', 'mrr_perdido', 'grupo', 'recuperou']].to_string())
