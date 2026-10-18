# Passo 2: gera o Excel com painéis, filtros e fórmulas
# Gerado com Claude (IA) - desafio G4 data-001-churn
# Pastas: G4_DESAFIO/dados (CSVs originais, só leitura) -> G4_DESAFIO/analise
from pathlib import Path
RAIZ = Path(__file__).resolve().parents[3]          # G4_DESAFIO
DADOS = str(RAIZ / 'dados') + '/'
TRAB = RAIZ / "analise" / "DADOS_ORIGINAIS" / "scripts" / "_trabalho"; TRAB.mkdir(parents=True, exist_ok=True)
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
B=pd.read_pickle(TRAB/'bases.pkl'); K,S,C,U,T=B['K'],B['S'],B['C'],B['U'],B['T']
OUT=str(RAIZ/'analise'/'DADOS_ORIGINAIS'/'RavenStack_Analise_Churn.xlsx')
FONT='Arial'; BLUE='2A78D6'; ORANGE='EB6834'; AQUA='1BAF7A'; GRAY='52514E'
F=lambda **k: Font(name=FONT,**k)
HDR_FILL=PatternFill('solid',fgColor='1F2937'); IN_FILL=PatternFill('solid',fgColor='FFF2CC'); TOT_FILL=PatternFill('solid',fgColor='EDEDED')
thin=Side(style='thin',color='D0D0D0'); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
USD='"US$ "#,##0;-"US$ "#,##0;"-"'; INT='#,##0;-#,##0;"-"'; PCT='0.0%;-0.0%;"-"'; DEC='#,##0.0;-#,##0.0;"-"'; DEC2='0.00'
wb=Workbook(); wb.remove(wb.active)
def S_(name): ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; return ws
leia=S_('LEIA-ME'); pm=S_('Painel_Mensal'); pseg=S_('Painel_Segmentos'); pf=S_('Painel_Features'); pc=S_('Painel_Comparativo')

# ---------- bases ----------
bases={}
LABELS={'Base_Contas':K,'Base_Assinaturas':S,'Base_Churn':C,'Base_Uso':U,'Base_Tickets':T}
DATEF={'signup_date','mes_cadastro','start_date','mes_inicio','end_date','mes_fim','end_date_calc','churn_date','mes_churn','usage_date','mes_uso','submitted_at','mes_ticket','closed_at','ultimo_churn'}
for name,df in LABELS.items():
    ws=wb.create_sheet(name); cols=list(df.columns)
    ws.append(cols)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v,float) and np.isnan(v)) or v is pd.NaT else (v.to_pydatetime() if isinstance(v,pd.Timestamp) else (v.item() if hasattr(v,'item') else v)) for v in row])
    n=len(df)+1
    for i,cname in enumerate(cols,1):
        ws.column_dimensions[L(i)].width=max(12,min(28,len(cname)+3))
        fmt=None
        if cname in DATEF: fmt='dd/mm/yyyy' if not cname.startswith('mes_') else 'mmm/yy'
        elif cname in('submitted_at','closed_at'): fmt='dd/mm/yyyy hh:mm'
        elif 'mrr' in cname or 'arr' in cname or 'refund' in cname: fmt=USD
        if fmt:
            for r in range(2,n+1): ws.cell(r,i).number_format=fmt
    tb=Table(displayName='t'+name.replace('Base_',''),ref=f'A1:{L(len(cols))}{n}')
    tb.tableStyleInfo=TableStyleInfo(name='TableStyleLight9',showRowStripes=True); ws.add_table(tb)
    ws.freeze_panes='B2'
    for c in ws[1]: c.font=F(bold=True,color='FFFFFF')
    bases[name]=(cols,n)
def R(base,col):
    cols,n=bases[base]; i=cols.index(col)+1; return f"{base}!${L(i)}$2:${L(i)}${n}"

# ---------- Listas ----------
ls=wb.create_sheet('Listas')
LIST={'Indústria':sorted(K.industry.unique()),'País':sorted(K.country.unique()),'Canal':sorted(K.referral_source.unique()),
      'Plano':['Basic','Pro','Enterprise'],'Cobrança':['monthly','annual'],'Definição de churn':['Evento de churn','churn_flag da conta']}
LREF={}
for j,(k,vals) in enumerate(LIST.items(),1):
    ls.cell(1,j,k).font=F(bold=True)
    v=vals if k=='Definição de churn' else ['(Todos)']+list(vals)
    for i,x in enumerate(v,2): ls.cell(i,j,x).font=F()
    LREF[k]=f"=Listas!${L(j)}$2:${L(j)}${len(v)+1}"
ls.sheet_state='visible'

# ---------- helpers ----------
def title(ws,t,sub):
    ws['A1']=t; ws['A1'].font=F(bold=True,size=16)
    ws['A2']=sub; ws['A2'].font=F(italic=True,color=GRAY,size=10)
def filt(ws,row,label,lst,note,default='(Todos)'):
    ws.cell(row,1,label).font=F(bold=True)
    c=ws.cell(row,2,default); c.fill=IN_FILL; c.font=F(color='0000FF',bold=True); c.border=BOX
    dv=DataValidation(type='list',formula1=LREF[lst],allow_blank=False); ws.add_data_validation(dv); dv.add(c)
    ws.cell(row,3,f'=IF(B{row}="(Todos)","*",B{row})').font=F(color='A0A0A0',size=8)
    ws.cell(row,4,note).font=F(italic=True,color=GRAY,size=9)
def header(ws,row,labels,start=1,height=45):
    for j,h in enumerate(labels,start):
        c=ws.cell(row,j,h); c.font=F(bold=True,color='FFFFFF',size=9); c.fill=HDR_FILL
        c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); c.border=BOX
    ws.row_dimensions[row].height=height
def put(ws,r,c,v,fmt=None,bold=False,fill=None):
    x=ws.cell(r,c,v); x.font=F(bold=bold,size=10); x.border=BOX
    if fmt: x.number_format=fmt
    if fill: x.fill=fill
    return x
def chart(kind,ws,ttl,cats,series,anchor,ytitle=None,fmt=None,w=17,h=8.5,colors=(BLUE,ORANGE,AQUA),titles=None):
    titles=list(titles or [])
    ch=BarChart() if kind=='bar' else LineChart()
    if kind=='bar': ch.type='col'; ch.grouping='clustered'; ch.gapWidth=60; ch.overlap=-10
    ch.title=ttl; ch.width=w; ch.height=h; ch.legend.position='b'
    if ytitle: ch.y_axis.title=ytitle
    if fmt: ch.y_axis.numFmt=fmt; ch.y_axis.number_format=fmt
    ch.y_axis.majorGridlines.spPr=GraphicalProperties(ln=LineProperties(solidFill='E1E0D9'))
    ch.x_axis.delete=False; ch.y_axis.delete=False
    for i,(ref,tl) in enumerate(series):
        ch.add_data(ref,titles_from_data=tl)
    ch.set_categories(cats)
    from openpyxl.chart.series import SeriesLabel
    for s_,(ref,tl) in zip(ch.series,series):
        if not tl and titles: s_.tx=SeriesLabel(v=titles.pop(0))
    for i,s in enumerate(ch.series):
        col=colors[i%len(colors)]
        if kind=='bar': s.graphicalProperties.solidFill=col; s.graphicalProperties.line.solidFill=col
        else: s.graphicalProperties.line.solidFill=col; s.graphicalProperties.line.width=22000; s.smooth=False; s.marker.symbol='none'
    ws.add_chart(ch,anchor); return ch

# ================= Painel_Mensal =================
ws=pm
title(ws,'Adesões, cancelamentos e churn mês a mês','Escolha os filtros nas células amarelas (lista suspensa). Tabela e gráficos recalculam automaticamente.')
filt(ws,4,'Indústria','Indústria','Aplica-se a todas as colunas')
filt(ws,5,'País','País','Aplica-se a todas as colunas')
filt(ws,6,'Canal de aquisição','Canal','Aplica-se a todas as colunas')
filt(ws,7,'Plano da assinatura','Plano','Só colunas de assinaturas, MRR e uso (D–J, O–Q). Eventos de churn e contas são por conta.')
filt(ws,8,'Cobrança (mensal/anual)','Cobrança','Só colunas de assinaturas, MRR e uso (D–J, O–Q).')
H=['Mês (data)','Mês','Novas contas (cadastro)','Novas assinaturas','MRR novo (US$)','Assinaturas encerradas','MRR cancelado (US$)','Saldo de MRR (novo − cancelado)','MRR ativo no início do mês','Churn de receita % (MRR cancelado ÷ MRR início)','Eventos de churn (contas)','MRR em risco nos eventos (US$)','Contas na base (início do mês)','Churn de contas % (eventos ÷ contas)','Uso total (usage_count)','Assinaturas ativas (início do mês)','Uso por assinatura ativa']
HR=11; header(ws,HR,H,height=60)
months=pd.date_range('2023-01-01','2024-12-01',freq='MS')
acc=lambda b: f'{R(b,"industry")},$C$4,{R(b,"country")},$C$5,{R(b,"referral_source")},$C$6'
subf=lambda b: acc(b)+f',{R(b,"plan_tier")},$C$7,{R(b,"billing_frequency")},$C$8'
BS='Base_Assinaturas'; BC='Base_Churn'; BK='Base_Contas'; BU='Base_Uso'; BT='Base_Tickets'
MN=['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
r0=HR+1
for i,m in enumerate(months):
    r=r0+i; a=f'$A{r}'
    put(ws,r,1,m.to_pydatetime(),'dd/mm/yyyy'); put(ws,r,2,f'{MN[m.month-1]}/{str(m.year)[2:]}')
    put(ws,r,3,f'=COUNTIFS({R(BK,"mes_cadastro")},{a},{acc(BK)})',INT)
    put(ws,r,4,f'=COUNTIFS({R(BS,"mes_inicio")},{a},{subf(BS)})',INT)
    put(ws,r,5,f'=SUMIFS({R(BS,"mrr_amount")},{R(BS,"mes_inicio")},{a},{subf(BS)})',USD)
    put(ws,r,6,f'=COUNTIFS({R(BS,"mes_fim")},{a},{subf(BS)})',INT)
    put(ws,r,7,f'=SUMIFS({R(BS,"mrr_amount")},{R(BS,"mes_fim")},{a},{subf(BS)})',USD)
    put(ws,r,8,f'=E{r}-G{r}',USD)
    put(ws,r,9,f'=SUMIFS({R(BS,"mrr_amount")},{R(BS,"start_date")},"<"&{a},{R(BS,"end_date_calc")},">="&{a},{subf(BS)})',USD)
    put(ws,r,10,f'=IF(I{r}>0,G{r}/I{r},"")',PCT)
    put(ws,r,11,f'=COUNTIFS({R(BC,"mes_churn")},{a},{acc(BC)})',INT)
    put(ws,r,12,f'=SUMIFS({R(BC,"mrr_em_risco")},{R(BC,"mes_churn")},{a},{acc(BC)})',USD)
    put(ws,r,13,f'=COUNTIFS({R(BK,"signup_date")},"<"&{a},{acc(BK)})',INT)
    put(ws,r,14,f'=IF(M{r}>0,K{r}/M{r},"")',PCT)
    put(ws,r,15,f'=SUMIFS({R(BU,"usage_count")},{R(BU,"mes_uso")},{a},{subf(BU)})',INT)
    put(ws,r,16,f'=COUNTIFS({R(BS,"start_date")},"<"&{a},{R(BS,"end_date_calc")},">="&{a},{subf(BS)})',INT)
    put(ws,r,17,f'=IF(P{r}>0,O{r}/P{r},"")',DEC)
rl=r0+len(months)-1; rt=rl+1
put(ws,rt,1,'',fill=TOT_FILL); put(ws,rt,2,'Total',bold=True,fill=TOT_FILL)
for c in [3,4,5,6,7,8,11,12,15]: put(ws,rt,c,f'=SUM({L(c)}{r0}:{L(c)}{rl})',USD if c in(5,7,8,12) else INT,bold=True,fill=TOT_FILL)
for c in [9,10,13,14,16,17]: put(ws,rt,c,'',fill=TOT_FILL)
ws.cell(rt+1,1,'MRR em risco = soma do MRR das assinaturas ativas da conta na data do evento de churn (os eventos não informam valor). Uso por assinatura: atenção, as datas de uso não são coerentes com as assinaturas (ver LEIA-ME).').font=F(italic=True,size=9,color=GRAY)
ws.column_dimensions['A'].hidden=False; ws.column_dimensions['A'].width=12; ws.column_dimensions['B'].width=14
for c in range(3,18): ws.column_dimensions[L(c)].width=14
ws.column_dimensions['A'].width=24
ws.freeze_panes=f'C{r0}'
cats=Reference(ws,min_col=2,min_row=r0,max_row=rl)
ser=lambda c: (Reference(ws,min_col=c,min_row=HR,max_row=rl),True)
cr=rt+3
ws.cell(cr-1,1,'Gráficos (seguem os filtros acima)').font=F(bold=True,size=12)
chart('bar',ws,'Adesões vs cancelamentos (qtd. de assinaturas)',cats,[ser(4),ser(6)],f'A{cr}',fmt='#,##0')
chart('bar',ws,'MRR novo vs MRR cancelado (US$)',cats,[ser(5),ser(7)],f'I{cr}',fmt='#,##0')
catsA=Reference(ws,min_col=2,min_row=r0+3,max_row=rl); catsJ=Reference(ws,min_col=2,min_row=r0+5,max_row=rl)
serA=lambda c,s0: (Reference(ws,min_col=c,min_row=s0,max_row=rl),False)
chart('line',ws,'Taxa de churn mensal (%), a partir de abr/23',catsA,[serA(10,r0+3),serA(14,r0+3)],f'A{cr+18}',fmt='0%',colors=(ORANGE,BLUE),titles=['Churn de receita %','Churn de contas %'])
chart('bar',ws,'Novas contas vs eventos de churn (contas)',cats,[ser(3),ser(11)],f'I{cr+18}',fmt='#,##0')
chart('line',ws,'Uso por assinatura ativa, a partir de jun/23',catsJ,[serA(17,r0+5)],f'A{cr+36}',fmt='#,##0',titles=['Uso por assinatura ativa'])
chart('bar',ws,'MRR cancelado (assinaturas) vs MRR em risco (eventos de churn)',cats,[ser(7),ser(12)],f'I{cr+36}',fmt='#,##0',colors=(ORANGE,'6250D6'))

# ================= Painel_Segmentos =================
ws=pseg
title(ws,'Churn por segmento e por motivo','Visão acumulada jan/23–dez/24. "Contas com evento" = conta com ao menos 1 registro em churn_events; "churn_flag" = marcação da tabela accounts.')
dims=[('Indústria','industry',LIST['Indústria']),('Plano inicial da conta','plano_inicial_conta',LIST['Plano']),('País','country',LIST['País']),('Canal de aquisição','referral_source',LIST['Canal'])]
HS=['Segmento','Contas','Contas com evento de churn','% contas com evento','Contas com churn_flag','% contas com churn_flag','Eventos de churn','MRR em risco nos eventos (US$)','MRR ativo atual (US$)','MRR cancelado – assinaturas encerradas (US$)','% MRR cancelado (cancelado ÷ ativo+cancelado)','MRR ativo médio por conta (US$)']
row=4; seg_blocks=[]
for lab,col,vals in dims:
    ws.cell(row,1,lab).font=F(bold=True,size=12); row+=1
    header(ws,row,HS); hr=row; row+=1; first=row
    for v in vals:
        k=f'$A{row}'
        put(ws,row,1,v,bold=True)
        put(ws,row,2,f'=COUNTIFS({R(BK,col)},{k})',INT)
        put(ws,row,3,f'=COUNTIFS({R(BK,col)},{k},{R(BK,"conta_com_evento_churn")},"Sim")',INT)
        put(ws,row,4,f'=IF(B{row}>0,C{row}/B{row},"")',PCT)
        put(ws,row,5,f'=COUNTIFS({R(BK,col)},{k},{R(BK,"churn_flag_conta")},"Sim")',INT)
        put(ws,row,6,f'=IF(B{row}>0,E{row}/B{row},"")',PCT)
        put(ws,row,7,f'=COUNTIFS({R(BC,col)},{k})',INT)
        put(ws,row,8,f'=SUMIFS({R(BC,"mrr_em_risco")},{R(BC,col)},{k})',USD)
        put(ws,row,9,f'=SUMIFS({R(BK,"mrr_ativo_atual")},{R(BK,col)},{k})',USD)
        put(ws,row,10,f'=SUMIFS({R(BS,"mrr_amount")},{R(BS,col)},{k},{R(BS,"status")},"Encerrada")',USD)
        put(ws,row,11,f'=IF(I{row}+J{row}>0,J{row}/(I{row}+J{row}),"")',PCT)
        put(ws,row,12,f'=IF(B{row}>0,I{row}/B{row},"")',USD)
        row+=1
    last=row-1
    put(ws,row,1,'Total',bold=True,fill=TOT_FILL)
    for c in [2,3,5,7,8,9,10]: put(ws,row,c,f'=SUM({L(c)}{first}:{L(c)}{last})',USD if c in(8,9,10) else INT,bold=True,fill=TOT_FILL)
    put(ws,row,4,f'=C{row}/B{row}',PCT,True,TOT_FILL); put(ws,row,6,f'=E{row}/B{row}',PCT,True,TOT_FILL)
    put(ws,row,11,f'=J{row}/(I{row}+J{row})',PCT,True,TOT_FILL); put(ws,row,12,f'=I{row}/B{row}',USD,True,TOT_FILL)
    seg_blocks.append((lab,hr,first,last)); row+=3
ws.column_dimensions['A'].width=24
for c in range(2,13): ws.column_dimensions[L(c)].width=15
# charts para segmentos à direita
cy=4
for lab,hr,first,last in seg_blocks:
    cats=Reference(ws,min_col=1,min_row=first,max_row=last)
    chart('bar',ws,f'{lab}: % contas com evento de churn vs % MRR cancelado',cats,[(Reference(ws,min_col=4,min_row=hr,max_row=last),True),(Reference(ws,min_col=11,min_row=hr,max_row=last),True)],f'N{cy}',fmt='0%',w=16,h=7.5)
    cy+=16
# motivos
row+=1; ws.cell(row,1,'Motivos de churn (reason_code)').font=F(bold=True,size=12); row+=1
reasons=sorted(C.reason_code.unique()); inds=LIST['Indústria']
HM=['Motivo','Eventos','% dos eventos','MRR em risco (US$)','Refund total (US$)','Eventos de reativação']+[f'Eventos – {i}' for i in inds]
header(ws,row,HM); mhr=row; row+=1; mf=row
for v in reasons:
    k=f'$A{row}'
    put(ws,row,1,v,bold=True)
    put(ws,row,2,f'=COUNTIFS({R(BC,"reason_code")},{k})',INT)
    put(ws,row,3,f'=B{row}/SUM($B${mf}:$B${mf+len(reasons)-1})',PCT)
    put(ws,row,4,f'=SUMIFS({R(BC,"mrr_em_risco")},{R(BC,"reason_code")},{k})',USD)
    put(ws,row,5,f'=SUMIFS({R(BC,"refund_amount_usd")},{R(BC,"reason_code")},{k})',USD)
    put(ws,row,6,f'=COUNTIFS({R(BC,"reason_code")},{k},{R(BC,"is_reactivation")},"Sim")',INT)
    for j,ind in enumerate(inds):
        put(ws,row,7+j,f'=COUNTIFS({R(BC,"reason_code")},{k},{R(BC,"industry")},"{ind}")',INT)
    row+=1
ml=row-1
chart('bar',ws,'Eventos de churn por motivo',Reference(ws,min_col=1,min_row=mf,max_row=ml),[(Reference(ws,min_col=2,min_row=mhr,max_row=ml),True)],f'N{cy}',fmt='#,##0',w=16,h=7.5)
cy+=16
chart('bar',ws,'MRR em risco por motivo (US$)',Reference(ws,min_col=1,min_row=mf,max_row=ml),[(Reference(ws,min_col=4,min_row=mhr,max_row=ml),True)],f'N{cy}',fmt='#,##0',w=16,h=7.5,colors=('6250D6',))
# motivo x feedback
row+=2; ws.cell(row,1,'Motivo (reason_code) × comentário do cliente (feedback_text)').font=F(bold=True,size=12); row+=1
fbs=sorted(C.feedback_text.unique())
header(ws,row,['Motivo']+fbs); row+=1
for v in reasons:
    put(ws,row,1,v,bold=True)
    for j,fb in enumerate(fbs): put(ws,row,2+j,f'=COUNTIFS({R(BC,"reason_code")},$A{row},{R(BC,"feedback_text")},"{fb}")',INT)
    row+=1
ws.cell(row,1,'Repare se o comentário do cliente bate com o motivo registrado: é um teste de qualidade do reason_code.').font=F(italic=True,size=9,color=GRAY)
ws.freeze_panes='B4'

# ================= Painel_Features =================
ws=pf
title(ws,'Features: uso, tempo, erros e relação com churn','Cada linha = uma feature. Colunas H–Q comparam contas COM churn vs SEM churn, conforme a definição escolhida em B7.')
filt(ws,4,'Indústria','Indústria','')
filt(ws,5,'Plano da assinatura','Plano','')
filt(ws,6,'Cobrança (mensal/anual)','Cobrança','')
filt(ws,7,'Definição de churn','Definição de churn','"Evento de churn" = conta tem registro em churn_events. "churn_flag da conta" = coluna churn_flag de accounts.',default='Evento de churn')
ws['C7']=None
uf=f'{R(BU,"industry")},$C$4,{R(BU,"plan_tier")},$C$5,{R(BU,"billing_frequency")},$C$6'
def g2(fn,rng,crit,grp):
    a=f'{fn}({rng+"," if rng else ""}{crit},{R(BU,"conta_com_evento_churn")},"{grp}",{uf})'
    b=f'{fn}({rng+"," if rng else ""}{crit},{R(BU,"churn_flag_conta")},"{grp}",{uf})'
    return f'IF($B$7="Evento de churn",{a},{b})'
HF=['Feature','Registros de uso','Uso total (usage_count)','Tempo total (horas)','Erros total','Erros por 100 usos','% registros em beta','Registros – contas COM churn','Registros – contas SEM churn','% registros de contas com churn','Uso médio por registro – COM churn','Uso médio por registro – SEM churn','Minutos médios por registro – COM churn','Minutos médios por registro – SEM churn','Erros por 100 usos – COM churn','Erros por 100 usos – SEM churn','Diferença de erros (COM − SEM)']
HR=10; header(ws,HR,HF,height=60); r0=HR+1
feats=[f'feature_{i}' for i in range(1,41)]
fc=R(BU,'feature_name')
for i,fe in enumerate(feats):
    r=r0+i; k=f'{fc},$A{r}'
    put(ws,r,1,fe,bold=True)
    put(ws,r,2,f'=COUNTIFS({k},{uf})',INT)
    put(ws,r,3,f'=SUMIFS({R(BU,"usage_count")},{k},{uf})',INT)
    put(ws,r,4,f'=SUMIFS({R(BU,"usage_duration_secs")},{k},{uf})/3600',DEC)
    put(ws,r,5,f'=SUMIFS({R(BU,"error_count")},{k},{uf})',INT)
    put(ws,r,6,f'=IF(C{r}>0,E{r}/C{r}*100,"")',DEC2)
    put(ws,r,7,f'=IF(B{r}>0,COUNTIFS({k},{R(BU,"is_beta_feature")},"Sim",{uf})/B{r},"")',PCT)
    put(ws,r,8,'='+g2('COUNTIFS','',k,'Sim'),INT)
    put(ws,r,9,'='+g2('COUNTIFS','',k,'Não'),INT)
    put(ws,r,10,f'=IF(B{r}>0,H{r}/B{r},"")',PCT)
    put(ws,r,11,f'=IF(H{r}>0,{g2("SUMIFS",R(BU,"usage_count"),k,"Sim")}/H{r},"")',DEC2)
    put(ws,r,12,f'=IF(I{r}>0,{g2("SUMIFS",R(BU,"usage_count"),k,"Não")}/I{r},"")',DEC2)
    put(ws,r,13,f'=IF(H{r}>0,{g2("SUMIFS",R(BU,"duracao_min"),k,"Sim")}/H{r},"")',DEC)
    put(ws,r,14,f'=IF(I{r}>0,{g2("SUMIFS",R(BU,"duracao_min"),k,"Não")}/I{r},"")',DEC)
    put(ws,r,15,f'=IFERROR({g2("SUMIFS",R(BU,"error_count"),k,"Sim")}/{g2("SUMIFS",R(BU,"usage_count"),k,"Sim")}*100,"")',DEC2)
    put(ws,r,16,f'=IFERROR({g2("SUMIFS",R(BU,"error_count"),k,"Não")}/{g2("SUMIFS",R(BU,"usage_count"),k,"Não")}*100,"")',DEC2)
    put(ws,r,17,f'=IF(AND(ISNUMBER(O{r}),ISNUMBER(P{r})),O{r}-P{r},"")',DEC2)
rl=r0+39; rt=rl+1
put(ws,rt,1,'Total',bold=True,fill=TOT_FILL)
for c in [2,3,4,5,8,9]: put(ws,rt,c,f'=SUM({L(c)}{r0}:{L(c)}{rl})',DEC if c==4 else INT,True,TOT_FILL)
put(ws,rt,6,f'=IF(C{rt}>0,E{rt}/C{rt}*100,"")',DEC2,True,TOT_FILL)
put(ws,rt,10,f'=IF(B{rt}>0,H{rt}/B{rt},"")',PCT,True,TOT_FILL)
put(ws,rt,11,f'=IF(H{rt}>0,SUMPRODUCT(K{r0}:K{rl},H{r0}:H{rl})/H{rt},"")',DEC2,True,TOT_FILL)
put(ws,rt,12,f'=IF(I{rt}>0,SUMPRODUCT(L{r0}:L{rl},I{r0}:I{rl})/I{rt},"")',DEC2,True,TOT_FILL)
put(ws,rt,13,f'=IF(H{rt}>0,SUMPRODUCT(M{r0}:M{rl},H{r0}:H{rl})/H{rt},"")',DEC,True,TOT_FILL)
put(ws,rt,14,f'=IF(I{rt}>0,SUMPRODUCT(N{r0}:N{rl},I{r0}:I{rl})/I{rt},"")',DEC,True,TOT_FILL)
for c in (7,15,16,17): put(ws,rt,c,'',fill=TOT_FILL)
ws.cell(rt+1,1,'Erros por 100 usos = error_count ÷ usage_count × 100. "Registro" = uma linha de feature_usage (uso diário de uma feature numa assinatura).').font=F(italic=True,size=9,color=GRAY)
ws.column_dimensions['A'].width=24
for c in range(2,18): ws.column_dimensions[L(c)].width=13
ws.freeze_panes=f'B{r0}'
cats=Reference(ws,min_col=1,min_row=r0,max_row=rl)
cr=rt+4; ws.cell(cr-1,1,'Gráficos (seguem os filtros acima)').font=F(bold=True,size=12)
chart('bar',ws,'Uso total por feature (usage_count)',cats,[(Reference(ws,min_col=3,min_row=HR,max_row=rl),True)],f'A{cr}',fmt='#,##0',w=34,h=8)
chart('bar',ws,'Erros por 100 usos: contas COM vs SEM churn',cats,[(Reference(ws,min_col=15,min_row=HR,max_row=rl),True),(Reference(ws,min_col=16,min_row=HR,max_row=rl),True)],f'A{cr+17}',fmt='0.0',w=34,h=8,colors=(ORANGE,BLUE))
chart('bar',ws,'Minutos médios por registro: contas COM vs SEM churn',cats,[(Reference(ws,min_col=13,min_row=HR,max_row=rl),True),(Reference(ws,min_col=14,min_row=HR,max_row=rl),True)],f'A{cr+34}',fmt='#,##0',w=34,h=8,colors=(ORANGE,BLUE))
chart('bar',ws,'% dos registros de uso vindos de contas com churn',cats,[(Reference(ws,min_col=10,min_row=HR,max_row=rl),True)],f'A{cr+51}',fmt='0%',w=34,h=8,colors=(ORANGE,))

# ================= Painel_Comparativo =================
ws=pc
title(ws,'Contas COM churn vs SEM churn: suporte, uso e receita','Médias por grupo. Use os filtros para isolar um segmento; a definição de churn muda o grupo.')
filt(ws,4,'Indústria','Indústria','')
filt(ws,5,'Plano inicial da conta','Plano','')
filt(ws,6,'País','País','')
filt(ws,7,'Definição de churn','Definição de churn','',default='Evento de churn'); ws['C7']=None
def kf(b): return f'{R(b,"industry")},$C$4,{R(b,"plano_inicial_conta")},$C$5,{R(b,"country")},$C$6'
def G(b,fn,rng,grp,extra=''):
    a=f'{fn}({rng+"," if rng else ""}{R(b,"conta_com_evento_churn")},"{grp}",{kf(b)}{extra})'
    c=f'{fn}({rng+"," if rng else ""}{R(b,"churn_flag_conta")},"{grp}",{kf(b)}{extra})'
    return f'IF($B$7="Evento de churn",{a},{c})'
header(ws,9,['Indicador','COM churn','SEM churn','Diferença (COM − SEM)','Como ler']); r=10
rows=[
 ('Contas',lambda g:'='+G(BK,'COUNTIFS','',g),INT,'Tamanho de cada grupo'),
 ('MRR ativo médio por conta (US$)',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"mrr_ativo_atual"),g)}/{G(BK,"COUNTIFS","",g)},"")',USD,'Receita atual de cada grupo'),
 ('MRR cancelado médio por conta (US$)',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"mrr_cancelado_total"),g)}/{G(BK,"COUNTIFS","",g)},"")',USD,'Soma das assinaturas encerradas'),
 ('Assinaturas por conta',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"qtd_assinaturas"),g)}/{G(BK,"COUNTIFS","",g)},"")',DEC2,''),
 ('Upgrades por conta',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"qtd_upgrades"),g)}/{G(BK,"COUNTIFS","",g)},"")',DEC2,'Assinaturas com upgrade_flag'),
 ('Downgrades por conta',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"qtd_downgrades"),g)}/{G(BK,"COUNTIFS","",g)},"")',DEC2,'Assinaturas com downgrade_flag'),
 ('Uso total por conta (usage_count)',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"uso_count_total"),g)}/{G(BK,"COUNTIFS","",g)},"")',DEC,''),
 ('Horas de uso por conta',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"uso_horas"),g)}/{G(BK,"COUNTIFS","",g)},"")',DEC,''),
 ('Erros por 100 usos',lambda g:f'=IFERROR({G(BK,"SUMIFS",R(BK,"erros_total"),g)}/{G(BK,"SUMIFS",R(BK,"uso_count_total"),g)}*100,"")',DEC2,''),
 ('Tickets por conta',lambda g:f'=IFERROR({G(BT,"COUNTIFS","",g)}/{G(BK,"COUNTIFS","",g)},"")',DEC2,'Carga de suporte'),
 ('Satisfação média (1–5)',lambda g:f'=IFERROR({G(BT,"SUMIFS",R(BT,"satisfaction_score"),g)}/{G(BT,"COUNTIFS","",g,","+R(BT,"satisfacao_respondida")+",\"Sim\"")},"")',DEC2,'Só tickets com resposta de satisfação'),
 ('% tickets com satisfação respondida',lambda g:f'=IFERROR({G(BT,"COUNTIFS","",g,","+R(BT,"satisfacao_respondida")+",\"Sim\"")}/{G(BT,"COUNTIFS","",g)},"")',PCT,'Quem não responde também é sinal'),
 ('Tempo médio de resolução (h)',lambda g:f'=IFERROR({G(BT,"SUMIFS",R(BT,"resolution_time_hours"),g)}/{G(BT,"COUNTIFS","",g)},"")',DEC,''),
 ('Primeira resposta média (min)',lambda g:f'=IFERROR({G(BT,"SUMIFS",R(BT,"first_response_time_minutes"),g)}/{G(BT,"COUNTIFS","",g)},"")',DEC,''),
 ('% tickets escalados',lambda g:f'=IFERROR({G(BT,"COUNTIFS","",g,","+R(BT,"escalation_flag")+",\"Sim\"")}/{G(BT,"COUNTIFS","",g)},"")',PCT,''),
 ('% tickets urgent',lambda g:f'=IFERROR({G(BT,"COUNTIFS","",g,","+R(BT,"priority")+",\"urgent\"")}/{G(BT,"COUNTIFS","",g)},"")',PCT,''),
 ('% tickets high',lambda g:f'=IFERROR({G(BT,"COUNTIFS","",g,","+R(BT,"priority")+",\"high\"")}/{G(BT,"COUNTIFS","",g)},"")',PCT,''),
]
for lab,fn,fmt,note in rows:
    put(ws,r,1,lab,bold=True); put(ws,r,2,fn('Sim'),fmt); put(ws,r,3,fn('Não'),fmt)
    put(ws,r,4,f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}-C{r},"")',fmt)
    x=put(ws,r,5,note); x.font=F(italic=True,size=9,color=GRAY); r+=1
ws.column_dimensions['A'].width=38
for c in 'BCD': ws.column_dimensions[c].width=16
ws.column_dimensions['E'].width=40
ws.freeze_panes='A10'

# ================= LEIA-ME =================
ws=leia; ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=120
lines=[('RavenStack: base de análise de churn',dict(bold=True,size=16)),
('Gerado a partir dos 5 CSVs em G4_DESAFIO\\dados (cópia; os originais não foram alterados). Todas as abas Painel_* usam fórmulas sobre as abas Base_*.',dict(italic=True,color=GRAY)),
('',{}),('COMO USAR',dict(bold=True,size=12)),
('• Células amarelas com texto azul = filtros (lista suspensa). "(Todos)" desliga o filtro. Tabelas e gráficos recalculam sozinhos.',{}),
('• As abas Base_* são Tabelas do Excel: dá para criar suas próprias tabelas dinâmicas e segmentações (Inserir > Tabela Dinâmica, escolhendo a tabela tContas, tAssinaturas, tChurn, tUso ou tTickets).',{}),
('• Painel_Mensal: adesões, cancelamentos, MRR e taxas de churn mês a mês.   • Painel_Segmentos: churn por indústria, plano, país, canal e motivo.',{}),
('• Painel_Features: uso/tempo/erros por feature, contas COM vs SEM churn.   • Painel_Comparativo: suporte, uso e receita de contas COM vs SEM churn.',{}),
('',{}),('DEFINIÇÕES (as três visões de churn)',dict(bold=True,size=12)),
('1) Eventos de churn (churn_events): 600 eventos em 352 contas, com data e motivo. Não trazem valor; usamos "MRR em risco" = MRR das assinaturas ativas da conta na data do evento.',{}),
('2) Churn de receita (assinaturas encerradas): 486 assinaturas com end_date. MRR cancelado = mrr_amount dessas assinaturas no mês do end_date.',{}),
('3) churn_flag da conta (accounts): 110 contas marcadas, sem data. Usado como rótulo/filtro de comparação (Painel_Features e Painel_Comparativo permitem alternar).',{}),
('• Adesão: "novas contas" = signup_date; "novas assinaturas" = start_date. • Taxa de churn de receita = MRR cancelado no mês ÷ MRR ativo no 1º dia do mês.',{}),
('• Taxa de churn de contas = eventos de churn no mês ÷ contas cadastradas antes do 1º dia do mês.',{}),
('',{}),('LIMITAÇÕES DOS DADOS (verificadas, base sintética)',dict(bold=True,size=12)),
('• As 3 marcações de churn não batem: 277 contas têm evento de churn com churn_flag = False; 96% das contas tinham assinaturas ativas no dia do último evento; só 7,8% dos eventos coincidem (±7 dias) com o fim de alguma assinatura.',{}),
('• Cada conta tem ~10 assinaturas (≈9 ativas simultâneas, misturando planos): assinatura = linha de contrato. MRR por seat é fixo: Basic 19, Pro 49, Enterprise 199 (US$). plan_tier da conta não se relaciona com o das assinaturas.',{}),
('• 77% dos registros de uso têm data anterior ao início da assinatura (53% anteriores ao cadastro da conta). Use o uso agregado por conta/feature; evite análises de "uso antes do churn" pela data.',{}),
('• Trials (is_trial) têm MRR = 0 e estão incluídos nas contagens de assinaturas.',{}),
]
for i,(tx,st) in enumerate(lines,1):
    c=ws.cell(i,2,tx); c.font=F(**st) if st else F(); c.alignment=Alignment(wrap_text=True,vertical='top')
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font and c.font.name!=FONT: c.font=Font(name=FONT,bold=c.font.bold,italic=c.font.italic,size=c.font.size,color=c.font.color)
wb.calculation.fullCalcOnLoad=True
wb.active=1
wb.save(OUT); print('ok')
